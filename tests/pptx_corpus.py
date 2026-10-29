"""A short-text, editable 12-slide enterprise-template acceptance deck."""
from pathlib import Path
import io
import zipfile


def make_template(root):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from lxml import etree
    root=Path(root);root.mkdir(parents=True,exist_ok=True);file=root/'enterprise-template.pptx'
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6]);brand=slide.shapes.add_textbox(Inches(.8),Inches(.18),Inches(5),Inches(.28));brand.text='NORTHSTAR  /  企业技术汇报'
    for run in brand.text_frame.paragraphs[0].runs:run.font.name='Noto Sans CJK SC';run.font.size=Pt(11);run.font.color.rgb=RGBColor.from_string('246BCE')
    prs.save(file)
    with zipfile.ZipFile(file) as z:parts={name:z.read(name) for name in z.namelist()}
    theme=etree.fromstring(parts['ppt/theme/theme1.xml']);ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    theme.xpath('//a:clrScheme/a:accent1/a:srgbClr',namespaces=ns)[0].set('val','246BCE')
    theme.xpath('//a:fontScheme/a:minorFont/a:ea',namespaces=ns)[0].set('typeface','Noto Sans CJK SC')
    parts['ppt/theme/theme1.xml']=etree.tostring(theme,xml_declaration=True,encoding='UTF-8')
    with zipfile.ZipFile(file,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    return file


def element(kind,sid,eid,**values):return dict(operation='element.add_'+kind,slide_id=sid,element=dict(id=eid,**values))


def title(sid,text):return element('text',sid,'title',text=text,style='slide_title',anchor='title')


def cards(sid,labels,bodies,prefix='card'):
    result=[]
    for i,(label,body) in enumerate(zip(labels,bodies)):
        x=.06+i*.31
        result.extend([
            element('shape',sid,f'{prefix}-bg-{i}',shape='rounded_rectangle',fill_token='card',layer='background',geometry=dict(x=x,y=.28,w=.26,h=.52)),
            element('text',sid,f'{prefix}-title-{i}',text=label,style='card_title',geometry=dict(x=x+.02,y=.32,w=.22,h=.1)),
            element('text',sid,f'{prefix}-body-{i}',text=body,style='body',geometry=dict(x=x+.02,y=.48,w=.22,h=.28)),
        ])
    for i in range(2):
        result.append(element('shape',sid,f'{prefix}-arrow-{i}',shape='arrow',fill_token='accent',layer='decoration',geometry=dict(x=.325+i*.31,y=.49,w=.035,h=.065)))
    return result


def plan(first_slide):
    result=[dict(operation='presentation.set_theme',theme='slate_dark')]
    cover=first_slide
    result.extend([dict(operation='slide.set_layout',slide_id=cover,layout='cover'),element('text',cover,'cover-title',text='AI Agent 工作台',style='cover_title',anchor='title'),element('text',cover,'cover-subtitle',text='平台建设方案与技术路线',style='body',anchor='subtitle'),element('text',cover,'cover-footer',text='2026 · 技术方案 / 项目汇报',style='caption',anchor='footer')])
    titles=['建设背景与业务目标','平台总体架构','核心能力与交付边界','执行与审批流程','技术路线与阶段计划','方案对比与选型依据','建设效益与交付指标','实施路径与里程碑','主要风险与保障机制','运营体系与责任分工','总结与下一步安排']
    layouts=['title_body','three_columns','comparison','title_body','data','title_body','data','three_columns','comparison','title_body','section']
    for number,(text,layout) in enumerate(zip(titles,layouts),2):
        sid=f'slide-{number:02d}'
        result.extend([dict(operation='slide.add',slide_id=sid,layout=layout),title(sid,text)])
        if number==2:result.append(element('text',sid,'background-body',anchor='body',text='统一入口：减少重复操作与信息分散\n过程可控：明确审批边界与交付状态\n成果可用：保留可编辑的真实业务文件',style='body'))
        elif number==3:result.extend(cards(sid,['交互层','执行层','工具资源层'],['Web 工作台\n消息与附件\n只读预览','主 Agent 驱动\n工具安全边界\n持久化会话','文档与演示\n数据与服务\n版本化成果'],prefix='architecture'))
        elif number==4:
            result.extend([element('text',sid,'left',anchor='left',text='交互与控制\n\n多轮对话与追加指令\n工具审批与暂停恢复\n历史记录与版本回溯',style='body'),element('text',sid,'right',anchor='right',text='工具与交付\n\n原子化文件操作\n页面质量检查\n真实可编辑的成果',style='body')])
        elif number==5:result.append(element('text',sid,'process-list',anchor='body',text='1. 接收用户目标\n2. 执行受控工具\n3. 检查并交付成果',style='body'))
        elif number==6:
            result.extend([element('chart',sid,'roadmap-chart',anchor='chart',chart=dict(type='line',categories=['Q1','Q2','Q3'],series=[dict(name='Demo progress',values=[35,70,100])])),element('text',sid,'roadmap-note',anchor='note',text='需求确认\n试点验证\n分批推广\n\n示例进度数据',style='body')])
        elif number==7:result.append(element('table',sid,'comparison-table',geometry=dict(x=.08,y=.28,w=.84,h=.48),rows=[['比较维度','现有方式','目标方案'],['执行过程','单次批量交付','原子化增量提交'],['反馈处理','完成后再修改','安全边界即时纠偏'],['成果验收','文件生成即结束','结构与视觉双检查']]))
        elif number==8:result.extend([element('chart',sid,'benefit-chart',anchor='chart',chart=dict(type='bar',categories=['Cycle','Coverage','Quality','Reuse'],series=[dict(name='Demo score',values=[70,85,90,80])])),element('text',sid,'benefit-note',anchor='note',text='图表为演示数据\n用于验证原生图表\n不代表真实业务指标',style='body')])
        elif number==9:result.extend(cards(sid,['试点验证','分批推广','持续运营'],['选择高频业务\n确定验收口径\n记录改进问题','完善培训指引\n分部门切换\n保留回退安排','跟踪服务指标\n维护规则版本\n持续业务复盘'],prefix='milestone'))
        elif number==10:result.extend([element('text',sid,'risk-left',anchor='left',text='主要风险\n\n需求范围持续变化\n模板与字体不一致\n外部服务不稳定',style='body'),element('text',sid,'risk-right',anchor='right',text='保障措施\n\n阶段目标与版本管理\n真实渲染与页面检查\n超时恢复与人工介入',style='body')])
        elif number==11:result.append(element('text',sid,'operations-body',anchor='body',text='业务负责人：确认目标与验收依据\n交付负责人：跟踪任务与问题闭环\n运营负责人：维护规则与服务质量',style='body'))
        else:result.extend([element('text',sid,'summary',anchor='subtitle',text='以受控执行连接业务目标与可用成果',style='body'),element('text',sid,'next',anchor='footer',text='下一步：完成试点、复核验收、分批推广',style='caption')])
        if number<12:result.append(element('text',sid,'page-number',anchor='footer',text=f'AI AGENT PLATFORM · 技术汇报                                      {number:02d}',style='caption'))
    return result
