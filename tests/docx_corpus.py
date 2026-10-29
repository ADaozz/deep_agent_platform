"""Formatting fixtures: short business content, never text padded to fill pages."""
from pathlib import Path

FORMAT_REQUIREMENTS = {
    'paper': 'A4 portrait, 210 × 297 mm',
    'margins': '25.4 mm on all four sides',
    'fonts': 'Noto Sans CJK SC for Chinese; Liberation Sans for Latin',
    'title': '20 pt bold, centered',
    'heading1': '16 pt bold, keep with next; 18 pt before / 8 pt after',
    'heading2': '13.5 pt bold, keep with next; 13 pt before / 6 pt after',
    'body': '11 pt, 1.5 line spacing, 6 pt after; every paragraph ≤ 180 characters',
    'tables': '3 columns with explicit widths, gray bold repeating header, 0.5 pt borders, unsplit rows/table',
    'images': 'PNG business workflow, 12 cm width, preserve aspect ratio, centered caption',
    'pagination': 'Natural pagination only; no inserted page breaks or filler paragraphs',
}

SUBJECTS = [
    ('项目背景与服务目标', '统一办理入口', '建立面向业务人员的统一工作入口，集中展示待办事项、办理规则和所需材料。以减少重复录入、缩短办理周期和明确责任边界作为首批建设目标。'),
    ('业务流程与职责分工', '明确协同责任', '由业务部门确认流程口径，运营团队维护办理规则，服务团队跟踪用户反馈。跨部门事项采用共同验收，避免在交接环节出现责任空白。'),
    ('数据治理与共享机制', '建立共享目录', '按照业务用途梳理数据目录，明确数据提供方、使用范围和更新周期。对关键字段建立质量检查，发现缺项时由责任部门补充确认。'),
    ('用户体验与服务渠道', '简化日常办理', '优先优化高频事项的提交、查询和反馈流程。保留必要的人工服务渠道，并通过可理解的状态说明，让用户清楚下一步需要采取的行动。'),
    ('试点范围与推广路径', '先试点后推广', '选择办理量适中、规则清晰的业务作为试点，安排固定联系人收集问题。试点稳定后，再按部门和事项分批推广，避免同时切换带来业务中断。'),
    ('风险管理与应急安排', '保障业务连续', '建立问题分级、升级处理和应急回退机制。遇到影响办理的异常时，及时说明影响范围、临时办理方式和预计恢复时间，并保留复盘记录。'),
    ('运营指标与持续改进', '关注业务成效', '定期查看办理时长、一次办结率和用户反馈，并结合业务量变化分析原因。对共性问题建立改进清单，明确负责人、完成时间和复核方式。'),
    ('培训支持与知识维护', '支持团队使用', '按不同岗位准备简明操作指引，以真实业务案例开展培训。对规则变化和常见问题及时更新说明，保证新成员也能理解办理要求。'),
    ('采购管理与资源保障', '控制投入节奏', '根据试点范围安排采购和服务资源，将交付成果、服务边界和验收方式写入约定。新增需求先评估业务价值和维护成本，再决定实施顺序。'),
    ('安全合规与权限管理', '保护业务信息', '按岗位配置最小必要权限，对涉及敏感信息的办理环节保留操作记录。定期复核离岗和岗位调整人员的权限，及时关闭不再需要的访问。'),
    ('交付验收与运行交接', '明确验收依据', '围绕业务场景验证流程、数据和服务响应，逐项记录测试结果。交接时同步提供运维联系人、问题处理流程和已知限制，避免验收与运营脱节。'),
    ('建设计划与阶段成果', '按阶段形成成果', '将建设任务分为需求确认、试点运行、问题修正和分批推广四个阶段。每个阶段都形成可检查的成果，并由业务负责人确认后进入下一阶段。'),
    ('资金计划与成本控制', '落实资金安排', '按建设阶段列出服务、培训和维护费用，并明确预算责任人。对超出原定范围的支出，先确认业务必要性，再按约定流程调整。'),
    ('服务质量与满意度评价', '开展服务评价', '结合办理记录和用户回访评价服务质量。对集中出现的问题开展专题分析，并把改进结果反馈给相关业务团队。'),
    ('合作单位与沟通机制', '建立沟通节奏', '为各参与单位设置固定沟通窗口，定期同步任务进展和待决事项。涉及业务规则调整的问题，保留确认记录并说明生效时间。'),
    ('成果归档与后续维护', '保证知识延续', '按项目阶段归档业务规则、验收记录和操作指引，明确版本和维护责任。项目结束后继续保留问题跟踪渠道，支持后续运营。'),
]


def make_assets(root):
    from PIL import Image, ImageDraw, ImageFont
    root = Path(root); (root / 'assets').mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (1200, 480), '#f4f7fb'); draw = ImageDraw.Draw(image)
    font_paths = list(Path('/usr/share/fonts').rglob('NotoSansCJK-Regular.ttc'))
    font = ImageFont.truetype(str(font_paths[0]), 30) if font_paths else ImageFont.load_default()
    for i, label in enumerate(['业务申请', '协同办理', '结果反馈']):
        x = 45 + i * 395
        draw.rounded_rectangle((x,150,x+310,330), radius=20, fill='#ffffff', outline='#567dad', width=3)
        draw.text((x+155,240), label, font=font, fill='#26384b', anchor='mm')
        if i < 2:
            draw.line((x+318,240,x+375,240), fill='#567dad', width=5)
            draw.polygon([(x+375,240),(x+357,228),(x+357,252)], fill='#567dad')
    image.save(root / 'assets' / 'workflow.png')


def formatted_operations(sections=12, images=True, tables=True):
    result = [dict(type='append_paragraph', block_id='title', style='Title', text='统一业务服务平台建设方案'),
        dict(type='append_paragraph', block_id='subtitle', text='业务评审稿 · 逐块生成与页面验收测试')]
    for i, (title, subheading, body) in enumerate(SUBJECTS[:sections], 1):
        result.extend([
            dict(type='append_paragraph', block_id=f'h-{i}', style='Heading1', text=f'{i}、{title}'),
            dict(type='append_paragraph', block_id=f'sh-{i}', style='Heading2', text=f'{i}.1 {subheading}'),
            dict(type='append_paragraph', block_id=f'p-{i}', text=body),
            dict(type='append_paragraph', block_id=f'item-{i}-1', text='• 业务负责人：确认办理范围、服务承诺与验收口径。'),
            dict(type='append_paragraph', block_id=f'item-{i}-2', text='• 交付负责人：跟踪阶段成果，记录问题与改进措施。'),
        ])
        if tables:
            result.append(dict(type='append_table', block_id=f'table-{i}', column_widths_twips=[2100,4200,2700], header=True,
                rows=[['工作事项', '交付结果', '验收依据'], [f'{i}-A 规则确认', f'{title}办理清单', '业务部门签认'],
                    [f'{i}-B 试点运行', f'{subheading}试点记录', '问题闭环记录'], [f'{i}-C 推广准备', f'{title}交接说明', '交接核对清单']]))
        if images:
            result.extend([dict(type='insert_image', block_id=f'image-{i}', path='assets/workflow.png', width_cm=12, alt=f'图{i}：{subheading}流程'),
                dict(type='append_paragraph', block_id=f'caption-{i}', text=f'图 {i}　{subheading}的业务协同流程'),
                dict(type='set_paragraph_format', block_id=f'caption-{i}', alignment='center', space_after_twips=120)])
    assert all(len(op.get('text', '')) <= 180 for op in result)
    return result


def make_template(root):
    """Complex source template fixture, not a competing generation path."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches
    path = Path(root) / 'source-template.docx'
    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.2677); section.page_height = Inches(11.6929)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
        section.header.paragraphs[0].text = '业务建设方案 · 内部评审'
        section.footer.paragraphs[0].text = '以正式验收版本为准'
    for style in ('Normal','Title','Heading 1','Heading 2'):
        doc.styles[style].font.name = 'Liberation Sans'
        doc.styles[style]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    doc.add_paragraph('模板保留区：项目管理办公室')
    cell = doc.add_table(rows=1, cols=1).cell(0,0); cell.text='模板审批区'
    cell.add_table(rows=1, cols=2).cell(0,0).text='业务确认'
    doc.save(path)
    return path
