import importlib.util
import io
import json
from pathlib import Path
import shutil
import pytest

SCRIPT=Path(__file__).resolve().parents[1]/'scripts'/'office_pptx.py'
spec=importlib.util.spec_from_file_location('ppt_test_engine',SCRIPT); m=importlib.util.module_from_spec(spec); spec.loader.exec_module(m)


def operation(store,kind,sid='s1',eid='title',**kwargs):
    state=store.status(); req=dict(presentation_id='deck',operation_id='op-'+str(state['revision']+1),base_revision=state['revision'],slide_id=sid,operation=kind,**kwargs)
    return store.apply_one(req)


def initialized(tmp_path):
    store=m.PresentationStore(tmp_path/'deck.pptx');store.init('deck');operation(store,'slide.add',layout='title_body');return store


def accept_fixture(store,sid='s1'):
    # Seed automatic check state to isolate storage/gate tests. No visual pass.
    with store.locked() as(c,state):
        state['reviews'][sid]=dict(automatic_passed=True,status='pending',visual_review='pending',revision=state['revision'],content_notes='fixture')
        store.save_state(c,state);c.commit()


def test_native_elements_editability_and_idempotency(tmp_path):
    from PIL import Image
    store=initialized(tmp_path)
    req=dict(presentation_id='deck',operation_id='heading',base_revision=1,slide_id='s1',operation='element.add_text',element=dict(id='heading',anchor='title',style='slide_title',text='原生对象构建'))
    result=store.apply_one(req);assert store.apply_one(req)==result
    with pytest.raises(ValueError):store.apply_one(dict(req,operation_id='stale'))
    with pytest.raises(ValueError):store.apply_one(dict(req,element=dict(req['element'],text='different')))
    operation(store,'element.add_shape',element=dict(id='card',geometry=dict(x=.06,y=.3,w=.3,h=.4),shape='rectangle',layer='background'))
    operation(store,'element.add_table',element=dict(id='table',geometry=dict(x=.06,y=.3,w=.4,h=.4),rows=[['事项','说明'],['A','原生表格']]))
    operation(store,'element.add_chart',element=dict(id='chart',geometry=dict(x=.53,y=.3,w=.4,h=.4),chart=dict(type='bar',categories=['A','B'],series=[dict(name='Score',values=[12,28])])))
    Image.new('RGB',(200,100),'blue').save(tmp_path/'photo.png')
    operation(store,'element.add_image',element=dict(id='picture',geometry=dict(x=.7,y=.05,w=.2,h=.1777778),path='photo.png'))
    prs=m.Presentation(io.BytesIO(store.snapshot(store.status()['revision'])))
    assert len(prs.slides[0].shapes)==5
    assert any(shape.has_table for shape in prs.slides[0].shapes)
    assert any(shape.has_chart for shape in prs.slides[0].shapes)
    chart=next(shape.chart for shape in prs.slides[0].shapes if shape.has_chart)
    assert tuple(chart.series[0].values)==(12.0,28.0)
    assert len(prs.slides[0].shapes[0].text)>0
    m.validator.validate_package(store.path.read_bytes())


def test_failed_op_restore_and_slide_gate(tmp_path):
    store=initialized(tmp_path); before=store.path.read_bytes()
    with pytest.raises(ValueError):operation(store,'element.add_text',element=dict(id='bad',text='x',geometry=dict(x=.9,y=.1,w=.5,h=.1)))
    assert store.path.read_bytes()==before and store.status()['revision']==1
    with pytest.raises(ValueError):operation(store,'slide.add',sid='s2',layout='data')
    operation(store,'element.add_text',element=dict(id='title',anchor='title',text='可回滚的内容'))
    assert store.restore(1,2,'restore')['revision']==3
    assert store.path.read_bytes()==store.snapshot(1)
    with pytest.raises(ValueError):store.finalize(3)


def test_twenty_slides_and_twenty_elements(tmp_path):
    store=initialized(tmp_path)
    for i in range(20):
        operation(store,'element.add_shape',element=dict(id=f'shape-{i}',geometry=dict(x=.05+(i%5)*.18,y=.25+(i//5)*.15,w=.12,h=.1),shape='rectangle',layer='decoration'))
    assert len(store.status()['manifest']['slides'][0]['elements'])==20
    for i in range(2,21):
        accept_fixture(store,f's{i-1}')
        operation(store,'slide.add',sid=f's{i}',layout='title_body')
        operation(store,'element.add_text',sid=f's{i}',element=dict(id='title',anchor='title',text=f'第 {i} 页'))
    slides=store.status()['manifest']['slides'];assert [s['slide_id'] for s in slides]==[f's{i}' for i in range(1,21)]


def test_theme_and_historical_edits_invalidate_only_affected_slides(tmp_path):
    store=initialized(tmp_path);operation(store,'element.add_text',element=dict(id='title',anchor='title',text='第一页'))
    accept_fixture(store);operation(store,'slide.add',sid='s2',layout='title_body');operation(store,'element.add_text',sid='s2',element=dict(id='title',anchor='title',text='第二页'));accept_fixture(store,'s2')
    operation(store,'element.update',sid='s1',element=dict(id='title',text='第一页面向业务的描述'))
    slides=store.status()['manifest']['slides'];assert slides[0]['review']['status']=='writing' and slides[1]['review']['automatic_passed']
    operation(store,'presentation.set_theme',theme='slate_dark')
    assert all(s['review']['status']=='writing' for s in store.status()['manifest']['slides'])
    assert all(s['background']=='14263D' for s in store.status()['manifest']['slides'])


def test_template_masters_theme_and_resources_preserved(tmp_path):
    from pptx import Presentation
    import zipfile
    template=tmp_path/'template.pptx';prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_textbox(100000,100000,2000000,400000).text='企业模板保留对象';prs.save(template)
    store=m.PresentationStore(tmp_path/'deck.pptx');store.init('deck',template)
    sid=store.status()['manifest']['slides'][0]['slide_id']
    operation(store,'element.add_text',sid=sid,element=dict(id='new',geometry=dict(x=.2,y=.6,w=.5,h=.2),text='新增原生文字'))
    with zipfile.ZipFile(template) as a,zipfile.ZipFile(store.path) as b:
        for name in a.namelist():
            if name.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/','ppt/media/')):assert a.read(name)==b.read(name),name
    assert store.status()['manifest']['slides'][0]['complex']


def test_geometry_overlap_and_text_overflow_checks(tmp_path):
    store=initialized(tmp_path)
    operation(store,'element.add_text',element=dict(id='a',geometry=dict(x=.1,y=.3,w=.3,h=.02),text='需要明确控制的文字内容'*8))
    operation(store,'element.add_text',element=dict(id='b',geometry=dict(x=.1,y=.3,w=.3,h=.1),text='第二个文本框'))
    issues=m.validator.layout_issues(store.status()['manifest']['slides'][0])
    assert any(i['type']=='text_overflow' for i in issues)
    assert any(i['type']=='unexpected_overlap' for i in issues)


@pytest.mark.skipif(not shutil.which('libreoffice'),reason='LibreOffice Impress required')
def test_actual_impress_render_and_pending_visual(tmp_path):
    store=initialized(tmp_path);operation(store,'element.add_text',element=dict(id='title',anchor='title',style='slide_title',text='AI Agent 平台建设方案'))
    result=store.check_slide('s1','逐字检查标题为 AI Agent 平台建设方案。')
    assert result['review']['automatic_passed'],result
    assert result['review']['visual_review']=='pending'
    # Visual review is advisory: delivery only needs the automatic and content checks.
    event=store.finalize(store.status()['revision'])
    assert event['finalized'] and event['aesthetics']['s1']['status']=='pending' and event['outstanding_critical']=={}


def test_targeted_table_chart_data_and_line_updates_remain_native(tmp_path):
    store=initialized(tmp_path)
    operation(store,'element.add_table',element=dict(id='table',geometry=dict(x=.05,y=.25,w=.4,h=.3),rows=[['A','B'],['1','2']]))
    operation(store,'element.update',element=dict(id='table',rows=[['业务','说明'],['试点','已确认']]))
    operation(store,'element.add_chart',element=dict(id='chart',geometry=dict(x=.55,y=.25,w=.4,h=.4),chart=dict(type='bar',categories=['A','B'],series=[dict(name='Score',values=[1,2])])))
    operation(store,'element.update',element=dict(id='chart',chart=dict(type='bar',categories=['A','B'],series=[dict(name='Score',values=[3,4])])))
    operation(store,'element.add_shape',element=dict(id='line',shape='line',layer='decoration',geometry=dict(x=.1,y=.8,w=.8,h=.001)))
    prs=m.Presentation(io.BytesIO(store.path.read_bytes()));slide=prs.slides[0]
    assert m.find_element(slide,'table').table.cell(1,0).text=='试点'
    assert tuple(m.find_element(slide,'chart').chart.series[0].values)==(3.0,4.0)
    assert m.renderer.metadata(m.find_element(slide,'line'))['shape']=='line'
