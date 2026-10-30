"""Reliability of template fill, palette, logo protection, links and delivery slides."""
import importlib.util
import io
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def load(name):
    spec = importlib.util.spec_from_file_location('pptx_rel_' + name, SCRIPTS / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


office = load('office_pptx')
sample = load('sample_template')
fill = load('template_fill')


def accept_all(store):
    with store.locked() as (connection, state):
        for slide in store.view(connection, state)['manifest']['slides']:
            state['reviews'][slide['slide_id']] = dict(
                automatic_passed=True, status='pending', visual_review='pending',
                revision=state['revision'], content_notes='fixture',
                fingerprint=f'fp-{slide["slide_id"]}-{state["revision"]}')
        store.save_state(connection, state)
        connection.commit()


def apply(store, kind, **kwargs):
    state = store.status()
    return store.apply_one(dict(
        presentation_id='deck', operation_id=f'op-{state["revision"] + 1}',
        base_revision=state['revision'], operation=kind, **kwargs))


def solid_values(shape):
    return [node.get('val').upper() for node in shape._element.iter(A + 'srgbClr') if node.get('val')]


def test_shape_fill_keeps_schema_order_geometry_and_line():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    rounded = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    rounded.fill.gradient()
    rounded.fill.gradient_angle = 45
    rounded.line.color.rgb = RGBColor(0x11, 0x22, 0x33)
    rounded.line.width = Pt(2)
    box = (rounded.left, rounded.top, rounded.width, rounded.height)
    fill.set_shape_fill(rounded, '707780')
    assert [etree.QName(child).localname for child in rounded._element.spPr] == ['xfrm', 'prstGeom', 'solidFill', 'ln']
    assert (rounded.left, rounded.top, rounded.width, rounded.height) == box
    assert str(rounded.line.color.rgb) == '112233'
    rectangle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(4), Inches(1), Inches(1.2), Inches(0.8))
    rectangle.fill.solid()
    rectangle.fill.fore_color.rgb = RGBColor(0, 0, 255)
    rectangle.line.width = Pt(1.5)
    fill.set_shape_fill(rectangle, 'AABBCC')
    assert [etree.QName(child).localname for child in rectangle._element.spPr][:4] == ['xfrm', 'prstGeom', 'solidFill', 'ln']
    text_box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(0.6))
    text_box.text_frame.paragraphs[0].text = '边框'
    text_box.line.color.rgb = RGBColor(0x22, 0x22, 0x22)
    text_box.line.width = Pt(1)
    fill.set_shape_fill(text_box, 'F4F4F4')
    assert str(text_box.line.color.rgb) == '222222'
    picture = slide.shapes.add_picture(io.BytesIO(_png()), Inches(6), Inches(1), Inches(1), Inches(1))
    with pytest.raises(ValueError, match='blipFill'):
        fill.set_shape_fill(picture, '000000')


def test_imported_fill_failure_does_not_publish(tmp_path):
    deck = _bordered_shape(tmp_path / 'shapes.pptx')
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', deck)
    before = store.path.read_bytes()
    revision = store.status()['revision']
    with pytest.raises(ValueError, match='not restyled|blipFill|picture'):
        apply(store, 'element.set_style', slide_id='tpl-004', element={'id': 'photo', 'style': {'fill': '112233'}})
    assert store.path.read_bytes() == before
    assert store.status()['revision'] == revision
    apply(store, 'element.set_style', slide_id='tpl-004', element={'id': 'band', 'style': {'fill': '445566'}})
    opened = Presentation(store.path)
    band = next(shape for slide in opened.slides for shape in slide.shapes if shape.name == 'da:band')
    assert '445566' in solid_values(band)
    assert band.left > 0 and band.width > 0
    office.validator.validate_package(store.path.read_bytes())


def test_recolor_palette_applies_to_new_objects_and_restores(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', tmp_path / 'template.pptx')
    before = store.status()['revision']
    apply(store, 'presentation.recolor', palette={'accent': '707780', 'background': 'FFFFFF', 'text': '20242A', 'muted': '79818A'})
    assert store.status()['theme']['accent'] == '707780'
    assert store.status()['template_index']['slides']
    tokens = store.status()  # theme is synced; original tokens stay on the design record
    accept_all(store)
    apply(store, 'slide.add', slide_id='fresh', layout='title_body')
    apply(store, 'element.add_shape', slide_id='fresh', element={
        'id': 'swatch', 'geometry': {'x': 0.1, 'y': 0.3, 'w': 0.2, 'h': 0.2}, 'shape': 'rectangle', 'fill_token': 'accent'})
    apply(store, 'element.add_chart', slide_id='fresh', element={
        'id': 'bars', 'geometry': {'x': 0.4, 'y': 0.3, 'w': 0.5, 'h': 0.4},
        'chart': {'type': 'bar', 'categories': ['Q1', 'Q2'], 'series': [{'name': '收入', 'values': [3, 5]}]}})
    apply(store, 'element.add_table', slide_id='fresh', element={
        'id': 'grid', 'geometry': {'x': 0.1, 'y': 0.6, 'w': 0.25, 'h': 0.2}, 'rows': [['项', '值'], ['甲', '1']]})
    opened = Presentation(store.path)
    fresh = opened.slides[-1]
    swatch = next(shape for shape in fresh.shapes if shape.name == 'da:swatch')
    assert '707780' in solid_values(swatch)
    assert '246BCE' not in solid_values(swatch)
    chart = next(shape for shape in fresh.shapes if shape.has_chart)
    assert '707780' in etree.tostring(chart.chart._element).decode().upper()
    table = next(shape for shape in fresh.shapes if shape.has_table)
    header = ''.join(solid_values(table)[:1]) or etree.tostring(table._element).decode().upper()
    assert 'EDF4FC' not in header
    restored = store.restore(before, store.status()['revision'], 'restore-palette')
    assert restored['revision'] == store.status()['revision']
    assert store.status()['theme']['accent'] != '707780'
    rolled = Presentation(store.path)
    band = next(shape for slide in rolled.slides for shape in slide.shapes if shape.name == 'da:band')
    assert '0B2545' in solid_values(band)
    del tokens


def test_company_logo_survives_rename_clone_and_recolor(tmp_path):
    path = tmp_path / 'brand.pptx'
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    logo = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(11), Inches(0.4), Inches(0.8), Inches(0.8))
    logo.name = 'Company Logo'
    logo.fill.solid()
    logo.fill.fore_color.rgb = RGBColor.from_string('0B2545')
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.5), Inches(0.4), Inches(0.3), Inches(2))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string('0B2545')
    _paint_theme(presentation, '0B2545')
    presentation.save(path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', path)
    accept_all(store)
    source = store.status()['manifest']['slides'][0]['slide_id']
    apply(store, 'slide.clone', slide_id='copy', source_slide_id=source)
    apply(store, 'presentation.recolor', palette={'accent': '707780', 'background': 'FFFFFF', 'text': '20242A', 'muted': '79818A'})
    opened = Presentation(store.path)
    logos = [shape for slide in opened.slides for shape in slide.shapes if shape.name.startswith('da:shape-') and '0B2545' in solid_values(shape)]
    assert len(logos) == 2
    bands = [shape for slide in opened.slides for shape in slide.shapes if '707780' in solid_values(shape)]
    assert bands


def test_internal_slide_link_is_not_copied_as_a_new_slide(tmp_path):
    path = tmp_path / 'linked.pptx'
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    first._element.cSld.set('name', 'da-slide:origin')
    second._element.cSld.set('name', 'da-slide:target')
    button = first.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    rel_id = first.part.relate_to(second.part, RT.SLIDE)
    props = button._element.find('.//' + qn('p:cNvPr'))
    link = etree.SubElement(props, qn('a:hlinkClick'))
    link.set(qn('r:id'), rel_id)
    link.set('action', 'ppaction://hlinksldjump')
    presentation.save(path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', path)
    accept_all(store)
    apply(store, 'slide.clone', slide_id='origin-2', source_slide_id='origin')
    opened = Presentation(store.path)
    assert len(opened.slides) == 3
    clone = next(slide for slide in opened.slides if slide._element.cSld.get('name') == 'da-slide:origin-2')
    target_names = []
    for rel in clone.part.rels.values():
        if rel.reltype == RT.SLIDE and not rel.is_external:
            target_names.append(str(rel.target_part.partname))
    original_target = str(second.part.partname)
    assert target_names
    assert all(name == original_target or name.endswith(original_target.split('/')[-1]) for name in target_names)
    office.validator.validate_package(store.path.read_bytes())


def _import(store, slide_id, source_slide_id, template_id='consulting-brief', **kwargs):
    return apply(store, 'slide.import', slide_id=slide_id, source_slide_id=source_slide_id,
        template_id=template_id, **kwargs)


def test_blank_init_has_no_template_pages(tmp_path):
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    status = store.status()
    assert status['manifest']['slides'] == []
    assert status.get('imported_slides') == {}
    assert 'prototype_slides' not in status
    assert 'template_use' not in status


def test_import_from_public_library_does_not_touch_source(tmp_path):
    library = Path(__file__).resolve().parents[1] / 'templates' / 'consulting-brief' / 'template.pptx'
    before = library.read_bytes()
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    apply(store, 'slide.import', slide_id='page-1', source_slide_id='tpl-001', template_id='consulting-brief')
    assert library.read_bytes() == before
    source = store.status()['imported_slides']['page-1']
    assert source['template_id'] == 'consulting-brief'
    assert source['source_slide_id'] == 'tpl-001'
    opened = Presentation(store.path)
    assert any(shape.name == 'da:title' for shape in opened.slides[0].shapes)
    office.validator.validate_package(store.path.read_bytes())


def test_library_match_does_not_need_a_working_pptx():
    importer = load('template_importer')
    listed = importer.list_templates()
    brief = next(item for item in listed if item['template_id'] == 'consulting-brief')
    assert brief['source_sha256']
    assert brief['slides'][0]['preview']
    matched = importer.match_library('consulting-brief', 'data')
    assert matched['template_id'] == 'consulting-brief'
    assert matched['slides'][0]['slide_id'] == 'tpl-004'
    assert matched['slides'][0]['slots']['chart'] == 'chart'


def test_slide_import_delivers_only_generated_slides(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    sources = ['tpl-001', 'tpl-002', 'tpl-003', 'tpl-004']
    for index in range(10):
        accept_all(store)
        _import(store, f'page-{index + 1}', sources[index % 4], template_path='template.pptx')
    final_ids = [slide['slide_id'] for slide in store.status()['manifest']['slides']]
    assert final_ids == [f'page-{index + 1}' for index in range(10)]
    office.validator.validate_package(store.path.read_bytes())
    accept_all(store)
    store.finalize(store.status()['revision'])
    assert store.status()['finalized'] is True


def test_second_import_requires_check_slide_of_generated_pages(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    _import(store, 'page-1', 'tpl-001', template_path='template.pptx')
    revision = store.status()['revision']
    before = store.path.read_bytes()
    with pytest.raises(ValueError, match='Check all affected slides before creating the next slide'):
        _import(store, 'page-2', 'tpl-002', template_path='template.pptx')
    assert store.path.read_bytes() == before
    assert store.status()['revision'] == revision
    accept_all(store)
    _import(store, 'page-2', 'tpl-002', template_path='template.pptx')
    assert [slide['slide_id'] for slide in store.status()['manifest']['slides']] == ['page-1', 'page-2']


def test_import_failure_does_not_publish(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    before = store.path.read_bytes()
    revision = store.status()['revision']
    with pytest.raises(ValueError, match='Unknown or ambiguous stable slide ID'):
        _import(store, 'page-1', 'missing', template_path='template.pptx')
    assert store.path.read_bytes() == before
    assert store.status()['revision'] == revision
    _import(store, 'page-1', 'tpl-001', template_path='template.pptx')
    with pytest.raises(ValueError, match='not supported'):
        apply(store, 'presentation.discard_prototypes')
    assert store.status()['revision'] == revision + 1


def test_repeated_import_keeps_independent_charts(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    source_hash = (tmp_path / 'template.pptx').read_bytes()
    accept_all(store)
    _import(store, 'page-a', 'tpl-004', template_path='template.pptx')
    accept_all(store)
    _import(store, 'page-b', 'tpl-004', template_path='template.pptx')
    apply(store, 'chart.replace_data', slide_id='page-b', element_id='chart',
        chart={'categories': ['Q1', 'Q2', 'Q3'], 'series': [{'name': 'Revenue', 'values': [4, 5, 6]}]})
    opened = Presentation(store.path)
    values = {}
    parts = set()
    for slide in opened.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                values[slide._element.cSld.get('name')] = tuple(shape.chart.series[0].values)
                parts.add(str(shape.chart.part.partname))
    assert values['da-slide:page-a'] == (10, 20, 30)
    assert values['da-slide:page-b'] == (4, 5, 6)
    assert len(parts) == 2
    assert (tmp_path / 'template.pptx').read_bytes() == source_hash
    office.validator.validate_package(store.path.read_bytes())


def test_internal_jump_import_rebinds_or_refuses(tmp_path):
    path = _linked_template(tmp_path / 'linked.pptx')
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    before = store.path.read_bytes()
    revision = store.status()['revision']
    with pytest.raises(ValueError, match='unresolved_internal_link'):
        apply(store, 'slide.import', slide_id='page-001', source_slide_id='tpl-001',
            template_id='linked', template_path='linked.pptx')
    assert store.path.read_bytes() == before
    assert store.status()['revision'] == revision
    _import(store, 'page-002', 'tpl-002', template_id='linked', template_path='linked.pptx')
    accept_all(store)
    _import(store, 'page-001', 'tpl-001', template_id='linked', template_path='linked.pptx')
    opened = Presentation(store.path)
    nav = next(slide for slide in opened.slides if slide._element.cSld.get('name') == 'da-slide:page-001')
    assert _jump_ids(opened, nav) == ['page-002']
    office.validator.validate_package(store.path.read_bytes())


def test_restore_keeps_import_source_metadata(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    _import(store, 'page-1', 'tpl-001', template_path='template.pptx')
    imported_revision = store.status()['revision']
    apply(store, 'slot.fill', slide_id='page-1', slots={'title': '稍后还原', 'subtitle': '来源仍在'})
    store.restore(imported_revision, store.status()['revision'], 'restore-import')
    source = store.status()['imported_slides']['page-1']
    assert source['template_id'] == 'consulting-brief'
    assert source['source_slide_id'] == 'tpl-001'
    assert source['source_sha256']
    opened = Presentation(store.path)
    title = next(shape for shape in opened.slides[0].shapes if shape.name == 'da:title')
    assert '稍后还原' not in title.text_frame.text


def test_edit_existing_keeps_internal_jumps(tmp_path):
    path = _linked_template(tmp_path / 'linked.pptx')
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', path)
    opened = Presentation(store.path)
    origin = next(slide for slide in opened.slides if slide._element.cSld.get('name') == 'da-slide:tpl-001')
    assert _jump_ids(opened, origin) == ['tpl-002']
    with pytest.raises(ValueError, match='not supported'):
        apply(store, 'presentation.discard_prototypes', redirects={'tpl-002': 'tpl-002'})
    assert [slide['slide_id'] for slide in store.status()['manifest']['slides']] == ['tpl-001', 'tpl-002']


def test_deck_review_is_rejected_after_the_document_changes(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', tmp_path / 'template.pptx')
    accept_all(store)
    first = store.request_deck_review(store.status()['revision'])
    payload = dict(status='passed', summary='old', review_id='rev-old')
    assert store.record_deck_review(first['deck_review']['request_id'], payload) == dict(applied=True)
    assert store.review_status()['deck_review']['status'] == 'passed'
    store.record_human_decision(store.status()['revision'], 'accept', '看起来可以', 'human-old')
    assert store.status()['deck_review']['human']['decision'] == 'accept'
    accept_all(store)
    apply(store, 'slot.fill', slide_id='tpl-001', slots={'title': '新标题', 'subtitle': '修订'})
    deck = store.status()['deck_review']
    assert deck['status'] == 'stale'
    assert deck['request_id'] == first['deck_review']['request_id']
    assert store.record_deck_review(first['deck_review']['request_id'], dict(status='passed', summary='late', review_id='rev-late')) == dict(applied=False, reason='stale')
    assert store.status()['deck_review']['status'] == 'stale'
    assert store.status()['deck_review'].get('review_id') != 'rev-late'
    with pytest.raises(ValueError, match='No current deck review'):
        store.record_human_decision(store.status()['revision'], 'accept', '不能沿用旧版', 'human-new')
    accept_all(store)
    second = store.request_deck_review(store.status()['revision'])
    assert second['deck_review']['request_id'] != first['deck_review']['request_id']
    assert store.record_deck_review(first['deck_review']['request_id'], dict(status='passed', summary='first-again')) == dict(applied=False, reason='stale')
    assert store.record_deck_review(second['deck_review']['request_id'], dict(status='passed', summary='current', review_id='rev-new')) == dict(applied=True)
    assert store.review_status()['deck_review']['status'] == 'passed'
    assert store.review_status()['deck_review']['summary'] == 'current'


def test_deck_review_fingerprint_and_hash_mismatch_are_stale(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', tmp_path / 'template.pptx')
    accept_all(store)
    requested = store.request_deck_review(store.status()['revision'])
    request_id = requested['deck_review']['request_id']
    revision = store.status()['revision']
    with store.locked() as (connection, state):
        state['deck_review']['slide_fingerprints']['tpl-001'] = 'tampered'
        store.save_state(connection, state)
        connection.commit()
    assert store.record_deck_review(request_id, dict(status='passed', summary='bad-fp')) == dict(applied=False, reason='stale')
    assert store.status()['revision'] == revision
    with store.locked() as (connection, state):
        state['deck_review']['slide_fingerprints']['tpl-001'] = f'fp-tpl-001-{revision}'
        state['deck_review']['sha256'] = '0' * 64
        store.save_state(connection, state)
        connection.commit()
    assert store.record_deck_review(request_id, dict(status='passed', summary='bad-hash')) == dict(applied=False, reason='stale')
    with store.locked() as (connection, state):
        state['deck_review']['sha256'] = state['sha256']
        store.save_state(connection, state)
        connection.commit()
    assert store.record_deck_review(request_id, dict(status='passed', summary='ok', review_id='fresh')) == dict(applied=True)
    assert store.status()['revision'] == revision
    assert store.review_status()['deck_review']['summary'] == 'ok'


def _linked_template(path):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    first._element.cSld.set('name', 'da-slide:tpl-001')
    second._element.cSld.set('name', 'da-slide:tpl-002')
    _add_jump(first, second, 'nav')
    presentation.save(path)
    return path


def _add_jump(source, target, name='jump'):
    button = source.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    button.name = 'da:' + name
    rel_id = source.part.relate_to(target.part, RT.SLIDE)
    props = button._element.find('.//' + qn('p:cNvPr'))
    link = etree.SubElement(props, qn('a:hlinkClick'))
    link.set(qn('r:id'), rel_id)
    link.set('action', 'ppaction://hlinksldjump')
    return button


def _jump_ids(presentation, slide):
    by_part = {item.part: office.renderer.slide_identity(item) for item in presentation.slides}
    return [by_part[rel.target_part] for rel in slide.part.rels.values()
            if rel.reltype == RT.SLIDE and not rel.is_external]


def _png():
    from PIL import Image
    stream = io.BytesIO()
    Image.new('RGB', (32, 32), '#123456').save(stream, format='PNG')
    stream.seek(0)
    return stream.getvalue()


def _bordered_shape(path):
    # Reuse the sample data slide, which already has a picture named da:photo and a band.
    sample = load('sample_template')
    directory = path.parent / 'sample'
    sample.build(directory)
    return directory / 'template.pptx'


def _paint_theme(presentation, accent):
    theme = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    namespace = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    scheme = root.find('.//' + namespace + 'clrScheme')
    for child in scheme:
        if etree.QName(child).localname == 'accent1':
            for node in list(child):
                child.remove(node)
            color = etree.SubElement(child, namespace + 'srgbClr')
            color.set('val', accent)
    theme.blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
