"""Licensed project business template: import, fill, recolor, real Impress renders."""
import importlib.util
import json
import shutil
import time
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def load(name):
    spec = importlib.util.spec_from_file_location('pptx_biz_' + name, SCRIPTS / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


office = load('office_pptx')
sample = load('sample_template')
importer = load('template_importer')
baseline = load('render_baseline')
GRAY = {'accent': '707780', 'background': 'FFFFFF', 'text': '20242A', 'muted': '79818A'}
PLAN = (
    ('page-1', 'tpl-001', {'title': '企业项目汇报', 'subtitle': '沿用商务模板'}),
    ('page-2', 'tpl-003', {'title': '方案对比', 'left_title': '现状', 'left_body': '分散执行', 'right_title': '目标', 'right_body': '统一版式'}),
    ('page-3', 'tpl-004', {'title': '关键数据'}),
    ('page-4', 'tpl-003', {'title': '执行架构', 'left_title': '编排', 'left_body': '原子操作', 'right_title': '渲染', 'right_body': 'Impress'}),
    ('page-5', 'tpl-002', {'title': '范围'}),
    ('page-6', 'tpl-004', {'title': '进度'}),
    ('page-7', 'tpl-003', {'title': '风险对比', 'left_title': '风险', 'left_body': '格式漂移', 'right_title': '控制', 'right_body': '克隆填槽'}),
    ('page-8', 'tpl-002', {'title': '下一步'}),
    ('page-9', 'tpl-001', {'title': '总结', 'subtitle': '只交付生成页'}),
    ('page-10', 'tpl-002', {'title': '附录'}),
)


def accept_all(store):
    with store.locked() as (connection, state):
        for slide in store.view(connection, state)['manifest']['slides']:
            state['reviews'][slide['slide_id']] = dict(
                automatic_passed=True, status='pending', visual_review='pending',
                revision=state['revision'], content_notes='fixture',
                fingerprint=f'fp-{slide["slide_id"]}-{state["revision"]}')
        store.save_state(connection, state)
        connection.commit()


def apply(store, kind, **kwargs):
    state = store.status()
    return store.apply_one(dict(
        presentation_id='deck', operation_id=f'op-{state["revision"] + 1}',
        base_revision=state['revision'], operation=kind, **kwargs))


def native_inventory(path):
    presentation = Presentation(path)
    counts = dict(charts=0, tables=0, pictures=0, shapes=0)
    fonts, boxes = set(), []
    for slide in presentation.slides:
        for shape in slide.shapes:
            boxes.append((shape.name, int(shape.left), int(shape.top), int(shape.width), int(shape.height)))
            if shape.has_chart:
                counts['charts'] += 1
            elif shape.has_table:
                counts['tables'] += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts['pictures'] += 1
            else:
                counts['shapes'] += 1
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts.add(run.font.name)
    return counts, fonts, boxes


def contact_sheet(pages, dest, columns=5):
    from PIL import Image, ImageDraw
    images = [Image.open(path).convert('RGB') for path in pages]
    width = 240
    thumbs = [image.resize((width, round(image.height * width / image.width))) for image in images]
    height = max(thumb.height for thumb in thumbs)
    rows = (len(thumbs) + columns - 1) // columns
    sheet = Image.new('RGB', (columns * width + 24, rows * (height + 28) + 24), 'white')
    draw = ImageDraw.Draw(sheet)
    for index, thumb in enumerate(thumbs):
        row, column = divmod(index, columns)
        x, y = 12 + column * width, 12 + row * (height + 28)
        draw.text((x, y), f'{index + 1:02d}', fill='black')
        sheet.paste(thumb, (x, y + 16))
    sheet.save(dest)


def build_report(tmp_path, timings=None):
    sample.build(tmp_path)
    source_hash = (tmp_path / 'template.pptx').read_bytes()
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    listed = importer.list_templates()
    assert any(item['template_id'] == 'consulting-brief' for item in listed)
    store.init('deck')
    started = time.perf_counter()
    for slide_id, source, slots in PLAN:
        accept_all(store)
        mark = time.perf_counter()
        apply(store, 'slide.import', slide_id=slide_id, source_slide_id=source,
            template_id='consulting-brief', template_path='template.pptx')
        if timings is not None:
            timings.setdefault('import_s', []).append(time.perf_counter() - mark)
        apply(store, 'slot.fill', slide_id=slide_id, slots=slots)
    apply(store, 'presentation.recolor', palette=GRAY)
    apply(store, 'element.adjust', slide_id='page-4', element={'id': 'title', 'dx_in': 0.25})
    apply(store, 'chart.replace_data', slide_id='page-3', element_id='chart',
        chart={'categories': ['Q1', 'Q2', 'Q3'], 'series': [{'name': '进度', 'values': [4, 6, 9]}]})
    if timings is not None:
        timings['total_s'] = time.perf_counter() - started
        timings['pptx_bytes'] = store.path.stat().st_size
    assert (tmp_path / 'template.pptx').read_bytes() == source_hash
    return store


def test_business_template_builds_ten_generated_pages(tmp_path):
    timings = {}
    store = build_report(tmp_path, timings)
    ids = [slide['slide_id'] for slide in store.status()['manifest']['slides']]
    assert ids == [item[0] for item in PLAN]
    assert store.status()['theme']['accent'] == '707780'
    opened = Presentation(store.path)
    cover = opened.slides[0]
    assert _text(cover, 'da:title') == '企业项目汇报'
    logo = next(shape for shape in cover.shapes if shape.name == 'da:logo')
    assert '0B2545' in [node.get('val').upper() for node in logo._element.iter(A + 'srgbClr') if node.get('val')]
    fourth = opened.slides[3]
    title = next(shape for shape in fourth.shapes if shape.name == 'da:title')
    source = Presentation(tmp_path / 'template.pptx').slides[2]
    original = next(shape for shape in source.shapes if shape.name == 'da:title')
    assert title.left - original.left == 228600
    data = opened.slides[2]
    chart = next(shape for shape in data.shapes if shape.has_chart)
    assert tuple(chart.chart.series[0].values) == (4, 6, 9)
    photo = next(shape for shape in data.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert photo.width > 0
    counts, fonts, _ = native_inventory(store.path)
    assert counts['charts'] == 2 and counts['pictures'] == 2
    assert 'Liberation Sans' in fonts
    unmoved = next(shape for shape in opened.slides[1].shapes if shape.name == 'da:title')
    source_cmp = Presentation(tmp_path / 'template.pptx').slides[2]
    original_cmp = next(shape for shape in source_cmp.shapes if shape.name == 'da:title')
    assert unmoved.left == original_cmp.left
    sources = store.status()['imported_slides']
    assert sources['page-1']['source_slide_id'] == 'tpl-001'
    assert sources['page-3']['source_slide_id'] == 'tpl-004'
    accept_all(store)
    store.finalize(store.status()['revision'])
    office.validator.validate_package(store.path.read_bytes())
    assert timings['import_s']
    assert timings['total_s'] > 0


@pytest.mark.skipif(not shutil.which('libreoffice'), reason='LibreOffice Impress required')
def test_business_template_libreoffice_evidence(tmp_path):
    evidence = tmp_path / 'evidence'
    evidence.mkdir()
    sample.build(tmp_path / 'source')
    template = tmp_path / 'source' / 'template.pptx'
    before_pdf = baseline.render_pdf(template, tmp_path / 'before')
    baseline.save_page(before_pdf, evidence / 'before.png', index=0)
    store = build_report(tmp_path / 'work')
    filled_pdf = baseline.render_pdf(store.path, tmp_path / 'filled')
    baseline.save_page(filled_pdf, evidence / 'filled.png', index=0)
    baseline.save_page(filled_pdf, evidence / 'recolored.png', index=3)
    pages = []
    with __import__('pymupdf').open(filled_pdf) as document:
        assert len(document) == 10
        for index in range(len(document)):
            dest = evidence / f'delivery-{index + 1:02d}.png'
            baseline.save_page(filled_pdf, dest, index=index)
            pages.append(dest)
    contact_sheet(pages, evidence / 'final-sheet.png')
    baseline.comparison(evidence / 'before.png', evidence / 'recolored.png', evidence / 'comparison.png')
    counts, fonts, boxes = native_inventory(store.path)
    report = dict(
        renderer='LibreOffice Impress', source='templates/consulting-brief',
        slides=[slide['slide_id'] for slide in store.status()['manifest']['slides']],
        imported_slides=store.status()['imported_slides'],
        theme=store.status()['theme'], native=counts, fonts=sorted(fonts),
        title_moved=any(name == 'da:title' and box[0] > 0 for name, *box in boxes),
        images={name: dict(bytes=(evidence / name).stat().st_size)
            for name in ('before.png', 'filled.png', 'recolored.png', 'comparison.png', 'final-sheet.png')},
        note='Screenshots are LibreOffice PDF exports. Geometry and native objects are checked in XML, not by pixel match.')
    (evidence / 'validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    assert report['native']['charts'] == 2
    assert (evidence / 'final-sheet.png').stat().st_size > 1000


@pytest.mark.skipif(not shutil.which('libreoffice'), reason='LibreOffice Impress required')
def test_imported_cover_runs_real_check_slide(tmp_path):
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck')
    apply(store, 'slide.import', slide_id='page-1', source_slide_id='tpl-001',
        template_id='consulting-brief', template_path='template.pptx')
    apply(store, 'slot.fill', slide_id='page-1', slots={'title': '检查封面', 'subtitle': '真实渲染'})
    result = store.check_slide('page-1', 'Checked cover title and subtitle', intent='封面')
    assert result['review']['automatic_passed'] is True
    assert result['review']['rendered'] is True
    assert result['review']['content_notes']


def _text(slide, name):
    return next(shape.text_frame.text for shape in slide.shapes if shape.name == name)
