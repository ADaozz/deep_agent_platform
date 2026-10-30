"""Template import, clone, fill and recolor stay inside the atomic PPTX protocol."""
import importlib.util
import io
from pathlib import Path

import pytest
from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def load(name):
    spec = importlib.util.spec_from_file_location('pptx_' + name, SCRIPTS / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


office = load('office_pptx')
sample = load('sample_template')
importer = load('template_importer')


def accept_all(store):
    with store.locked() as (connection, state):
        for slide in store.view(connection, state)['manifest']['slides']:
            state['reviews'][slide['slide_id']] = dict(
                automatic_passed=True, status='pending', visual_review='pending',
                revision=state['revision'], content_notes='fixture')
        store.save_state(connection, state)
        connection.commit()


def apply(store, kind, **kwargs):
    state = store.status()
    request = dict(presentation_id='deck', operation_id=f'op-{state["revision"] + 1}',
        base_revision=state['revision'], operation=kind, **kwargs)
    return store.apply_one(request)


def colors(shape):
    return [node.get('val').upper() for node in shape._element.iter(A + 'srgbClr') if node.get('val')]


def open_deck(data):
    return Presentation(io.BytesIO(data))


def test_import_records_roles_and_refuses_unlicensed_vendor(tmp_path):
    sample.build(tmp_path / 'licensed')
    external = tmp_path / 'external'
    external.mkdir()
    (external / 'template.pptx').write_bytes((tmp_path / 'licensed' / 'template.pptx').read_bytes())
    manifest = importer.import_template(external / 'template.pptx', 'consulting-brief', license_note='external check')
    roles = [slide['role'] for slide in manifest['slides']]
    assert roles == ['cover', 'section', 'comparison', 'data']
    assert manifest['license']['redistributable'] is False
    assert manifest['license']['note'] == 'external check'
    skill_templates = Path(__file__).resolve().parents[1] / 'templates' / 'rejected' / 'manifest.json'
    with pytest.raises(ValueError, match='unlicensed'):
        importer.import_template(external / 'template.pptx', output=skill_templates, license_note='external check')


def test_clone_fill_recolor_and_chart_independence(tmp_path):
    from PIL import Image
    sample.build(tmp_path)
    store = office.PresentationStore(tmp_path / 'deck.pptx')
    store.init('deck', tmp_path / 'template.pptx')
    assert [slide['role'] for slide in store.status()['template_index']['slides']] == ['cover', 'section', 'comparison', 'data']
    plan = [('sec-a', 'tpl-002'), ('cmp-a', 'tpl-003'), ('data-a', 'tpl-004'),
        ('sec-b', 'tpl-002'), ('cmp-b', 'tpl-003'), ('data-b', 'tpl-004')]
    for slide_id, source in plan:
        accept_all(store)
        apply(store, 'slide.clone', slide_id=slide_id, source_slide_id=source, insert_after=source)
    slides = [slide['slide_id'] for slide in store.status()['manifest']['slides']]
    assert len(slides) == 10
    assert {'sec-a', 'cmp-a', 'data-a', 'data-b'} <= set(slides)
    fill = dict(presentation_id='deck', operation_id='fill-cover', base_revision=store.status()['revision'],
        slide_id='tpl-001', operation='slot.fill', slots={'title': 'Agent 执行架构', 'subtitle': '基于 Deep Agent CLI'})
    assert store.apply_one(fill) == store.apply_one(fill)
    deck = open_deck(store.snapshot(store.status()['revision']))
    title = next(shape for shape in deck.slides[0].shapes if shape.name == 'da:title')
    assert title.text_frame.text == 'Agent 执行架构'
    assert any(node.get('b') == '1' for node in title._element.iter(A + 'rPr'))
    before_move = {slide._element.cSld.get('name'): next(shape.left for shape in slide.shapes if shape.name == 'da:title')
        for slide in deck.slides if any(shape.name == 'da:title' for shape in slide.shapes)}
    apply(store, 'element.adjust', slide_id='tpl-002', element={'id': 'title', 'dx_in': 0.25})
    moved = open_deck(store.snapshot(store.status()['revision']))
    after_move = {slide._element.cSld.get('name'): next(shape.left for shape in slide.shapes if shape.name == 'da:title')
        for slide in moved.slides if any(shape.name == 'da:title' for shape in slide.shapes)}
    assert after_move['da-slide:tpl-002'] - before_move['da-slide:tpl-002'] == 228600
    assert after_move['da-slide:tpl-003'] == before_move['da-slide:tpl-003']
    apply(store, 'chart.replace_data', slide_id='data-b', element_id='chart',
        chart={'categories': ['Q1', 'Q2', 'Q3'], 'series': [{'name': 'Revenue', 'values': [4, 5, 6]}]})
    charted = open_deck(store.snapshot(store.status()['revision']))
    values = {}
    for slide in charted.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                values[slide._element.cSld.get('name')] = tuple(shape.chart.series[0].values)
    assert values['da-slide:tpl-004'] == (10, 20, 30)
    assert values['da-slide:data-b'] == (4, 5, 6)
    assert len({str(shape.chart.part.partname) for slide in charted.slides for shape in slide.shapes if shape.has_chart}) >= 2
    Image.new('RGB', (400, 200), '#335577').save(tmp_path / 'replacement.png')
    apply(store, 'image.replace', slide_id='data-b', element_id='photo', path='replacement.png')
    before_color = store.path.read_bytes()
    with pytest.raises(ValueError, match='capacity conflict'):
        apply(store, 'slot.fill', slide_id='tpl-004', slots={'footnote': '这段脚注明显超过了原有文本框能够容纳的行数。'})
    assert store.path.read_bytes() == before_color
    with pytest.raises(ValueError, match='palette'):
        apply(store, 'presentation.recolor', palette={'accent': '707780'})
    assert store.path.read_bytes() == before_color
    colored = apply(store, 'presentation.recolor', palette={'accent': '707780', 'background': 'FFFFFF', 'text': '20242A', 'muted': '79818A'})
    assert colored['recolor']['recolor_complete'] is False
    assert any(item['color'] == 'C45C26' for item in colored['recolor']['unhandled'])
    report = importer.consistency_report(before_color, store.path.read_bytes())
    assert report['geometry_changes'] == []
    assert report['text_changes'] == []
    assert report['pictures_before'] == report['pictures_after']
    assert report['charts_before'] == report['charts_after']
    final = open_deck(store.snapshot(store.status()['revision']))
    logo = next(shape for shape in final.slides[0].shapes if shape.name == 'da:logo')
    band = next(shape for slide in final.slides for shape in slide.shapes if shape.name == 'da:band')
    note = next(shape for slide in final.slides for shape in slide.shapes if shape.name == 'da:note')
    assert '0B2545' in colors(logo)
    assert '707780' in colors(band)
    assert 'C45C26' in colors(note)
    office.validator.validate_package(store.path.read_bytes())
