import copy
import importlib.util
import io
import json
from pathlib import Path
import sys
import zipfile

import pytest

ROOT=Path(__file__).resolve().parents[1]
SCRIPT=ROOT/'scripts'/'office_pptx.py'
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
from design import design_quality,layout_registry,layout_resolver,profile_registry
spec=importlib.util.spec_from_file_location('ppt_design_engine',SCRIPT)
engine=importlib.util.module_from_spec(spec);spec.loader.exec_module(engine)
renderer=engine.renderer


def op(store,kind,sid=None,**kwargs):
    state=store.status();request=dict(presentation_id='deck',operation_id='op-'+str(state['revision']+1),base_revision=state['revision'],operation=kind)
    if sid:request['slide_id']=sid
    request.update(kwargs)
    return store.apply_one(request)


def initialized(tmp_path,profile='corporate_blue',template=None):
    store=engine.PresentationStore(tmp_path/'deck.pptx')
    store.init('deck',template,profile=profile)
    return store


def accept(store,sid):
    with store.locked() as(c,state):
        state['reviews'][sid]=dict(automatic_passed=True,status='pending',visual_review='pending',revision=state['revision'],content_notes='fixture')
        store.save_state(c,state);c.commit()


def slide_view(store,sid='s1'):
    return next(slide for slide in store.status()['manifest']['slides'] if slide['slide_id']==sid)


def add_component(store,element,sid='s1'):
    return op(store,'element.add_component',sid=sid,element=element)


def base_component(**overrides):
    value=dict(id='flow',component='process.chevron_flow',params={'steps':4},
        texts={'step1':'需求确认','step2':'方案设计','step3':'试点验证','step4':'规模推广'},
        geometry=dict(x=.06,y=.58,w=.88,h=.88*13.333333/5/7.5))
    value.update(overrides)
    return value


# ---------------------------------------------------------------- profiles ---

def test_profiles_are_data_driven_and_extendable(tmp_path):
    registry=profile_registry.load_registry()
    assert registry.ids()==['consulting','corporate_blue','financial','minimal_keynote','slate_dark','tech_dark']
    assert registry.palette('corporate_blue')['accent']=='246BCE'
    assert registry.style('slate_dark','body')['font_size']==20
    consulting=registry.get('consulting')
    assert consulting['extends']=='corporate_blue'
    assert consulting['colors']['accent']=='238F94'
    assert consulting['type_scale']['body']['size']==18
    assert consulting['quality']==registry.get('corporate_blue')['quality']
    assert registry.get('tech_dark')['colors']['background']=='080D18'
    assert set(registry.style('financial','caption'))>= {'font_family','font_size','bold','color','line_height'}


def write_profile(directory,name,base='corporate_blue',**overrides):
    data=json.loads((ROOT/'design'/'profiles'/(base+'.json')).read_text())
    data['id']=name;data['name']='Test '+name;data['extends']=None
    for key,value in overrides.items():data[key]=value
    (directory/f'{name}.json').write_text(json.dumps(data),encoding='utf-8')
    return data


def test_profile_schema_rejects_invalid_definitions(tmp_path):
    import shutil
    directory=tmp_path/'profiles';directory.mkdir()
    shutil.copy(ROOT/'design'/'profiles'/'corporate_blue.json',directory/'corporate_blue.json')
    data=write_profile(directory,'broken')
    data['colors']['accent']='blue'
    (directory/'broken.json').write_text(json.dumps(data))
    with pytest.raises(profile_registry.ProfileError):profile_registry.ProfileRegistry(directory)
    (directory/'broken.json').unlink()
    data=write_profile(directory,'broken',type_scale={'body':{'size':20,'bold':False,'color_token':'text','line_height':1.2}})
    (directory/'broken.json').write_text(json.dumps(data))
    with pytest.raises(profile_registry.ProfileError):profile_registry.ProfileRegistry(directory)
    (directory/'broken.json').unlink()
    data=write_profile(directory,'cycle_a',extends='cycle_b')
    (directory/'cycle_a.json').write_text(json.dumps(data))
    data=write_profile(directory,'cycle_b',extends='cycle_a')
    (directory/'cycle_b.json').write_text(json.dumps(data))
    with pytest.raises(profile_registry.ProfileError):profile_registry.ProfileRegistry(directory)
    with pytest.raises(profile_registry.ProfileError):profile_registry.load_registry().get('not_a_profile')


def test_new_profile_applies_without_engine_changes(tmp_path):
    custom=tmp_path/'profiles';custom.mkdir()
    import shutil
    shutil.copy(ROOT/'design'/'profiles'/'corporate_blue.json',custom/'corporate_blue.json')
    data=json.loads((ROOT/'design'/'profiles'/'corporate_blue.json').read_text())
    data.update(id='warm_report',name='Warm report',extends='corporate_blue')
    data['colors']=dict(data['colors'],background='FFF9F2',accent='B4532A',text='3B2418')
    (custom/'warm_report.json').write_text(json.dumps(data),encoding='utf-8')
    registry=profile_registry.ProfileRegistry(custom)
    assert registry.palette('warm_report')['accent']=='B4532A'
    assert registry.style('warm_report','body')['color']=='3B2418'


# ----------------------------------------------------------------- layouts ---

def test_legacy_and_semantic_layouts_are_registered():
    registry=layout_registry.load_layouts()
    assert set(layout_registry.LEGACY_LAYOUTS)<=set(registry.ids())
    for name in ('statement','hero','bento','evidence','kpi','timeline','architecture','matrix'):
        assert name in registry.ids()
    key,slots=registry.variant_for('bento',3)
    assert key=='3' and len([name for name,slot in slots.items() if slot['role']=='module'])==3
    assert registry.variant_for('bento',5)==(None,None)
    assert registry.get('kpi')['module_counts']==[3,4]


def test_layout_resolution_is_deterministic_and_profile_aware():
    profiles=profile_registry.load_registry();layouts=layout_registry.load_layouts()
    content={'blocks':[{'type':'text','chars':120},{'type':'chart'}],'modules':0}
    first=layout_resolver.plan(profiles,layouts,'consulting','data',content)
    second=layout_resolver.plan(profiles,layouts,'consulting','data',content)
    assert first==second
    assert first['layout']=='data'
    corporate=layout_resolver.plan(profiles,layouts,'corporate_blue','kpi',{'blocks':[{'type':'text','chars':30}],'modules':4})
    consulting=layout_resolver.plan(profiles,layouts,'consulting','kpi',{'blocks':[{'type':'text','chars':30}],'modules':4})
    assert corporate['layout']==consulting['layout']=='kpi'
    assert corporate['slots']['metric1']['rect']==[.06,.22,.2,.14]
    assert consulting['slots']['metric1']['rect'][0]>corporate['slots']['metric1']['rect'][0]
    assert corporate['slots']['metric1']['style_token']=='metric'
    with pytest.raises(layout_resolver.ResolutionError):
        layout_resolver.plan(profiles,layouts,'corporate_blue','data',{'blocks':[{'type':'text','chars':4000},{'type':'chart'}]})
    with pytest.raises(layout_resolver.ResolutionError):
        layout_resolver.plan(profiles,layouts,'corporate_blue','modules',{'blocks':[{'type':'text','chars':40}],'modules':7})


def test_layout_variants_store_modules_and_migrate_losslessly(tmp_path):
    store=initialized(tmp_path)
    op(store,'slide.add',sid='s1',layout='kpi',modules=4)
    op(store,'element.add_text',sid='s1',element=dict(id='m1',anchor='metric1',text='38%'))
    slide=slide_view(store)
    assert slide['layout']=='kpi' and slide['layout_modules']==4 and slide['layout_variant']=='4'
    op(store,'slide.set_layout',sid='s1',layout='kpi',modules=3)
    slide=slide_view(store)
    assert slide['layout_modules']==3 and slide['layout_variant']=='3'
    assert next(e for e in slide['elements'] if e['id']=='m1')['geometry']['w']==.26
    op(store,'element.add_text',sid='s1',element=dict(id='note',anchor='note',text='说明'))
    with pytest.raises(ValueError):op(store,'slide.set_layout',sid='s1',layout='statement')
    assert {e['id'] for e in slide_view(store)['elements']}=={'m1','note'}


# -------------------------------------------------------------- components ---

def test_component_registry_is_reviewed_and_safe():
    import component_validator as validator
    registry=validator.load_registry()
    assert len(registry['components'])>=20
    entry=validator.get(registry,'process.chevron_flow')
    assert entry['variants']['4']['svg'].endswith('chevron_flow_4.svg')
    assert entry['baked_text'] is False
    with pytest.raises(validator.ComponentError):validator.get(registry,'process.unknown')
    bad=b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10" viewBox="0 0 10 10"><script>alert(1)</script></svg>'
    with pytest.raises(validator.ComponentError):validator.validate_svg(bad,'test',('accent',),())
    external=b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10" viewBox="0 0 10 10"><rect width="10" height="10" fill="url(https://example.com/x)"/></svg>'
    with pytest.raises(validator.ComponentError):validator.validate_svg(external,'test',('accent',),())
    text=b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10" viewBox="0 0 10 10"><text x="1" y="5">hi</text></svg>'
    with pytest.raises(validator.ComponentError):validator.validate_svg(text,'test',('accent',),())
    stray=b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10" viewBox="0 0 10 10"><rect width="10" height="10" fill="{accent}"/></svg>'
    with pytest.raises(validator.ComponentError):validator.validate_svg(stray,'test',('accent',),())


def test_component_parameters_text_and_previews_are_validated():
    import component_renderer as renderer_module
    import component_validator as validator
    profiles=profile_registry.load_registry();registry=validator.load_registry()
    result=renderer_module.render(registry,profiles,'corporate_blue','process.chevron_flow',{'steps':5},{'step1':'A'})
    assert result['variant']=='5'
    repeat=renderer_module.render(registry,profiles,'corporate_blue','process.chevron_flow',{'steps':5},{'step1':'A'})
    assert result['png']==repeat['png'] and result['svg_sha256']==repeat['svg_sha256']
    with pytest.raises(ValueError):renderer_module.render(registry,profiles,'corporate_blue','process.chevron_flow',{'steps':9},{})
    with pytest.raises(ValueError):renderer_module.render(registry,profiles,'corporate_blue','process.chevron_flow',{'unknown':1},{})
    with pytest.raises(ValueError):renderer_module.render(registry,profiles,'corporate_blue','process.chevron_flow',{}, {'step1':'x'*40})
    with pytest.raises(ValueError):renderer_module.render(registry,profiles,'corporate_blue','process.chevron_flow',{}, {'nope':'x'})
    entry=validator.get(registry,'architecture.layer_stack')
    assert entry['variants']['3']['svg'].endswith('layer_stack_3.svg')
    assert renderer_module.render(registry,profiles,'financial','icons.check',{})['png'][:8]==b'\x89PNG\r\n\x1a\n'


def test_component_instances_keep_native_text_and_traceability(tmp_path):
    store=initialized(tmp_path)
    op(store,'slide.add',sid='s1',layout='title_body')
    add_component(store,base_component())
    slide=slide_view(store)
    ids={element['id'] for element in slide['elements']}
    assert {'flow','flow.step1','flow.step4'}<=ids
    flow=next(element for element in slide['elements'] if element['id']=='flow')
    assert flow['type']=='component' and flow['component']['params']=={'steps':4}
    assert flow['component']['svg_template_sha256']==engine.component_validator_get('process.chevron_flow')['variants']['4']['svg_sha256']
    child=next(element for element in slide['elements'] if element['id']=='flow.step1')
    assert child['type']=='text' and child['text']=='需求确认' and child['component_slot']=='step1'
    assert flow['vector'] is True
    import hashlib
    with zipfile.ZipFile(store.path) as package:
        media=[name for name in package.namelist() if name.startswith('ppt/media/') and name.endswith('.svg')]
        assert len(media)==1
        assert hashlib.sha256(package.read(media[0])).hexdigest()==flow['component']['svg_sha256']
        assert b'image/svg+xml' in package.read('[Content_Types].xml')
    with zipfile.ZipFile(store.path) as package:
        assert b'asvg:svgBlip' in package.read('ppt/slides/slide1.xml')
    prs=engine.Presentation(io.BytesIO(store.path.read_bytes()))
    assert next(shape for shape in prs.slides[0].shapes if shape.name=='da:flow').image.content_type=='image/png'
    op(store,'element.update',sid='s1',element=dict(id='flow',texts={'step2':'方案复核'}))
    slide=slide_view(store)
    assert {e['id']:e.get('text') for e in slide['elements'] if e.get('component_instance')}=={
        'flow.step1':'需求确认','flow.step2':'方案复核','flow.step3':'试点验证','flow.step4':'规模推广'}
    snapshot_revision=store.status()['revision']
    op(store,'element.update',sid='s1',element=dict(id='flow',params={'steps':5},texts={'step1':'A','step2':'B','step3':'C','step4':'D','step5':'E'}))
    slide=slide_view(store)
    flow=next(element for element in slide['elements'] if element['id']=='flow')
    assert flow['component']['variant']=='5' and flow['component']['version']=='1.0.0'
    store.restore(snapshot_revision,store.status()['revision'],'restore-component')
    slide=slide_view(store)
    flow=next(element for element in slide['elements'] if element['id']=='flow')
    assert flow['component']['variant']=='4' and next(e for e in slide['elements'] if e['id']=='flow.step2')['text']=='方案复核'
    with zipfile.ZipFile(store.path) as package:
        assert len([name for name in package.namelist() if name.endswith('.svg')])==1
    op(store,'element.update',sid='s1',element=dict(id='flow',params={'steps':5},texts={'step1':'A','step2':'B','step3':'C','step4':'D','step5':'E'}))
    slide=slide_view(store)
    assert {element['id'] for element in slide['elements'] if element.get('component_instance')}=={'flow.step1','flow.step2','flow.step3','flow.step4','flow.step5'}
    with zipfile.ZipFile(store.path) as package:
        assert len([name for name in package.namelist() if name.endswith('.svg')])==1
    op(store,'element.remove',sid='s1',element=dict(id='flow'))
    assert not [element for element in slide_view(store)['elements'] if element.get('component_instance') or element['id']=='flow']
    with zipfile.ZipFile(store.path) as package:
        assert not [name for name in package.namelist() if name.endswith('.svg')]


def test_component_calls_reject_bad_geometry_and_versions(tmp_path):
    store=initialized(tmp_path)
    op(store,'slide.add',sid='s1',layout='title_body')
    with pytest.raises(ValueError):add_component(store,base_component(component='process.missing'))
    with pytest.raises(ValueError):add_component(store,base_component(params={'steps':6}))
    with pytest.raises(ValueError):add_component(store,base_component(version='9.9.9'))
    with pytest.raises(ValueError):add_component(store,base_component(geometry=dict(x=.1,y=.1,w=.5,h=.1)))
    with pytest.raises(ValueError):add_component(store,base_component(texts={'step1':'x'*400}))
    with pytest.raises(ValueError):add_component(store,base_component(geometry=dict(x=.6,y=.6,w=.2,h=.2)))
    assert not slide_view(store)['elements']


def test_profile_switch_rerenders_components_and_keeps_native_text(tmp_path):
    store=initialized(tmp_path,'consulting')
    op(store,'slide.add',sid='s1',layout='architecture',modules=4)
    add_component(store,base_component(component='architecture.layer_stack',params={'layers':4},
        texts={'layer1':'交互层','layer2':'执行层','layer3':'能力层','layer4':'资源层'},
        geometry=dict(x=.06,y=.22,w=.5,h=.5*13.333333/1.25/7.5)))
    before=next(element for element in slide_view(store)['elements'] if element['id']=='flow')
    op(store,'presentation.set_profile',profile='tech_dark')
    slide=slide_view(store)
    after=next(element for element in slide['elements'] if element['id']=='flow')
    assert slide['background']=='080D18'
    assert after['component']['svg_sha256']!=before['component']['svg_sha256']
    assert after['component']['profile_id']=='tech_dark'
    assert next(element for element in slide['elements'] if element['id']=='flow.layer1')['text']=='交互层'
    assert all(item['review']['status']=='writing' for item in store.status()['manifest']['slides'])


def test_manifest_exposes_profile_tokens_and_layouts(tmp_path):
    store=initialized(tmp_path,'minimal_keynote')
    op(store,'slide.add',sid='s1',layout='statement')
    manifest=store.status()['manifest']
    assert manifest['profile']['id']=='minimal_keynote'
    assert manifest['profile']['colors']['accent']=='7496D5'
    assert manifest['profile']['type_scale']['statement']['size']==44
    assert manifest['slides'][0]['layout']=='statement'
    assert manifest['slides'][0]['profile_id']=='minimal_keynote'


# ---------------------------------------------------------- design quality ---

def text_element(eid='t1',text='内容',size=20,color='17324D',bold=False,rect=(.1,.1,.5,.2),token='body',anchor=None,family='Noto Sans CJK SC'):
    return dict(id=eid,type='text',text=text,geometry=dict(zip(('x','y','w','h'),rect)),layer='content',
        style=dict(font_family=family,font_size=size,bold=bold,color=color,line_height=1.22),style_token=token,
        anchor=anchor,source='generated')


def quality_slide(elements,background='FFFFFF',layout='title_body'):
    return dict(slide_id='s1',elements=elements,background=background,layout=layout)


def test_design_quality_detects_small_text_and_low_contrast():
    profile=profile_registry.load_registry().get('corporate_blue')
    issues=design_quality.review(quality_slide([text_element(size=7)]),profile)
    assert any(issue['type']=='font_too_small' and issue['severity']=='error' for issue in issues)
    issues=design_quality.review(quality_slide([text_element(color='F2F2F2')]),profile)
    assert any(issue['type']=='contrast_insufficient' and issue['severity']=='error' for issue in issues)
    issues=design_quality.review(quality_slide([text_element()]),profile)
    assert not [issue for issue in issues if issue['type']=='contrast_insufficient']
    summary=design_quality.summary(issues)
    assert summary['errors']==0 and set(summary)=={'errors','warnings','advisories'}


def test_design_quality_respects_slot_capacity_and_components():
    profiles=profile_registry.load_registry();components=engine.registry()
    profile=profiles.get('corporate_blue')
    slots={'body':{'capacity':{'chars':10,'lines':1},'role':'body','rect':[.1,.1,.5,.2]}}
    issues=design_quality.review(quality_slide([text_element(text='超过容量限制的文字内容',anchor='body')]),profile,components,slots)
    assert any(issue['type']=='slot_capacity_exceeded' for issue in issues)
    component=dict(id='c',type='component',geometry=dict(x=.1,y=.1,w=.5,h=.2),layer='content',source='generated',
        component=dict(component_id='process.chevron_flow',version='1.0.0',variant='4',params={'steps':4},
            texts={'step1':'x'*99},svg_sha256='x',svg_template_sha256='x'))
    issues=design_quality.review(quality_slide([component]),profile,components)
    assert any(issue['type']=='component_capacity_exceeded' and issue['severity']=='error' for issue in issues)
    component['component']['component_id']='process.missing'
    issues=design_quality.review(quality_slide([component]),profile,components)
    assert any(issue['type']=='component_unregistered' and issue['severity']=='error' for issue in issues)


def test_decoration_overlap_is_intentional_but_content_overlap_is_not(tmp_path):
    store=initialized(tmp_path)
    op(store,'slide.add',sid='s1',layout='title_body')
    op(store,'element.add_text',sid='s1',element=dict(id='a',geometry=dict(x=.1,y=.3,w=.4,h=.2),text='第一段'))
    op(store,'element.add_text',sid='s1',element=dict(id='b',geometry=dict(x=.1,y=.3,w=.4,h=.2),text='第二段'))
    issues=engine.validator.layout_issues(slide_view(store))
    assert any(issue['type']=='unexpected_overlap' for issue in issues)
    op(store,'element.add_shape',sid='s1',element=dict(id='wash',geometry=dict(x=.05,y=.25,w=.5,h=.3),shape='rectangle',fill_token='card',layer='decoration'))
    issues=engine.validator.layout_issues(slide_view(store))
    assert not [issue for issue in issues if issue.get('element_id')=='wash' or 'wash' in issue.get('elements',[])]


def test_design_quality_flags_unreadable_margins_and_chart_metadata():
    profile=profile_registry.load_registry().get('corporate_blue')
    element=text_element(rect=(.01,.1,.2,.1))
    issues=design_quality.review(quality_slide([element]),profile)
    assert any(issue['type']=='safe_area_violation' and issue['severity']=='warning' for issue in issues)
    chart=dict(id='c',type='chart',geometry=dict(x=.1,y=.3,w=.5,h=.4),layer='content',source='generated',
        chart=dict(type='bar',categories=['A','B'],series=[dict(name='',values=[1,2])]))
    issues=design_quality.review(quality_slide([chart]),profile)
    assert any(issue['type']=='chart_series_unlabeled' and issue['severity']=='error' for issue in issues)
    chart['chart']['series'][0]['name']='交付量'
    issues=design_quality.review(quality_slide([chart]),profile)
    assert any(issue['type']=='chart_unit_missing' and issue['severity']=='advisory' for issue in issues)


# ------------------------------------------------------------- integration ---

def brand_template(tmp_path,accent='123456',font='WenQuanYi Micro Hei'):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from lxml import etree
    template=tmp_path/'template.pptx';prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6]);brand=slide.shapes.add_textbox(Inches(.8),Inches(.18),Inches(5),Inches(.28));brand.text='企业模板品牌'
    for run in brand.text_frame.paragraphs[0].runs:run.font.name='Noto Sans CJK SC';run.font.size=Pt(11)
    prs.save(template)
    with zipfile.ZipFile(template) as z:parts={name:z.read(name) for name in z.namelist()}
    theme=etree.fromstring(parts['ppt/theme/theme1.xml']);ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    theme.xpath('//a:clrScheme/a:accent1/a:srgbClr',namespaces=ns)[0].set('val',accent)
    theme.xpath('//a:fontScheme/a:minorFont/a:ea',namespaces=ns)[0].set('typeface',font)
    parts['ppt/theme/theme1.xml']=etree.tostring(theme,xml_declaration=True,encoding='UTF-8')
    with zipfile.ZipFile(template,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    return template


def test_profile_precedence_for_templates_and_restore(tmp_path):
    template=brand_template(tmp_path)
    store=initialized(tmp_path,'consulting',template)
    sid=store.status()['manifest']['slides'][0]['slide_id']
    first=op(store,'element.add_text',sid=sid,element=dict(id='note',geometry=dict(x=.2,y=.6,w=.5,h=.2),text='新增原生文字'))
    status=store.status();assert status['theme']['accent']=='123456'
    element=next(element for element in store.status()['manifest']['slides'][0]['elements'] if element['id']=='note')
    assert element['style']['font_family']=='WenQuanYi Micro Hei'
    with zipfile.ZipFile(template) as a,zipfile.ZipFile(store.path) as b:
        for name in a.namelist():
            if name.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/','ppt/media/')):assert a.read(name)==b.read(name),name
    op(store,'element.add_text',sid=sid,element=dict(id='note2',geometry=dict(x=.2,y=.8,w=.5,h=.1),text='第二段'))
    restored=store.restore(first['revision'],store.status()['revision'],'restore-design')
    assert restored['revision']==store.status()['revision']
    assert 'note2' not in {element['id'] for element in store.status()['manifest']['slides'][0]['elements']}
    assert store.status()['profile_id']=='consulting'
    assert store.status()['theme']['accent']=='123456'


def test_conflicting_template_accent_is_reported_not_mixed(tmp_path):
    template=brand_template(tmp_path,accent='246BCE')
    store=initialized(tmp_path,'slate_dark',template)
    status=store.status()
    assert status['theme']['accent']=='66B5FF'
    conflicts=status['manifest']['conflicts']
    assert conflicts and conflicts[0]['kind']=='template_accent_contrast'
    assert conflicts[0]['accent']=='246BCE' and conflicts[0]['kept_accent']=='66B5FF'
    sid=status['manifest']['slides'][0]['slide_id']
    op(store,'element.add_text',sid=sid,element=dict(id='title',geometry=dict(x=.1,y=.1,w=.5,h=.2),text='深色标题',style='card_title'))
    element=next(element for element in store.status()['manifest']['slides'][0]['elements'] if element['id']=='title')
    assert element['style']['color']=='66B5FF'


def test_dark_profile_tables_and_charts_stay_readable(tmp_path):
    store=initialized(tmp_path,'slate_dark')
    op(store,'slide.add',sid='s1',layout='title_body')
    op(store,'element.add_text',sid='s1',element=dict(id='title',anchor='title',text='深色表格与图表',style='slide_title'))
    op(store,'element.add_table',sid='s1',element=dict(id='table',geometry=dict(x=.08,y=.26,w=.84,h=.3),
        rows=[['维度','现状','目标'],['周期','单次批量','原子增量']]))
    op(store,'element.add_chart',sid='s1',element=dict(id='chart',geometry=dict(x=.08,y=.6,w=.5,h=.24),
        chart=dict(type='bar',categories=['A','B'],series=[dict(name='进度',values=[3,5])],unit='%')))
    result=store.check_slide('s1','核对表格内容与图表数据。')
    assert result['review']['automatic_passed'],result['review']['layout_issues']
    prs=engine.Presentation(io.BytesIO(store.snapshot(store.status()['revision'])))
    table=next(shape for shape in prs.slides[0].shapes if shape.has_table).table
    assert str(table.cell(1,0).fill.fore_color.rgb)=='14263D'
    assert str(table.cell(0,0).fill.fore_color.rgb)=='203B58'
    assert str(table.cell(1,0).text_frame.paragraphs[0].runs[0].font.color.rgb)=='FFFFFF'
    chart=next(shape for shape in prs.slides[0].shapes if shape.has_chart).chart
    assert str(chart.value_axis.tick_labels.font.color.rgb)=='BCD0E4'
    assert str(chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.color.rgb)=='BCD0E4'


def test_state_created_before_the_design_system_still_edits(tmp_path):
    store=initialized(tmp_path)
    op(store,'slide.add',sid='s1',layout='title_body')
    op(store,'element.add_text',sid='s1',element=dict(id='title',anchor='title',text='旧版状态'))
    with store.locked() as(c,state):
        for key in ('profile','profile_id','theme_name','layout_slots','layout_variants','layout_modules'):
            state.pop(key,None)
        state['theme']=dict(background='FFFFFF',text='17324D',accent='246BCE',muted='536B82',card='EDF4FC')
        state['reviews']={'s1':dict(automatic_passed=True,status='passed',visual_review='passed',content_notes='legacy')}
        store.save_state(c,state);c.commit()
    op(store,'element.add_text',sid='s1',element=dict(id='body',anchor='body',text='迁移后的新增正文'))
    slide=slide_view(store)
    assert {'title','body'}<={element['id'] for element in slide['elements']}
    assert store.status()['profile_id']=='corporate_blue'
    assert slide['layout_variant'] is None
    assert slide['review']['status']=='writing'
