"""Native PowerPoint flow / architecture diagrams: shapes for nodes, bound connectors for edges.

The caller declares nodes, edges, relation types and a direction. This module
owns all geometry: rank assignment, node placement, connection sites,
connector kind, arrowheads (from the relation), orthogonal routing and simple
obstacle avoidance. Free-form connector coordinates are never accepted, and an
edge that cannot be drawn clearly is rejected instead of being approximated.
"""
import math

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from design import design_quality, diagram_presets

EMU_PER_IN = 914400
SITES = {'top': 0, 'left': 1, 'bottom': 2, 'right': 3}
NODE_KEYS = {'id', 'label', 'detail', 'rank', 'emphasis'}
EDGE_KEYS = {'id', 'source', 'target', 'relation', 'label'}
SPEC_KEYS = {'id', 'geometry', 'anchor', 'direction', 'preset', 'nodes', 'edges'}
PATCH_KEYS = {'add_nodes', 'remove_nodes', 'update_nodes', 'add_edges', 'remove_edges'}
TOLERANCE_IN = .01

HOST = {}


def configure(**callbacks):
    """office_pptx injects renderer/style helpers to avoid a circular import."""
    HOST.update(callbacks)


def _renderer():
    return HOST['renderer']


# ---------------------------------------------------------------- validation

def _text(value, limit, field):
    if not isinstance(value, str) or not value.strip() or len(value) > limit or '\n' in value:
        raise ValueError(f'{field} must be one line of 1 to {limit} characters')
    return value.strip()


def normalize(spec):
    """Validated, canonical diagram description. Relations are mandatory."""
    if not isinstance(spec, dict):
        raise ValueError('diagram must be an object')
    unknown = set(spec) - SPEC_KEYS
    if unknown:
        raise ValueError('Unsupported diagram field(s): ' + ', '.join(sorted(unknown)) + '; connector coordinates are computed, never supplied')
    identity = HOST['identifier'](spec.get('id'))
    preset_name = spec.get('preset') or ('layered' if spec.get('direction', 'TB') == 'TB' else 'flow')
    preset = diagram_presets.preset(preset_name)
    direction = spec.get('direction', preset['direction'])
    if direction not in ('TB', 'LR'):
        raise ValueError('direction must be TB (top to bottom) or LR (left to right)')
    limits = diagram_presets.LIMITS
    nodes = spec.get('nodes')
    if not isinstance(nodes, list) or not 2 <= len(nodes) <= limits['nodes']:
        raise ValueError(f'A diagram needs 2 to {limits["nodes"]} nodes')
    seen = set()
    result_nodes = []
    for item in nodes:
        if not isinstance(item, dict) or set(item) - NODE_KEYS:
            raise ValueError('Node fields are id, label, detail, rank and emphasis')
        node_id = HOST['identifier'](item.get('id'))
        if node_id in seen:
            raise ValueError(f'Duplicate node id {node_id}')
        seen.add(node_id)
        node = dict(id=node_id, label=_text(item.get('label'), limits['label_chars'], f'node {node_id} label'))
        if item.get('detail') is not None:
            node['detail'] = _text(item['detail'], limits['detail_chars'], f'node {node_id} detail')
        if 'rank' in item:
            if type(item['rank']) is not int or not 0 <= item['rank'] < limits['ranks']:
                raise ValueError(f'node {node_id} rank must be an integer from 0 to {limits["ranks"] - 1}')
            node['rank'] = item['rank']
        if item.get('emphasis') is not None:
            if type(item['emphasis']) is not bool:
                raise ValueError('emphasis must be true or false')
            node['emphasis'] = item['emphasis']
        result_nodes.append(node)
    if 0 < sum('rank' in node for node in result_nodes) < len(result_nodes):
        raise ValueError('Give every node a rank or none; partial ranks are ambiguous')
    edges = spec.get('edges', [])
    if not isinstance(edges, list) or len(edges) > limits['edges']:
        raise ValueError(f'A diagram holds at most {limits["edges"]} edges')
    result_edges = []
    pairs = {}
    edge_ids = set()
    for item in edges:
        if not isinstance(item, dict) or set(item) - EDGE_KEYS:
            raise ValueError('Edge fields are id, source, target, relation and label; coordinates are not accepted')
        source, target = item.get('source'), item.get('target')
        if source not in seen or target not in seen:
            raise ValueError(f'Edge {item.get("id") or (source, target)} must connect two declared nodes')
        if source == target:
            raise ValueError(f'Edge from {source} to itself has no meaning')
        if 'relation' not in item:
            raise ValueError(f'Edge {source} -> {target} needs a relation ({", ".join(sorted(diagram_presets.RELATIONS))})')
        diagram_presets.relation(item['relation'])
        edge_id = HOST['identifier'](item.get('id') or f'e-{source}-{target}')
        if edge_id in edge_ids or edge_id in seen:
            raise ValueError(f'Duplicate edge id {edge_id}')
        edge_ids.add(edge_id)
        key = frozenset((source, target))
        if key in pairs:
            raise ValueError(f'Nodes {source} and {target} already have edge {pairs[key]}; use one bidirectional edge or a label instead of parallel connectors')
        pairs[key] = edge_id
        edge = dict(id=edge_id, source=source, target=target, relation=item['relation'])
        if item.get('label') is not None:
            edge['label'] = _text(item['label'], limits['edge_label_chars'], f'edge {edge_id} label')
        result_edges.append(edge)
    result = dict(id=identity, preset=preset_name, direction=direction, nodes=result_nodes, edges=result_edges)
    if spec.get('anchor'):
        result['anchor'] = spec['anchor']
    elif spec.get('geometry'):
        result['geometry'] = spec['geometry']
    else:
        raise ValueError('A diagram needs geometry or a layout anchor for its area')
    return result


def assign_ranks(spec):
    nodes = spec['nodes']
    if all('rank' in node for node in nodes):
        return {node['id']: node['rank'] for node in nodes}
    forward = [(e['source'], e['target']) for e in spec['edges']
               if diagram_presets.RELATIONS[e['relation']]['directed'] and not diagram_presets.RELATIONS[e['relation']]['back']]
    order = [node['id'] for node in nodes]
    incoming = {node: 0 for node in order}
    for _, target in forward:
        incoming[target] += 1
    rank = {node: 0 for node in order}
    queue = [node for node in order if incoming[node] == 0]
    visited = 0
    while queue:
        current = queue.pop(0)
        visited += 1
        for source, target in forward:
            if source == current:
                rank[target] = max(rank[target], rank[current] + 1)
                incoming[target] -= 1
                if incoming[target] == 0:
                    queue.append(target)
    if visited != len(order):
        raise ValueError('Directed edges form a cycle; mark the backward edge as relation "returns" or give every node an explicit rank')
    if max(rank.values()) >= diagram_presets.LIMITS['ranks']:
        raise ValueError(f'The diagram is deeper than {diagram_presets.LIMITS["ranks"]} ranks; split it across slides')
    return rank


# ---------------------------------------------------------------- text fit

def _em(text):
    return sum(1.0 if ord(char) > 0x2E7F else .56 for char in text)


def _lines(text, size_pt, width_in):
    capacity = max(width_in, .05)
    return sum(max(1, math.ceil(_em(line) * size_pt / 72 * 1.05 / capacity)) for line in text.split('\n'))


def text_height(node, label_style, detail_style, width_in):
    usable = width_in - .18
    height = _lines(node['label'], label_style['font_size'], usable) * label_style['font_size'] * label_style.get('line_height', 1.2) / 72
    if node.get('detail'):
        height += .04 + _lines(node['detail'], detail_style['font_size'], usable) * detail_style['font_size'] * detail_style.get('line_height', 1.2) / 72
    return height + .12


# ---------------------------------------------------------------- layout

def label_size(text, size_pt):
    return _em(text) * size_pt / 72 + .14, size_pt * 1.3 / 72 + .06


def layout(spec, ranks, area, preset, label_style, detail_style, edge_label_pt=12):
    """Node rectangles in inches. Ranks run along the main axis of the direction.

    Gaps grow to hold edge labels: same-rank labels widen the node gap, cross-rank
    labels raise the minimum rank gap.
    """
    preset = dict(preset)
    vertical_layout = spec['direction'] == 'TB'
    for edge in spec['edges']:
        if not edge.get('label'):
            continue
        width, height = label_size(edge['label'], edge_label_pt)
        if ranks[edge['source']] == ranks[edge['target']]:
            preset['node_gap_in'] = max(preset['node_gap_in'], (width if vertical_layout else height) + .2)
        else:
            preset['rank_gap_in'] = max(preset['rank_gap_in'], (height if vertical_layout else width) + .16)
    groups = {}
    for node in spec['nodes']:
        groups.setdefault(ranks[node['id']], []).append(node)
    rank_ids = sorted(groups)
    if any(len(items) > diagram_presets.LIMITS['per_rank'] for items in groups.values()):
        raise ValueError(f'At most {diagram_presets.LIMITS["per_rank"]} nodes share a rank; split the diagram')
    x0, y0, width, height = area
    vertical = spec['direction'] == 'TB'
    main, cross = (height, width) if vertical else (width, height)
    main_min, main_max = (preset['min_node_h_in'], preset['max_node_h_in']) if vertical else (preset['min_node_w_in'], preset['max_node_w_in'])
    cross_min, cross_max = (preset['min_node_w_in'], preset['max_node_w_in']) if vertical else (preset['min_node_h_in'], preset['max_node_h_in'])
    count = len(rank_ids)
    gap = preset['rank_gap_in']
    main_size = min(main_max, (main - (count - 1) * gap) / count)
    if main_size + 1e-9 < main_min:
        raise ValueError(f'The diagram area is too small for {count} ranks; enlarge its geometry or split the content')
    rank_gap = min(max(gap, (main - count * main_size) / (count - 1)), 1.2) if count > 1 else 0
    start = (main - (count * main_size + (count - 1) * rank_gap)) / 2

    def cross_size(n):
        return min(cross_max, (cross - (n - 1) * preset['node_gap_in']) / n)
    uniform = min(cross_size(len(items)) for items in groups.values()) if preset['uniform'] else None
    rects = {}
    for index, rank in enumerate(rank_ids):
        items = groups[rank]
        size = uniform if uniform is not None else cross_size(len(items))
        if size + 1e-9 < cross_min:
            raise ValueError(f'Rank {rank} holds {len(items)} nodes, more than the diagram width allows; split the rank or enlarge the area')
        row = len(items) * size + (len(items) - 1) * preset['node_gap_in']
        offset = (cross - row) / 2
        main_pos = start + index * (main_size + rank_gap)
        for position, node in enumerate(items):
            cross_pos = offset + position * (size + preset['node_gap_in'])
            if vertical:
                rect = (x0 + cross_pos, y0 + main_pos, size, main_size)
            else:
                rect = (x0 + main_pos, y0 + cross_pos, main_size, size)
            if text_height(node, label_style, detail_style, rect[2]) > rect[3] + 1e-6:
                raise ValueError(f'Node {node["id"]} text does not fit a {rect[2]:.2f} x {rect[3]:.2f} in node; shorten its label/detail or enlarge the diagram')
            rects[node['id']] = rect
    return rects


# ---------------------------------------------------------------- routing

def site_point(rect, site):
    x, y, w, h = rect
    return {'top': (x + w / 2, y), 'left': (x, y + h / 2), 'bottom': (x + w / 2, y + h), 'right': (x + w, y + h / 2)}[site]


def _segments(points):
    return list(zip(points, points[1:]))


def _segment_hits_rect(a, b, rect, inset=.03):
    x, y, w, h = rect
    left, top, right, bottom = x + inset, y + inset, x + w - inset, y + h - inset
    if right <= left or bottom <= top:
        return False
    (x1, y1), (x2, y2) = a, b
    if abs(x1 - x2) < 1e-9:
        return left < x1 < right and min(y1, y2) < bottom and max(y1, y2) > top
    if abs(y1 - y2) < 1e-9:
        return top < y1 < bottom and min(x1, x2) < right and max(x1, x2) > left
    # Liang-Barsky for completeness.
    t0, t1 = 0.0, 1.0
    dx, dy = x2 - x1, y2 - y1
    for p, q in ((-dx, x1 - left), (dx, right - x1), (-dy, y1 - top), (dy, bottom - y1)):
        if abs(p) < 1e-12:
            if q < 0:
                return False
        else:
            t = q / p
            if p < 0:
                t0 = max(t0, t)
            else:
                t1 = min(t1, t)
            if t0 > t1:
                return False
    return True


def _collinear_overlap(first, second):
    for a, b in _segments(first):
        for c, d in _segments(second):
            if abs(a[0] - b[0]) < 1e-6 and abs(c[0] - d[0]) < 1e-6 and abs(a[0] - c[0]) < .02:
                if min(max(a[1], b[1]), max(c[1], d[1])) - max(min(a[1], b[1]), min(c[1], d[1])) > .02:
                    return True
            if abs(a[1] - b[1]) < 1e-6 and abs(c[1] - d[1]) < 1e-6 and abs(a[1] - c[1]) < .02:
                if min(max(a[0], b[0]), max(c[0], d[0])) - max(min(a[0], b[0]), min(c[0], d[0])) > .02:
                    return True
    return False


def _candidates(edge, rects, ranks, direction):
    """Connection-site pairs and lanes, most natural first."""
    s, t = rects[edge['source']], rects[edge['target']]
    rs, rt = ranks[edge['source']], ranks[edge['target']]
    vertical = direction == 'TB'
    result = []
    if rs != rt:
        forward = rt > rs
        if vertical:
            first, last = ('bottom', 'top') if forward else ('top', 'bottom')
        else:
            first, last = ('right', 'left') if forward else ('left', 'right')
        a, b = site_point(s, first), site_point(t, last)
        axis = 1 if vertical else 0
        lo, hi = sorted((a[axis], b[axis]))
        if hi - lo < .16 or (b[axis] - a[axis] > 0) != forward:
            raise ValueError(f'Nodes {edge["source"]} and {edge["target"]} are too close to connect clearly')
        # Midway first; for rank-skipping edges also the gaps next to the target and the source.
        span = b[axis] - a[axis]
        step = min(.15 / abs(span), 1 / 3)
        lanes = [a[axis] + span / 2, b[axis] - span * step, a[axis] + span * step]
        for lane in lanes:
            if not lo + .06 <= lane <= hi - .06:
                continue
            if vertical:
                points = [a, b] if abs(a[0] - b[0]) < TOLERANCE_IN else [a, (a[0], lane), (b[0], lane), b]
            else:
                points = [a, b] if abs(a[1] - b[1]) < TOLERANCE_IN else [a, (lane, a[1]), (lane, b[1]), b]
            result.append((first, last, points))
    else:
        if vertical:
            first, last = ('right', 'left') if s[0] < t[0] else ('left', 'right')
        else:
            first, last = ('bottom', 'top') if s[1] < t[1] else ('top', 'bottom')
        a, b = site_point(s, first), site_point(t, last)
        if abs(a[0] - b[0]) + abs(a[1] - b[1]) < .16:
            raise ValueError(f'Nodes {edge["source"]} and {edge["target"]} are too close to connect clearly')
        result.append((first, last, [a, b]))
    return result


def route(spec, rects, ranks, obstacles):
    """One orthogonal route per edge, or ValueError naming what blocks it."""
    routes = {}
    for edge in spec['edges']:
        blocked = None
        for first, last, points in _candidates(edge, rects, ranks, spec['direction']):
            hit = None
            for node_id, rect in rects.items():
                if node_id in (edge['source'], edge['target']):
                    continue
                if any(_segment_hits_rect(a, b, rect) for a, b in _segments(points)):
                    hit = 'node ' + node_id
                    break
            if hit is None:
                for name, rect in obstacles:
                    if any(_segment_hits_rect(a, b, rect) for a, b in _segments(points)):
                        hit = 'text ' + name
                        break
            if hit is None:
                for other_id, other in routes.items():
                    other_edge = next(e for e in spec['edges'] if e['id'] == other_id)
                    shared = {edge['source'], edge['target']} & {other_edge['source'], other_edge['target']}
                    if not shared and _collinear_overlap(points, other['points']):
                        hit = 'edge ' + other_id
                        break
            if hit is None:
                routes[edge['id']] = dict(points=points, start_site=first, end_site=last)
                break
            blocked = blocked or hit
        if edge['id'] not in routes:
            raise ValueError(f'Edge {edge["id"]} ({edge["source"]} -> {edge["target"]}) cannot be drawn without crossing {blocked or "other nodes"}; '
                             'reorder or re-rank the nodes, move the text, or express this relation in a note')
    return routes


def label_rects(points, text, size_pt):
    """Candidate label boxes beside the connector's middle segment."""
    width, height = label_size(text, size_pt)
    (x1, y1), (x2, y2) = _segments(points)[len(points) // 2 - 1] if len(points) > 2 else _segments(points)[0]
    mx, my = (x1 + x2) / 2, (y1 + y2) / 2
    if abs(x1 - x2) < 1e-6:
        return [(mx + .06, my - height / 2, width, height), (mx - .06 - width, my - height / 2, width, height)]
    return [(mx - width / 2, my - height - .03, width, height), (mx - width / 2, my + .03, width, height)]


# ---------------------------------------------------------------- OOXML

def _set_xfrm(connector, points, prs):
    a, b = points[0], points[-1]
    element = connector._element
    xfrm = element.spPr.find(qn('a:xfrm'))
    for key in ('flipH', 'flipV', 'rot'):
        xfrm.attrib.pop(key, None)
    cx, cy = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
    dx, dy = abs(b[0] - a[0]), abs(b[1] - a[1])
    elbow = len(points) == 4
    vertical_first = elbow and abs(points[1][0] - a[0]) < 1e-9
    adjust = None
    if vertical_first:
        w, h = dy, dx
        xfrm.set('rot', '5400000')
        if a[0] < cx:
            xfrm.set('flipV', '1')
        if a[1] > cy:
            xfrm.set('flipH', '1')
        adjust = (points[1][1] - a[1]) / (b[1] - a[1])
    else:
        w, h = dx, dy
        if b[0] < a[0]:
            xfrm.set('flipH', '1')
        if b[1] < a[1]:
            xfrm.set('flipV', '1')
        if elbow:
            adjust = (points[1][0] - a[0]) / (b[0] - a[0])
    off, ext = xfrm.find(qn('a:off')), xfrm.find(qn('a:ext'))
    off.set('x', str(int(round((cx - w / 2) * EMU_PER_IN))))
    off.set('y', str(int(round((cy - h / 2) * EMU_PER_IN))))
    ext.set('cx', str(int(round(w * EMU_PER_IN))))
    ext.set('cy', str(int(round(h * EMU_PER_IN))))
    if adjust is not None and abs(adjust - .5) > 1e-4:
        av = element.spPr.find(qn('a:prstGeom')).find(qn('a:avLst'))
        guide = etree.SubElement(av, qn('a:gd'))
        guide.set('name', 'adj1')
        guide.set('fmla', 'val %d' % int(round(adjust * 100000)))


def _bind(connector, start_shape, start_site, end_shape, end_site):
    props = connector._element.find('.//' + qn('p:cNvCxnSpPr'))
    for tag in ('a:stCxn', 'a:endCxn'):
        for old in props.findall(qn(tag)):
            props.remove(old)
    start = etree.SubElement(props, qn('a:stCxn'))
    start.set('id', str(start_shape.shape_id))
    start.set('idx', str(SITES[start_site]))
    end = etree.SubElement(props, qn('a:endCxn'))
    end.set('id', str(end_shape.shape_id))
    end.set('idx', str(SITES[end_site]))


def _arrows(connector, arrow):
    line = connector._element.spPr.find(qn('a:ln'))
    for tag in ('a:headEnd', 'a:tailEnd'):
        for old in line.findall(qn(tag)):
            line.remove(old)
    if arrow in ('both',):
        head = etree.SubElement(line, qn('a:headEnd'))
        head.set('type', 'triangle'); head.set('w', 'med'); head.set('len', 'med')
    if arrow in ('end', 'both'):
        tail = etree.SubElement(line, qn('a:tailEnd'))
        tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')


def _readable(fill, candidates, minimum):
    best = max(candidates, key=lambda color: design_quality.contrast(color, fill) or 0)
    ratio = design_quality.contrast(best, fill) or 0
    if ratio < minimum:
        raise ValueError(f'No profile text colour reaches {minimum}:1 contrast on node fill #{fill}')
    return best


# ---------------------------------------------------------------- build

def _area(spec, state, prs, anchors):
    element = {key: spec[key] for key in ('geometry', 'anchor') if key in spec}
    g = HOST['geometry'](element, anchors)
    width_in, height_in = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    return (g['x'] * width_in, g['y'] * height_in, g['w'] * width_in, g['h'] * height_in), g


def _rect_in(shape):
    return (shape.left / EMU_PER_IN, shape.top / EMU_PER_IN, shape.width / EMU_PER_IN, shape.height / EMU_PER_IN)


def diagram_shapes(slide, diagram_id):
    renderer = _renderer()
    return [shape for shape in slide.shapes if renderer.metadata(shape).get('diagram_id') == diagram_id]


def obstacles_for(slide, diagram_id):
    renderer = _renderer()
    result = []
    for shape in slide.shapes:
        meta = renderer.metadata(shape)
        if meta.get('diagram_id') == diagram_id:
            continue
        has_text = shape.has_text_frame and shape.text_frame.text.strip()
        if has_text and meta.get('layer', 'content') != 'background':
            result.append((renderer.element_identity(shape), _rect_in(shape)))
    return result


def read_spec(slide, diagram_id):
    renderer = _renderer()
    nodes, edges, frame, metas = [], [], None, {}
    for shape in diagram_shapes(slide, diagram_id):
        meta = renderer.metadata(shape)
        metas[renderer.element_identity(shape)] = meta
        if meta.get('diagram_role') == 'node':
            frame = frame or meta.get('diagram')
            node = dict(id=meta['node_id'], label=meta['label'])
            for key in ('detail', 'emphasis'):
                if meta.get(key) is not None:
                    node[key] = meta[key]
            if meta.get('declared_rank') is not None:
                node['rank'] = meta['declared_rank']
            nodes.append((meta.get('order', 0), node))
        elif meta.get('diagram_role') == 'edge':
            edge = dict(id=meta['edge_id'], source=meta['source'], target=meta['target'], relation=meta['relation'])
            if meta.get('label'):
                edge['label'] = meta['label']
            edges.append((meta.get('order', 0), edge))
    if frame is None:
        raise ValueError(f'Unknown diagram {diagram_id}')
    spec = dict(id=diagram_id, preset=frame['preset'], direction=frame['direction'],
        nodes=[node for _, node in sorted(nodes, key=lambda item: item[0])],
        edges=[edge for _, edge in sorted(edges, key=lambda item: item[0])])
    spec.update({key: frame[key] for key in ('geometry', 'anchor') if frame.get(key)})
    return spec, metas


def _remove(slide, diagram_id):
    shapes = diagram_shapes(slide, diagram_id)
    if not shapes:
        return None
    parent = shapes[0]._element.getparent()
    index = min(parent.index(shape._element) for shape in shapes)
    for shape in shapes:
        parent.remove(shape._element)
    return index


def _reinsert(slide, created, index):
    if index is None:
        return
    parent = created[0]._element.getparent()
    for offset, shape in enumerate(created):
        parent.remove(shape._element)
        parent.insert(index + offset, shape._element)


def build(slide, spec, state, prs, anchors, profile, previous=None, rects=None, insert_at=None):
    """Create nodes, then bound connectors and labels; returns the created shapes."""
    renderer, native_style, style_for = HOST['renderer'], HOST['native_style'], HOST['style_for']
    previous = previous or {}
    preset = diagram_presets.preset(spec['preset'])
    area, g = _area(spec, state, prs, anchors)
    label_style = dict(style_for(state, preset['label_token']))
    detail_style = dict(style_for(state, preset['detail_token']))
    ranks = assign_ranks(spec)
    keep_positions = rects is not None
    if rects is None:
        rects = layout(spec, ranks, area, preset, label_style, detail_style, style_for(state, 'caption')['font_size'])
    obstacles = obstacles_for(slide, spec['id'])
    for name, rect in obstacles:
        for node_id, node_rect in rects.items():
            if _overlap(_inset(node_rect, .02), _inset(rect, .02)):
                raise ValueError(f'Diagram node {node_id} would cover text {name}; give the diagram a free area')
    routes = route(spec, rects, ranks, obstacles)
    palette = state['theme']
    quality = profile.get('quality') or {}
    background = palette['background']
    connector_color = palette.get(preset['connector_token'], palette['muted'])
    if (design_quality.contrast(connector_color, background) or 0) < 3.0:
        connector_color = _readable(background, [palette['muted'], palette['text']], 3.0)
    frame = dict(id=spec['id'], preset=spec['preset'], direction=spec['direction'])
    frame.update({key: spec[key] for key in ('geometry', 'anchor') if key in spec})
    created, by_node = [], {}
    group = 'diagram:' + spec['id']
    radius = (profile.get('shape') or {}).get('corner_radius_in', .08)
    for order, node in enumerate(spec['nodes']):
        x, y, w, h = rects[node['id']]
        element_id = spec['id'] + '.' + node['id']
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.adjustments[0] = max(0.0, min(.5, radius / max(min(w, h), .01)))
        shape.name = 'da:' + element_id
        fill_token = preset['emphasis_fill'] if node.get('emphasis') else preset['node_fill']
        fill = palette.get(fill_token, palette['card'])
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        native_style.flat(shape)
        label_color = _readable(fill, [palette['text'], palette['background'], 'FFFFFF'], float(quality.get('min_contrast_normal', 4.5)))
        detail_color = palette['muted'] if (design_quality.contrast(palette['muted'], fill) or 0) >= float(quality.get('min_contrast_normal', 4.5)) else label_color
        frame_text = shape.text_frame
        frame_text.clear()
        frame_text.word_wrap = True
        frame_text.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame_text.margin_left = frame_text.margin_right = Inches(.08)
        frame_text.margin_top = frame_text.margin_bottom = Inches(.04)
        blocks = [(node['label'], label_style, label_color, True)]
        if node.get('detail'):
            blocks.append((node['detail'], detail_style, detail_color, False))
        for index, (text, style, color, bold) in enumerate(blocks):
            paragraph = frame_text.paragraphs[0] if index == 0 else frame_text.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.text = text
            if index:
                paragraph.space_before = Pt(2)
            for run in paragraph.runs:
                run.font.name = style['font_family']
                run.font.size = Pt(style['font_size'])
                run.font.bold = bold
                run.font.color.rgb = RGBColor.from_string(color)
        old = previous.get(element_id, {})
        meta = dict(kind='shape', shape='rounded_rectangle', layer='content', overlap_group=group, diagram_id=spec['id'],
            diagram_role='node', node_id=node['id'], label=node['label'], text='\n'.join(text for text, *_ in blocks),
            text_blocks=[dict(text=text, font_size=style['font_size'], line_height=style.get('line_height', 1.2)) for text, style, *_ in blocks],
            resolved_style=dict(label_style, color=label_color, align='center'), fill=fill, fill_token=fill_token, rank=ranks[node['id']],
            order=order, diagram=frame)
        for key in ('detail', 'emphasis'):
            if node.get(key) is not None:
                meta[key] = node[key]
        if node.get('rank') is not None:
            meta['declared_rank'] = node['rank']
        for key in ('style_overrides', 'preset') + (('manual_geometry',) if keep_positions else ()):
            if old.get(key):
                meta[key] = old[key]
        if meta.get('style_overrides') or meta.get('preset'):
            native_style.apply_shape(shape, meta, state)
        renderer.set_metadata(shape, meta)
        created.append(shape)
        by_node[node['id']] = shape
    width_in, height_in = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    label_size = dict(style_for(state, 'caption'))
    occupied = [rect for rect in rects.values()] + [rect for _, rect in obstacles]
    for order, edge in enumerate(spec['edges']):
        info = routes[edge['id']]
        relation = diagram_presets.RELATIONS[edge['relation']]
        points = info['points']
        kind = MSO_CONNECTOR.STRAIGHT if len(points) == 2 else MSO_CONNECTOR.ELBOW
        connector = slide.shapes.add_connector(kind, Emu(0), Emu(0), Emu(1), Emu(1))
        element_id = spec['id'] + '.' + edge['id']
        connector.name = 'da:' + element_id
        _set_xfrm(connector, points, prs)
        native_style._set_line(connector.line, connector_color, preset['connector_width_pt'], relation['dash'])
        _arrows(connector, relation['arrow'])
        _bind(connector, by_node[edge['source']], info['start_site'], by_node[edge['target']], info['end_site'])
        old = previous.get(element_id, {})
        meta = dict(kind='connector', layer='content', overlap_group=group, diagram_id=spec['id'], diagram_role='edge',
            edge_id=edge['id'], source=edge['source'], target=edge['target'], relation=edge['relation'], arrow=relation['arrow'],
            source_element=spec['id'] + '.' + edge['source'], target_element=spec['id'] + '.' + edge['target'],
            route=[[round(px / width_in, 6), round(py / height_in, 6)] for px, py in points], fill=connector_color, order=order)
        if edge.get('label'):
            meta['label'] = edge['label']
        for key in ('style_overrides', 'preset'):
            if old.get(key):
                meta[key] = old[key]
        if meta.get('style_overrides'):
            native_style.apply_shape(connector, meta, state)
        renderer.set_metadata(connector, meta)
        created.append(connector)
        if edge.get('label'):
            placed = None
            for rect in label_rects(points, edge['label'], label_size['font_size']):
                inside = 0 <= rect[0] and 0 <= rect[1] and rect[0] + rect[2] <= width_in and rect[1] + rect[3] <= height_in
                if inside and not any(_overlap(rect, other) for other in occupied):
                    placed = rect
                    break
            if placed is None:
                raise ValueError(f'Edge label {edge["label"]!r} has no free space next to its connector; shorten it or enlarge the diagram gaps')
            occupied.append(placed)
            box = slide.shapes.add_textbox(Inches(placed[0]), Inches(placed[1]), Inches(placed[2]), Inches(placed[3]))
            box.name = 'da:' + element_id + '.label'
            style = dict(label_size, color=_readable(background, [palette['muted'], palette['text']], 4.5), align='center')
            HOST['apply_text'](box, edge['label'], style)
            renderer.set_metadata(box, dict(kind='text', layer='content', overlap_group=group, diagram_id=spec['id'], diagram_role='edge_label',
                edge_id=edge['id'], text=edge['label'], style='caption', resolved_style=style, align='center'))
            created.append(box)
    _reinsert(slide, created, insert_at)
    return created


def _overlap(a, b):
    return min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]) > 1e-6 and min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]) > 1e-6


def _inset(rect, amount):
    x, y, w, h = rect
    return (x + amount, y + amount, max(0.0, w - 2 * amount), max(0.0, h - 2 * amount))


def _merge(spec, payload):
    """diagram.update: full replacement of nodes/edges, or an explicit patch."""
    spec = dict(spec, nodes=[dict(node) for node in spec['nodes']], edges=[dict(edge) for edge in spec['edges']])
    unknown = set(payload) - SPEC_KEYS - PATCH_KEYS
    if unknown:
        raise ValueError('Unsupported diagram.update field(s): ' + ', '.join(sorted(unknown)))
    for key in ('direction', 'preset'):
        if key in payload:
            spec[key] = payload[key]
    if 'geometry' in payload or 'anchor' in payload:
        spec.pop('geometry', None); spec.pop('anchor', None)
        spec.update({key: payload[key] for key in ('geometry', 'anchor') if payload.get(key)})
    if 'nodes' in payload:
        spec['nodes'] = payload['nodes']
    if 'edges' in payload:
        spec['edges'] = payload['edges']
    removed = set(payload.get('remove_nodes') or [])
    if removed:
        unknown_nodes = removed - {node['id'] for node in spec['nodes']}
        if unknown_nodes:
            raise ValueError('Cannot remove unknown node(s): ' + ', '.join(sorted(unknown_nodes)))
    removed_edges = set(payload.get('remove_edges') or [])
    known_edges = {edge['id'] for edge in spec['edges']}
    if removed_edges - known_edges:
        raise ValueError('Cannot remove unknown edge(s): ' + ', '.join(sorted(removed_edges - known_edges)))
    spec['edges'] = [edge for edge in spec['edges'] if edge['id'] not in removed_edges]
    dangling = [edge['id'] for edge in spec['edges'] if edge['source'] in removed or edge['target'] in removed]
    if dangling:
        raise ValueError('Removing node(s) would leave edge(s) ' + ', '.join(dangling) + ' without an endpoint; remove or rebind those edges explicitly')
    spec['nodes'] = [node for node in spec['nodes'] if node['id'] not in removed]
    for update in payload.get('update_nodes') or []:
        if not isinstance(update, dict) or update.get('id') not in {node['id'] for node in spec['nodes']}:
            raise ValueError('update_nodes entries must name an existing node id')
        node = next(node for node in spec['nodes'] if node['id'] == update['id'])
        for key, value in update.items():
            if key not in NODE_KEYS:
                raise ValueError(f'Unsupported node field {key}')
            if value is None:
                node.pop(key, None)
            else:
                node[key] = value
    spec['nodes'].extend(payload.get('add_nodes') or [])
    spec['edges'].extend(payload.get('add_edges') or [])
    return spec


def apply(slide, op, payload, state, prs, anchors, profile):
    renderer = _renderer()
    if op == 'diagram.add':
        spec = normalize(payload)
        if diagram_shapes(slide, spec['id']):
            raise ValueError(f'Diagram {spec["id"]} already exists; use diagram.update')
        wanted = {spec['id'] + '.' + node['id'] for node in spec['nodes']} | {spec['id'] + '.' + edge['id'] for edge in spec['edges']}
        if wanted & {renderer.element_identity(shape) for shape in slide.shapes}:
            raise ValueError('Diagram element ids collide with existing elements')
        return build(slide, spec, state, prs, anchors, profile)
    identity = payload.get('id')
    current, metas = read_spec(slide, identity)
    spec = normalize(_merge(current, payload))
    rects = None
    # Manually placed nodes stay put when only relations change; structural changes re-layout.
    if any(meta.get('manual_geometry') for meta in metas.values()) and set(payload) <= {'id', 'edges', 'add_edges', 'remove_edges'}:
        rects = _node_rects(slide, identity)
    index = _remove(slide, identity)
    return build(slide, spec, state, prs, anchors, profile, previous=metas, rects=rects, insert_at=index)


def rebuild(slide, diagram_id, state, prs, anchors):
    """Profile change: re-apply tokens; manual node positions are kept."""
    spec, metas = read_spec(slide, diagram_id)
    rects = None
    if any(meta.get('manual_geometry') for meta in metas.values()):
        rects = _node_rects(slide, diagram_id)
    index = _remove(slide, diagram_id)
    return build(slide, normalize(spec), state, prs, anchors, state['profile'], previous=metas, rects=rects, insert_at=index)


def _node_rects(slide, diagram_id):
    renderer = _renderer()
    return {renderer.metadata(shape)['node_id']: _rect_in(shape) for shape in diagram_shapes(slide, diagram_id)
            if renderer.metadata(shape).get('diagram_role') == 'node'}


def reroute(slide, diagram_id, state, prs):
    """After a node is adjusted, keep every node where it is and route the edges again."""
    renderer = _renderer()
    rects = _node_rects(slide, diagram_id)
    spec, metas = read_spec(slide, diagram_id)
    for shape in diagram_shapes(slide, diagram_id):
        meta = renderer.metadata(shape)
        if meta.get('diagram_role') == 'node' and not meta.get('manual_geometry'):
            meta['manual_geometry'] = True
            renderer.set_metadata(shape, meta)
            metas[renderer.element_identity(shape)] = meta
    for a, first in rects.items():
        for b, second in rects.items():
            if a < b and _overlap(first, second):
                raise ValueError(f'Nodes {a} and {b} would overlap; adjust less')
    anchors = HOST['anchors'](state, HOST['slide_id'](slide))
    index = _remove(slide, diagram_id)
    return build(slide, normalize(spec), state, prs, anchors, state['profile'], previous=metas, rects=rects, insert_at=index)


def style_element(slide, shape, meta, element, op, state, prs, anchors):
    """Style overrides on a node or edge persist across re-layout of the diagram."""
    renderer, native_style = HOST['renderer'], HOST['native_style']
    kind = 'connector' if meta.get('diagram_role') == 'edge' else 'shape'
    if meta.get('diagram_role') == 'edge_label':
        kind = 'text'
    if op == 'element.set_style':
        if set(element) - {'id', 'style'}:
            raise ValueError('element.set_style accepts id and style only')
        overrides = dict(meta.get('style_overrides') or {})
        additions = {key: value for key, value in (element.get('style') or {}).items() if value is not None}
        for key in [key for key, value in (element.get('style') or {}).items() if value is None]:
            overrides.pop(key, None)
        if additions:
            overrides.update(native_style.validate_style(kind, additions, state))
        if overrides:
            meta['style_overrides'] = overrides
        else:
            meta.pop('style_overrides', None)
    else:
        if element.get('preset') is None:
            meta.pop('preset', None)
        else:
            native_style.preset_style(element['preset'], kind)
            meta['preset'] = element['preset']
    renderer.set_metadata(shape, meta)
    if kind == 'text':
        style = native_style.text_style(meta, state)
        HOST['apply_text'](shape, meta['text'], style)
        meta['resolved_style'] = style
        renderer.set_metadata(shape, meta)
        return
    rebuild(slide, meta['diagram_id'], state, prs, anchors)
