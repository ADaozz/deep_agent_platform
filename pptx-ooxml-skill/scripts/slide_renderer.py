"""Structured editable-element preview. Complex template objects stay opaque."""
import base64
import json
import math
from pptx.enum.shapes import MSO_SHAPE_TYPE

DEFAULT_STYLE=dict(font_family='Noto Sans CJK SC',font_size=20,color='17324D')


def slide_identity(slide):
    name=slide._element.cSld.get('name','')
    return name[9:] if name.startswith('da-slide:') else f'template-{slide.slide_id}'


def element_identity(shape):
    return shape.name[3:] if shape.name.startswith('da:') else f'template-{shape.shape_id}'


def metadata(shape):
    nodes=shape._element.xpath('.//p:cNvPr')
    if nodes:
        try:return json.loads(nodes[0].get('descr','{}'))
        except (ValueError,TypeError):pass
    return {}


def set_metadata(shape,data):
    shape._element.xpath('.//p:cNvPr')[0].set('descr',json.dumps(data,ensure_ascii=False))


def profile_view(state):
    """Design tokens for the quick preview; never page content."""
    profile=state.get('profile')
    if not profile:return None
    return dict(id=profile.get('id'),version=profile.get('version'),name=profile.get('name'),label=profile.get('label'),
        colors=profile.get('colors'),fonts=profile.get('fonts'),type_scale=profile.get('type_scale'),
        spacing=profile.get('spacing'),density=profile.get('density'),tables=profile.get('tables'),
        charts=profile.get('charts'),quality=profile.get('quality'))


MANIFEST_KEYS=('text','shape','rows','chart','fill','line','alt','component','component_instance','component_slot','align','vector',
    'rotation','preset','style_overrides','manual_geometry','text_color','diagram_id','diagram_role','node_id','edge_id',
    'source','target','relation','label','detail','arrow','route','icon','source_element','target_element','text_blocks',
    'emphasis','rank','edge_id','chart_style')
A_NS='{http://schemas.openxmlformats.org/drawingml/2006/main}'
P_NS='{http://schemas.openxmlformats.org/presentationml/2006/main}'


def connector_view(shape,by_shape_id,width,height):
    """Real endpoints, bound shapes and arrowheads of a p:cxnSp connector."""
    element=shape._element
    if element.tag!=P_NS+'cxnSp':return None
    props=element.find('.//'+P_NS+'cNvCxnSpPr')
    def bound(tag):
        node=props.find(A_NS+tag) if props is not None else None
        if node is None:return None
        target=by_shape_id.get(node.get('id'))
        return dict(element_id=target,site=int(node.get('idx','0'))) if target else dict(element_id=None,site=int(node.get('idx','0')),missing=True)
    line=element.find('.//'+A_NS+'ln')
    def arrow(tag):
        node=line.find(A_NS+tag) if line is not None else None
        return node is not None and node.get('type','none')!='none'
    geometry=element.find('.//'+A_NS+'prstGeom')
    xfrm=element.find('.//'+A_NS+'xfrm')
    x,y,w,h=shape.left,shape.top,shape.width,shape.height
    angle=math.radians(int(xfrm.get('rot','0'))/60000) if xfrm is not None else 0.0
    points=[]
    for sx,sy in ((-w/2,-h/2),(w/2,h/2)):
        if xfrm is not None and xfrm.get('flipH') in ('1','true'):sx=-sx
        if xfrm is not None and xfrm.get('flipV') in ('1','true'):sy=-sy
        px=sx*math.cos(angle)-sy*math.sin(angle)+x+w/2;py=sx*math.sin(angle)+sy*math.cos(angle)+y+h/2
        points.append((round(px/width,6),round(py/height,6)))
    begin,end=points
    return dict(begin=begin,end=end,start=bound('stCxn'),finish=bound('endCxn'),
        arrow_start=arrow('headEnd'),arrow_end=arrow('tailEnd'),geometry=geometry.get('prst') if geometry is not None else None)


def manifest(prs,state):
    result=[]; width,height=prs.slide_width,prs.slide_height
    for slide in prs.slides:
        sid=slide_identity(slide); elements=[]; complex_slide=bool(state.get('template_complex',{}).get(sid))
        by_shape_id={str(shape.shape_id):element_identity(shape) for shape in slide.shapes}
        for shape in slide.shapes:
            meta=metadata(shape); generated=bool(meta.get('kind'))
            g={k:round(float(v),6) for k,v in zip(('x','y','w','h'),(shape.left/width,shape.top/height,shape.width/width,shape.height/height))}
            e=dict(id=element_identity(shape),geometry=g,layer=meta.get('layer','content'),anchor=meta.get('anchor'),
                overlap_group=meta.get('overlap_group'),style=meta.get('resolved_style',DEFAULT_STYLE),
                style_token=meta.get('style'),source='generated' if generated else 'template')
            connection=connector_view(shape,by_shape_id,width,height)
            if connection:e['connector']=connection
            if generated:
                e['type']=meta['kind']
                for key in MANIFEST_KEYS:
                    if key in meta:e[key]=meta[key]
                if e['type'] in ('image','component'):e['image_type']=shape.image.content_type
                if e['type']=='chart' and (e['chart'].get('type')=='pie' or any(v<0 for series in e['chart']['series'] for v in series['values'])):complex_slide=True
            elif shape.has_text_frame:
                # An untouched template's rich text/theme can be more complicated
                # than the quick renderer. Use the actual rendered slide instead.
                e.update(type='complex',text=shape.text); complex_slide=True
            else:
                e.update(type='complex'); complex_slide=True
            elements.append(e)
        variant=state.get('layout_variants',{}).get(sid)
        result.append(dict(slide_id=sid,elements=elements,complex=complex_slide,
            background=state.get('backgrounds',{}).get(sid,'FFFFFF'),layout=state.get('layouts',{}).get(sid,'title_body'),
            layout_variant=variant if variant not in (None,'0') else None,layout_modules=state.get('layout_modules',{}).get(sid),
            profile_id=state.get('profile_id'),
            review=state.get('reviews',{}).get(sid,dict(status='writing',visual_review='pending')),
            formal=state.get('renders',{}).get(sid)))
    return dict(width=width,height=height,aspect_ratio=width/height,profile=profile_view(state),
        conflicts=list(state.get('design_conflicts') or []),slides=result)
