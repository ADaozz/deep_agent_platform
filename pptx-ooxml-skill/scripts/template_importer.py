"""Read a PPTX template into a manifest the agent can correct and reuse.

The importer does not download templates and does not copy an unlicensed file
into this skill. A manifest records slide roles, fill slots, theme tokens and
precheck warnings. ``index_loaded`` writes stable ``da-slide:`` / ``da:`` names
so later operations do not depend on PowerPoint's numeric shape ids.

The template library is independent of any working PPTX: pages are selected
from ``templates/`` (or another root) and copied in with ``slide.import``.
"""
import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
ID = re.compile(r'^[\w.-]{1,80}$')
SKILL_ROOT = Path(__file__).resolve().parents[1]


def file_sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def list_templates(root=None):
    base = Path(root) if root else SKILL_ROOT / 'templates'
    found = []
    if not base.is_dir():
        return found
    for manifest_path in sorted(base.glob('*/manifest.json')):
        try:
            manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        except (OSError, json.JSONDecodeError):
            continue
        package = manifest_path.parent / 'template.pptx'
        recorded = manifest.get('source_sha256')
        actual = file_sha256(package) if package.is_file() else None
        found.append(dict(
            template_id=manifest.get('template_id') or manifest_path.parent.name,
            path=str(package),
            manifest=str(manifest_path),
            license=manifest.get('license'),
            source_sha256=actual or recorded,
            source_sha256_recorded=recorded,
            slide_size=manifest.get('slide_size'),
            tokens=manifest.get('tokens') or {},
            slides=[_catalog_slide(slide, manifest.get('slide_size')) for slide in manifest.get('slides', [])],
        ))
    return found


def _catalog_slide(slide, default_size=None):
    slide_id = slide.get('slide_id')
    return dict(
        slide_id=slide_id,
        role=slide.get('role'),
        preview=slide.get('preview') or (f'previews/{slide_id}.png' if slide_id else None),
        slots=dict(slide.get('slots') or {}),
        native_objects=dict(slide.get('native_objects') or {}),
        slide_size=slide.get('slide_size') or default_size,
        preserve=list(slide.get('preserve') or []),
    )


def load_template(template_id, root=None):
    """Return ``(pptx_path, catalog_entry, source_sha256)`` for a library template."""
    if not ID.fullmatch(template_id or ''):
        raise ValueError('Invalid template_id')
    matches = [item for item in list_templates(root) if item['template_id'] == template_id]
    if not matches:
        raise ValueError('Unknown template_id ' + template_id)
    entry = matches[0]
    package = Path(entry['path'])
    if not package.is_file():
        raise ValueError('Template file not found: ' + str(package))
    actual = file_sha256(package)
    recorded = entry.get('source_sha256_recorded')
    if recorded and recorded != actual:
        raise ValueError('Template source_sha256 does not match the file; refusing to use a replaced template')
    return package, entry, actual


def import_template(path, template_id=None, output=None, license_note=''):
    """Analyze ``path`` and optionally write a manifest. The PPTX itself is not copied."""
    source = Path(path)
    if not source.is_file():
        raise ValueError('Template file not found')
    prs = Presentation(str(source))
    manifest = analyze(prs, template_id or source.stem, source=str(source), license_note=license_note)
    if output:
        dest = Path(output)
        if dest.suffix.lower() != '.json':
            dest.mkdir(parents=True, exist_ok=True)
            dest = dest / 'manifest.json'
        _refuse_unlicensed_vendor(dest, manifest)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        manifest = dict(manifest, manifest_path=str(dest))
    return manifest


def index_loaded(prs, template_path):
    """Assign stable ids from an optional sidecar manifest and return the manifest."""
    source = Path(template_path)
    manifest = analyze(prs, source.stem, source=str(source))
    sidecar = _sidecar(source)
    if sidecar:
        manifest = merge_manifest(manifest, sidecar)
    apply_stable_ids(prs, manifest)
    manifest['slots_by_slide'] = {
        slide['slide_id']: dict(slide.get('slots') or {}) for slide in manifest['slides']
    }
    return manifest


def analyze(prs, template_id, source='', license_note=''):
    if not ID.fullmatch(template_id or ''):
        template_id = 'template'
    slides = []
    for index, slide in enumerate(prs.slides):
        slides.append(_describe_slide(slide, index, prs))
    tokens = _tokens(prs)
    fonts = _fonts(prs)
    installed = _installed_families()
    missing = sorted(face for face in fonts if face.casefold() not in installed and not face.startswith('+'))
    warnings = []
    if missing:
        warnings.append('missing fonts: ' + ', '.join(missing))
    if _has_smartart(prs):
        warnings.append('SmartArt is preserved but not a fill slot')
    if shutil.which('soffice') is None and shutil.which('libreoffice') is None:
        warnings.append('LibreOffice is not available to verify render compatibility')
    size = dict(width=int(prs.slide_width), height=int(prs.slide_height))
    for slide in slides:
        slide['preview'] = f'previews/{slide["slide_id"]}.png'
        slide['slide_size'] = size
        slide['native_objects'] = _native_objects(slide['elements'])
    return dict(
        schema='deep-agent-template/v1',
        kind='pptx',
        template_id=template_id,
        source=source,
        source_sha256=file_sha256(source) if source and Path(source).is_file() else None,
        slide_size=size,
        license=_license(source, license_note),
        tokens=tokens,
        precheck=dict(
            slides=len(slides),
            charts=sum(1 for slide in slides for item in slide['elements'] if item['kind'] == 'chart'),
            images=sum(1 for slide in slides for item in slide['elements'] if item['kind'] == 'image'),
            tables=sum(1 for slide in slides for item in slide['elements'] if item['kind'] == 'table'),
            fonts=fonts,
            missing_fonts=missing,
            warnings=warnings,
        ),
        slides=slides,
    )


def merge_manifest(draft, sidecar):
    """Sidecar roles and slot indexes override the heuristic draft. Slots stay editable."""
    if not isinstance(sidecar, dict):
        raise ValueError('Template manifest must be a JSON object')
    merged = dict(draft)
    if sidecar.get('template_id'):
        merged['template_id'] = sidecar['template_id']
    if sidecar.get('license'):
        merged['license'] = sidecar['license']
    overrides = list(sidecar.get('slides') or [])
    slides = []
    for index, slide in enumerate(draft['slides']):
        item = dict(slide)
        extra = overrides[index] if index < len(overrides) else None
        if extra:
            if extra.get('role'):
                item['role'] = extra['role']
            if extra.get('slide_id') and ID.fullmatch(extra['slide_id']):
                item['slide_id'] = extra['slide_id']
            if isinstance(extra.get('slots'), dict):
                item['slots'], item['elements'] = _slots_from_override(item, extra['slots'])
        slides.append(item)
    merged['slides'] = slides
    return merged


def apply_stable_ids(prs, manifest):
    for slide, spec in zip(prs.slides, manifest['slides']):
        slide_id = spec['slide_id']
        if not ID.fullmatch(slide_id):
            raise ValueError('Invalid template slide id ' + str(slide_id))
        slide._element.cSld.set('name', 'da-slide:' + slide_id)
        used = set()
        for element in spec['elements']:
            element_id = element['id']
            if not ID.fullmatch(element_id) or element_id in used:
                raise ValueError('Invalid or duplicate template element id ' + str(element_id))
            used.add(element_id)
            prs_slide_shapes = list(slide.shapes)
            if not 0 <= element['index'] < len(prs_slide_shapes):
                raise ValueError('Template manifest shape index is outside the slide')
            prs_slide_shapes[element['index']].name = 'da:' + element_id


def match_library(template_id=None, role=None, root=None):
    """Search the template library. A working PPTX is not required."""
    templates = list_templates(root)
    if template_id:
        templates = [item for item in templates if item['template_id'] == template_id]
        if not templates:
            raise ValueError('Unknown template_id ' + str(template_id))
    payload = []
    for item in templates:
        slides = list(item.get('slides') or [])
        if role:
            slides = [slide for slide in slides if slide.get('role') == role]
        if not slides and role:
            continue
        payload.append(dict(
            template_id=item['template_id'],
            source_sha256=item.get('source_sha256'),
            slide_size=item.get('slide_size'),
            slides=slides,
        ))
    if template_id:
        chosen = payload[0] if payload else dict(
            template_id=template_id, source_sha256=templates[0].get('source_sha256'),
            slide_size=templates[0].get('slide_size'), slides=[])
        return dict(chosen, role=role)
    return dict(role=role, templates=payload)


def match(status, role, root=None):
    """Library search. ``status`` is ignored; kept so older callers still work."""
    del status
    return match_library(role=role, root=root)


def consistency_report(before, after):
    """Compare two PPTX packages after a template edit. Positions are EMU tuples."""
    left, right = _geometry(before), _geometry(after)
    text_changes, geometry_changes, chart_changes = [], [], []
    for index, (source, dest) in enumerate(zip(left, right)):
        by_id = {item['id']: item for item in dest}
        for item in source:
            other = by_id.get(item['id'])
            if other is None:
                continue
            if item['text'] != other['text']:
                text_changes.append(dict(slide=index, id=item['id']))
            if item['box'] != other['box']:
                geometry_changes.append(dict(slide=index, id=item['id']))
            if item['chart'] != other['chart']:
                chart_changes.append(dict(slide=index, id=item['id']))
    return dict(
        slides_before=len(left), slides_after=len(right),
        pictures_before=sum(item['picture'] for slide in left for item in slide),
        pictures_after=sum(item['picture'] for slide in right for item in slide),
        charts_before=sum(item['chart'] is not None for slide in left for item in slide),
        charts_after=sum(item['chart'] is not None for slide in right for item in slide),
        text_changes=text_changes, geometry_changes=geometry_changes, chart_changes=chart_changes,
    )


def _describe_slide(slide, index, prs):
    name = slide._element.cSld.get('name') or ''
    slide_id = name[9:] if name.startswith('da-slide:') else f'tpl-{index + 1:03d}'
    elements = []
    for shape_index, shape in enumerate(slide.shapes):
        kind = _kind(shape)
        props = shape._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        original_name = shape.name
        original_descr = props.get('descr') or '' if props is not None else ''
        label = f'{original_name} {original_descr}'.lower()
        semantic = 'logo' if 'logo' in label else ('brand' if 'brand' in label else None)
        elements.append(dict(
            id=_element_id(shape, shape_index),
            index=shape_index,
            kind=kind,
            name=shape.name,
            original_name=original_name,
            original_descr=original_descr,
            semantic_role=semantic,
            preserve=semantic in ('logo', 'brand'),
            text=(shape.text_frame.text if shape.has_text_frame and not shape.has_table else '')[:200],
            x=shape.left / prs.slide_width,
            y=shape.top / prs.slide_height,
            w=shape.width / prs.slide_width,
            h=shape.height / prs.slide_height,
        ))
    slots = _slots(elements)
    return dict(
        slide_id=slide_id,
        index=index,
        role=_role(index, elements, slots),
        layout=getattr(slide.slide_layout, 'name', ''),
        slots=slots,
        elements=elements,
        preserve=[element['id'] for element in elements if element.get('preserve')],
    )


def _native_objects(elements):
    counts = dict(chart=0, table=0, image=0, shape=0, text=0)
    for element in elements:
        kind = element.get('kind')
        if kind in counts:
            counts[kind] += 1
        else:
            counts['shape'] += 1
    return {kind: count for kind, count in counts.items() if count}


def _kind(shape):
    if shape.has_chart:
        return 'chart'
    if shape.has_table:
        return 'table'
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return 'image'
    if shape.has_text_frame and shape.text_frame.text.strip():
        return 'text'
    return 'shape'


def _element_id(shape, index):
    if shape.name.startswith('da:') and ID.fullmatch(shape.name[3:]):
        return shape.name[3:]
    return f'shape-{index + 1:03d}'


def _slots(elements):
    texts = [element for element in elements if element['kind'] == 'text']
    texts.sort(key=lambda item: (item['y'], item['x']))
    slots = {}
    title = next((item for item in texts if item['y'] < 0.35 and item['w'] > 0.25), None)
    if title:
        slots['title'] = title['id']
    columns = [item for item in texts if item is not title]
    left = [item for item in columns if item['x'] + item['w'] / 2 < 0.48]
    right = [item for item in columns if item['x'] + item['w'] / 2 > 0.52]
    if left and right:
        _name_column(slots, left, 'left')
        _name_column(slots, right, 'right')
    else:
        rest = [item for item in texts if item is not title]
        if rest and 'subtitle' not in slots and rest[0]['y'] < 0.45:
            slots['subtitle'] = rest[0]['id']
            rest = rest[1:]
        for number, item in enumerate(rest, 1):
            slots['body' if number == 1 else f'body_{number}'] = item['id']
    for element in elements:
        if element['kind'] == 'chart':
            slots.setdefault('chart', element['id'])
        elif element['kind'] == 'table':
            slots.setdefault('table', element['id'])
        elif element['kind'] == 'image' and 'logo' not in element['id'].lower():
            slots.setdefault('picture', element['id'])
    return slots


def _name_column(slots, items, side):
    items = sorted(items, key=lambda item: item['y'])
    if not items:
        return
    slots[f'{side}_title'] = items[0]['id']
    if len(items) > 1:
        slots[f'{side}_body'] = items[1]['id']


def _role(index, elements, slots):
    kinds = {element['kind'] for element in elements}
    if 'chart' in kinds:
        return 'data'
    if 'table' in kinds:
        return 'table'
    if 'left_body' in slots or 'left_title' in slots and 'right_title' in slots:
        return 'comparison'
    text_count = sum(element['kind'] == 'text' for element in elements)
    if index == 0:
        return 'cover'
    if text_count <= 1:
        return 'section'
    return 'content'


def _slots_from_override(slide, slots):
    elements = [dict(element) for element in slide['elements']]
    by_index = {element['index']: element for element in elements}
    by_id = {element['id']: element for element in elements}
    resolved = {}
    for name, target in slots.items():
        if not ID.fullmatch(name):
            raise ValueError('Invalid slot name ' + str(name))
        if isinstance(target, int):
            element = by_index.get(target)
            if element is None:
                raise ValueError('Slot shape index is outside the slide')
            element['id'] = name
            resolved[name] = name
        elif isinstance(target, str):
            element = by_id.get(target) or by_id.get(name)
            if element is None:
                raise ValueError('Slot does not match a shape: ' + target)
            resolved[name] = element['id']
        else:
            raise ValueError('Slot targets must be a shape index or element id')
    return resolved, elements


def _tokens(prs):
    tokens = {}
    for rel in prs.part.package.iter_rels():
        if rel.is_external or rel.reltype != RT.THEME:
            continue
        scheme = _part_xml(rel.target_part).find('.//' + A + 'clrScheme')
        if scheme is None:
            continue
        for child in scheme:
            srgb = child.find(A + 'srgbClr')
            sysclr = child.find(A + 'sysClr')
            value = srgb.get('val') if srgb is not None else (sysclr.get('lastClr') if sysclr is not None else None)
            if value:
                tokens[etree.QName(child).localname] = value.upper()
        break
    return tokens


def _part_xml(part):
    element = getattr(part, '_element', None)
    return element if element is not None else etree.fromstring(part.blob)


def _fonts(prs):
    faces = []
    for rel in prs.part.package.iter_rels():
        if rel.is_external or rel.reltype != RT.THEME:
            continue
        for node in _part_xml(rel.target_part).iter(A + 'latin', A + 'ea', A + 'cs'):
            face = node.get('typeface')
            if face and face not in faces:
                faces.append(face)
        break
    return faces


def _installed_families():
    if shutil.which('fc-list') is None:
        return set()
    try:
        result = subprocess.run(['fc-list', ':family'], capture_output=True, text=True, timeout=8, check=False)
    except (OSError, subprocess.TimeoutExpired):
        return set()
    families = set()
    for line in result.stdout.splitlines():
        for part in line.split(','):
            families.add(part.strip().casefold())
    return families


def _has_smartart(prs):
    return any(not rel.is_external and str(rel.reltype).endswith('/diagramData') for rel in prs.part.package.iter_rels())


def _license(source, note):
    sidecar = _sidecar(Path(source)) if source else None
    if sidecar and isinstance(sidecar.get('license'), dict):
        recorded = dict(sidecar['license'])
        recorded.setdefault('redistributable', False)
        return recorded
    return dict(
        status='recorded' if note else 'external',
        redistributable=False,
        note=note or 'External file. Do not copy it into the skill unless redistribution is explicitly allowed.',
        source=source,
    )


def _sidecar(path):
    manifest = Path(path).parent / 'manifest.json'
    if not manifest.is_file():
        return None
    try:
        return json.loads(manifest.read_text(encoding='utf-8'))
    except json.JSONDecodeError as exc:
        raise ValueError('Template manifest.json is not valid JSON') from exc


def _refuse_unlicensed_vendor(dest, manifest):
    root = (SKILL_ROOT / 'templates').resolve()
    target = dest.resolve()
    if target == root or root in target.parents:
        if not (manifest.get('license') or {}).get('redistributable'):
            raise ValueError('Refusing to write an unlicensed template manifest into the skill templates directory')


def _geometry(data):
    presentation = Presentation(__import__('io').BytesIO(data))
    slides = []
    for slide in presentation.slides:
        items = []
        for shape in slide.shapes:
            chart = None
            if shape.has_chart:
                chart = [tuple(float(value) for value in series.values) for series in shape.chart.series]
            text = ''
            if shape.has_text_frame and not shape.has_table:
                text = shape.text_frame.text
            items.append(dict(
                id=shape.name, text=text, box=(shape.left, shape.top, shape.width, shape.height),
                chart=chart, picture=shape.shape_type == MSO_SHAPE_TYPE.PICTURE,
            ))
        slides.append(items)
    return slides
