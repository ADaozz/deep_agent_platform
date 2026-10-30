#!/usr/bin/env python3
"""Element-atomic PPTX engine, using the platform's shared artifact transaction store.

Design behaviour lives in ``design/``: profiles provide tokens, layouts provide
semantic slots, components provide reviewed SVG instances. This module only
executes validated atomic operations; it never plans pages or draws new SVG.
"""
import argparse
import copy
import hashlib
import importlib.util
import io
import json
import math
from pathlib import Path
import sys
import tempfile
import time

from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.package import Part
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

SKILL_ROOT=Path(__file__).resolve().parents[1]
if str(SKILL_ROOT) not in sys.path:sys.path.insert(0,str(SKILL_ROOT))
from design import design_quality, layout_registry, layout_resolver, profile_registry


def load(name):
    path=Path(__file__).with_name(name+'.py')
    if name=='artifact_runtime' and not path.exists():path=Path(__file__).resolve().parents[2]/'platform_app'/'artifact_runtime.py'
    spec=importlib.util.spec_from_file_location(name,path); mod=importlib.util.module_from_spec(spec); spec.loader.exec_module(mod); return mod

shared=load('artifact_runtime'); renderer=load('slide_renderer'); validator=load('slide_validator'); component_renderer=load('component_renderer')

DEFAULT_PROFILE='corporate_blue'
SVG_BLIP_URI='{96DAC541-7B7A-43D3-8B79-37D633B846F1}'
SVG_NAMESPACE='http://schemas.microsoft.com/office/drawing/2016/SVG/main'
# Legacy constants stay importable for older callers; new code reads profiles.
THEMES={
 'corporate_blue':dict(background='FFFFFF',text='17324D',accent='246BCE',muted='536B82',card='EDF4FC'),
 'slate_dark':dict(background='14263D',text='FFFFFF',accent='66B5FF',muted='BCD0E4',card='203B58')}
STYLES={'slide_title':dict(font_size=30,bold=True,color_token='text'),'cover_title':dict(font_size=40,bold=True,color_token='text'),
 'body':dict(font_size=20,bold=False,color_token='text'),'caption':dict(font_size=12,bold=False,color_token='muted'),'card_title':dict(font_size=24,bold=True,color_token='accent')}

_PROFILES=_LAYOUTS=_COMPONENTS=None


def profiles():
    global _PROFILES
    if _PROFILES is None:_PROFILES=profile_registry.load_registry()
    return _PROFILES


def layouts():
    global _LAYOUTS
    if _LAYOUTS is None:_LAYOUTS=layout_registry.load_layouts()
    return _LAYOUTS


def registry():
    global _COMPONENTS
    if _COMPONENTS is None:_COMPONENTS=component_renderer.load()
    return _COMPONENTS


def identifier(value):return shared.identifier(value)


def component_validator_get(component_id):return component_renderer.components.get(registry(),component_id)


def palette_of(profile):
    return {key:value for key,value in profile['colors'].items() if key!='chart_series'}


def apply_profile(state,profile_id):
    profile=profiles().get(profile_id)
    state.update(profile_id=profile_id,profile=profile,theme_name=profile_id,theme=palette_of(profile))
    conflicts=[]
    accent=state.get('template_accent')
    if accent:
        background=state['theme']['background'];card=state['theme']['card']
        ratios=[design_quality.contrast(accent,background),design_quality.contrast(accent,card)]
        if min(ratios)>=3.0:state['theme']['accent']=accent
        else:conflicts.append(dict(kind='template_accent_contrast',accent=accent,profile=profile_id,
            contrast=min(ratios),kept_accent=state['theme']['accent'],
            message='Inherited template accent '+accent+' is unreadable on profile '+profile_id+'; the profile accent is used instead'))
    state['design_conflicts']=conflicts
    state['font_family']=state.get('template_font') or profile['fonts']['body']
    return profile


def profile_of(state):
    if 'profile' not in state:
        apply_profile(state,state.get('profile_id') or state.get('theme_name') or DEFAULT_PROFILE)
    return state['profile']


def style_for(state,token):
    profile=profile_of(state)
    if token not in profile['type_scale']:
        raise ValueError(f'Unknown style token {token!r}; available: '+', '.join(profile['type_scale']))
    item=profile['type_scale'][token]
    return dict(style_token=token,font_family=state.get('template_font') or profile['fonts'][item['font']],
        font_size=item['size'],bold=item['bold'],color=state['theme'].get(item['color_token'],profile['colors'][item['color_token']]),
        line_height=item['line_height'],letter_spacing=item.get('letter_spacing'))


def slot_views(state,layout_id,modules):
    key,defs=layouts().variant_for(layout_id,modules)
    if key is None:raise ValueError(f'Layout {layout_id} has no variant for {modules} modules; available: '+', '.join(map(str,layouts().get(layout_id)['module_counts'])))
    views={name:layout_resolver.slot_view(slot,profile_of(state),lambda token:style_for(state,token)) for name,slot in defs.items()}
    return key,views


def set_slide_layout(state,sid,layout_id,modules):
    if not profiles().allows_layout(state['profile_id'],layout_id):raise ValueError(f'Profile {state["profile_id"]} does not enable layout {layout_id}')
    key,views=slot_views(state,layout_id,modules)
    state.setdefault('layouts',{})[sid]=layout_id
    state.setdefault('layout_variants',{})[sid]=key
    state.setdefault('layout_modules',{})[sid]=modules
    state.setdefault('layout_slots',{})[sid]={name:view['rect'] for name,view in views.items()}
    return views


def check_anchor_content(state,sid,anchor,content):
    """Reject content that the declared slot cannot hold, before mutating."""
    layout_id=state.get('layouts',{}).get(sid)
    if not layout_id or not anchor:return
    if anchor in state.get('slots',{}).get(sid,{}):return
    modules=state.get('layout_modules',{}).get(sid,0)
    try:key,definitions=layouts().variant_for(layout_id,modules)
    except ValueError:return
    if not definitions or anchor not in definitions:return
    problem=layouts().capacity_error(layout_id,modules,anchor,content)
    if problem:raise ValueError(problem)


def normalize_state(state):
    """Upgrade a pre-design-system state without rewriting its package."""
    if 'profile' not in state:
        apply_profile(state,state.get('profile_id') or state.get('theme_name') or DEFAULT_PROFILE)
        # Reviews produced before the design system never ran the new checks.
        state['reviews']={};state['renders']={};state['repair_attempts']={}
    for key in ('slots','layouts','layout_variants','layout_modules','layout_slots'):
        state.setdefault(key,{})
    for sid,layout_id in list(state['layouts'].items()):
        if state['layout_slots'].get(sid):continue
        try:
            modules=state['layout_modules'].get(sid,layouts().get(layout_id)['default_modules'])
            key,views=slot_views(state,layout_id,modules)
        except ValueError:continue
        state['layout_variants'][sid]=key;state['layout_modules'][sid]=modules
        state['layout_slots'][sid]={name:view['rect'] for name,view in views.items()}


def anchors_for(state,sid):
    merged=dict(state.get('layout_slots',{}).get(sid,{}))
    merged.update(state.get('slots',{}).get(sid,{}))
    return merged


def resolve_style(state,element,views=None):
    token=element.get('style') or element.get('style_token')
    if not token and views:
        anchor=element.get('anchor');slot=views.get(anchor) if anchor else None
        token=slot['style_token'] if slot else None
    return style_for(state,token or 'body')


def geometry(element,anchors):
    if element.get('anchor'):
        if element['anchor'] not in anchors:raise ValueError('Unknown layout anchor')
        g=dict(zip(('x','y','w','h'),anchors[element['anchor']]))
    else:g=element.get('geometry',{})
    if set(g)!=set(('x','y','w','h')) or any(type(v) not in (int,float) or not math.isfinite(v) for v in g.values()):raise ValueError('Normalized finite geometry required')
    if min(g['x'],g['y'])<0 or min(g['w'],g['h'])<=0 or g['x']+g['w']>1.00001 or g['y']+g['h']>1.00001:raise ValueError('Element outside slide canvas')
    return g


def apply_text(shape,text,style):
    if not isinstance(text,str) or len(text)>1200:raise ValueError('One bounded text element, maximum 1200 characters')
    frame=shape.text_frame; frame.clear(); frame.word_wrap=True
    frame.margin_left=frame.margin_right=Pt(3); frame.margin_top=frame.margin_bottom=Pt(2)
    align={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(style.get('align'))
    for i,line in enumerate(text.split('\n')):
        p=frame.paragraphs[0] if i==0 else frame.add_paragraph(); p.text=line; p.space_after=Pt(4)
        p.line_spacing=max(.9,min(2.0,round(float(style.get('line_height',1.2))/1.2,3)))
        if align is not None:p.alignment=align
        for run in p.runs:
            run.font.name=style['font_family']; run.font.size=Pt(style['font_size']); run.font.bold=style['bold']; run.font.color.rgb=RGBColor.from_string(style['color'])
            rpr=run._r.get_or_add_rPr(); east=OxmlElement('a:ea'); east.set('typeface',style['font_family']); rpr.append(east)
            if style.get('letter_spacing'):rpr.set('spc',str(int(round(float(style['letter_spacing'])*100))))


def chart_data(chart):
    categories=chart['categories'];series=chart['series']
    if not 1<=len(categories)<=12 or any(not isinstance(label,str) or len(label)>80 for label in categories) or not 1<=len(series)<=4:
        raise ValueError('Bounded chart categories/series required')
    values=CategoryChartData();values.categories=categories
    for item in series:
        if not isinstance(item.get('name'),str) or len(item['name'])>80 or len(item['values'])!=len(categories) or any(type(n)not in (int,float) or not math.isfinite(n) for n in item['values']):
            raise ValueError('Bounded finite chart series required')
        values.add_series(item['name'],item['values'])
    return values


def positions(g,prs):return (round(g['x']*prs.slide_width),round(g['y']*prs.slide_height),round(g['w']*prs.slide_width),round(g['h']*prs.slide_height))


def set_geometry(shape,g,prs):shape.left,shape.top,shape.width,shape.height=positions(g,prs)


def check_color(color):
    if not isinstance(color,str) or len(color)!=6 or any(c not in '0123456789abcdefABCDEF' for c in color):raise ValueError('Invalid RGB color')


def serialize(prs):
    out=io.BytesIO(); prs.save(out); data=out.getvalue(); validator.validate_package(data); return data


def find_slide(prs,sid):
    values=[s for s in prs.slides if renderer.slide_identity(s)==sid]
    if len(values)!=1:raise ValueError('Unknown or ambiguous stable slide ID')
    return values[0]


def find_element(slide,eid):
    values=[s for s in slide.shapes if renderer.element_identity(s)==eid]
    if len(values)!=1:raise ValueError('Unknown or ambiguous element ID')
    return values[0]


def component_children(slide,instance_id):
    return [shape for shape in slide.shapes if renderer.metadata(shape).get('component_instance')==instance_id]


def svg_rel_id(shape):
    """Relationship id of the vector source attached to a picture, if any."""
    blip=shape._element.blipFill.blip
    ext_lst=blip.find(qn('a:extLst'))
    if ext_lst is None:return None
    node=ext_lst.find('.//{%s}svgBlip'%SVG_NAMESPACE)
    return node.get(qn('r:embed')) if node is not None else None


def attach_svg_source(shape,svg_bytes):
    """Embed the reviewed SVG as the picture's vector source with PNG fallback.

    PowerPoint 2016+ uses the SVG; LibreOffice and older viewers use the
    fallback PNG. No SVG-to-shape conversion is performed.
    """
    package=shape.part.package
    part=Part(package.next_partname('/ppt/media/image%d.svg'),'image/svg+xml',package,bytes(svg_bytes))
    relationship=shape.part.relate_to(part,RT.IMAGE)
    blip=shape._element.blipFill.blip
    ext_lst=blip.find(qn('a:extLst'))
    if ext_lst is None:ext_lst=etree.SubElement(blip,qn('a:extLst'))
    extension=etree.SubElement(ext_lst,qn('a:ext'));extension.set('uri',SVG_BLIP_URI)
    svg_blip=etree.SubElement(extension,'{%s}svgBlip'%SVG_NAMESPACE,nsmap={'asvg':SVG_NAMESPACE})
    svg_blip.set(qn('r:embed'),relationship)
    return relationship


def add_svg_picture(slide,png_path,svg_bytes,left,top,width,height):
    shape=slide.shapes.add_picture(str(png_path),left,top,width,height)
    attach_svg_source(shape,svg_bytes)
    return shape


def component_size_in(geometry_value,prs):
    return (geometry_value['w']*prs.slide_width/914400,geometry_value['h']*prs.slide_height/914400)


def check_component_geometry(entry,geometry_value,prs,aspect_key='aspect_ratio'):
    width_in,height_in=component_size_in(geometry_value,prs)
    aspect=width_in/height_in if height_in else 0
    if abs(aspect-entry[aspect_key])>entry[aspect_key]*.02:
        raise ValueError(f'Component geometry must keep the {entry[aspect_key]:.3f}:1 canvas ratio; received {aspect:.3f}:1')
    minimum=entry['min_size_in']
    if width_in+1e-6<minimum['w'] or height_in+1e-6<minimum['h']:
        raise ValueError(f'Component {entry["component_id"]} needs at least {minimum["w"]:.2f} x {minimum["h"]:.2f} in; received {width_in:.2f} x {height_in:.2f} in')


def add_component(slide,element,state,sid,prs,anchors,views):
    entry=component_validator_get(element['component'])
    g=geometry(element,anchors)
    check_component_geometry(entry,g,prs)
    result=component_renderer.render(registry(),profiles(),state['profile_id'],element['component'],
        element.get('params') or {},element.get('texts') or {},element.get('version'))
    if element.get('layer','content') not in ('content','background','decoration'):raise ValueError('Unknown layer intent')
    temporary=Path(tempfile.mkstemp(suffix='.png')[1])
    try:
        temporary.write_bytes(result['png'])
        shape=add_svg_picture(slide,temporary,result['svg'],*positions(g,prs))
    finally:
        temporary.unlink(missing_ok=True)
    shape.name='da:'+element['id']
    group=element.get('overlap_group') or 'component:'+element['id']
    meta=dict(kind='component',layer=element.get('layer','content'),anchor=element.get('anchor'),overlap_group=group,vector=True,
        component=dict(component_id=result['component_id'],version=result['version'],variant=result['variant'],
        params=result['params'],texts=result['texts'],svg_sha256=result['svg_sha256'],
        svg_template_sha256=result['svg_template_sha256'],profile_id=state['profile_id'],vector=True))
    renderer.set_metadata(shape,meta)
    children=[]
    for child in component_renderer.text_children(element['id'],entry,entry['variants'][result['variant']],result['texts'],g):
        if any(renderer.element_identity(item)==child['id'] for item in slide.shapes):raise ValueError('Duplicate stable element ID')
        textbox=slide.shapes.add_textbox(*positions(child['geometry'],prs))
        child_style=style_for(state,child['style']);child_style['align']=child['align']
        apply_text(textbox,child['text'],child_style)
        textbox.name='da:'+child['id']
        renderer.set_metadata(textbox,dict(kind='text',style=child['style'],text=child['text'],resolved_style=child_style,
            align=child['align'],layer='content',overlap_group=group,component_instance=element['id'],component_slot=child['slot']))
        children.append(child)
    meta['texts']={child['slot']:child['text'] for child in children}
    renderer.set_metadata(shape,meta)
    return shape,children


def rebuild_component(slide,shape,state,prs):
    """Re-render one component instance with its recorded version and parameters."""
    meta=renderer.metadata(shape); record=meta.get('component') or {}
    entry=component_validator_get(record.get('component_id'))
    result=component_renderer.render(registry(),profiles(),state['profile_id'],record['component_id'],
        record.get('params') or {},record.get('texts') or {},record.get('version'))
    left,top,width,height=shape.left,shape.top,shape.width,shape.height
    parent=shape._element.getparent();index=parent.index(shape._element);name=shape.name
    previous_svg=svg_rel_id(shape)
    g=dict(zip(('x','y','w','h'),(left/prs.slide_width,top/prs.slide_height,width/prs.slide_width,height/prs.slide_height)))
    check_component_geometry(entry,g,prs)
    temporary=Path(tempfile.mkstemp(suffix='.png')[1])
    try:
        temporary.write_bytes(result['png'])
        replacement=add_svg_picture(slide,temporary,result['svg'],left,top,width,height)
    finally:
        temporary.unlink(missing_ok=True)
    replacement.name=name
    parent.remove(replacement._element);parent.insert(index,replacement._element)
    meta['overlap_group']=meta.get('overlap_group') or 'component:'+record['component_id']
    meta['vector']=True
    meta['component']=dict(record,component_id=result['component_id'],version=result['version'],variant=result['variant'],
        params=result['params'],texts=result['texts'],svg_sha256=result['svg_sha256'],
        svg_template_sha256=result['svg_template_sha256'],profile_id=state['profile_id'],vector=True)
    renderer.set_metadata(replacement,meta)
    parent.remove(shape._element)
    if previous_svg:
        try:slide.part.drop_rel(previous_svg)
        except KeyError:pass
    return replacement,g


def sync_component_children(slide,instance_id,g,state,prs,record):
    entry=component_validator_get(record['component_id'])
    variant=entry['variants'][record['variant']]
    wanted=component_renderer.text_children(instance_id,entry,variant,record.get('texts') or {},g)
    existing={renderer.element_identity(shape):shape for shape in component_children(slide,instance_id)}
    keep=set()
    for child in wanted:
        keep.add(child['id']);style=style_for(state,child['style']);style['align']=child['align']
        shape=existing.get(child['id'])
        if shape is None:
            shape=slide.shapes.add_textbox(*positions(child['geometry'],prs));shape.name='da:'+child['id']
        set_geometry(shape,child['geometry'],prs)
        apply_text(shape,child['text'],style)
        renderer.set_metadata(shape,dict(kind='text',style=child['style'],text=child['text'],resolved_style=style,
            align=child['align'],layer='content',overlap_group='component:'+instance_id,component_instance=instance_id,component_slot=child['slot']))
    for element_id,shape in existing.items():
        if element_id not in keep:shape._element.getparent().remove(shape._element)


def apply_table_style(shape,state):
    profile=profile_of(state);tables=profile['tables'];palette=state['theme']
    for r,row in enumerate(shape.table.rows):
        for cell in row.cells:
            if r==0:fill_token=tables['header_fill_token']
            elif tables['zebra'] and r%2==0:fill_token='surface'
            else:fill_token='background'
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(palette.get(fill_token,profile['colors'][fill_token]))
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name=state['font_family'];run.font.size=Pt(tables['font_size'])
                    token=tables['header_text_token'] if r==0 else 'text'
                    run.font.color.rgb=RGBColor.from_string(palette.get(token,profile['colors'][token]))


def apply_chart_style(shape,state,unit=None):
    profile=profile_of(state);series_colors=profile['colors']['chart_series'];charts=profile['charts']
    palette=state['theme'];chart=shape.chart
    muted=RGBColor.from_string(palette['muted']);border=RGBColor.from_string(palette['border'])
    for index,series in enumerate(chart.series):
        color=RGBColor.from_string(series_colors[index%len(series_colors)])
        try:series.format.fill.solid();series.format.fill.fore_color.rgb=color
        except (ValueError,AttributeError):pass
        try:series.format.line.color.rgb=color
        except (ValueError,AttributeError):pass
    try:chart.has_title=False
    except (ValueError,AttributeError):pass
    try:
        single_series=len(chart.series)==1
        chart.has_legend=(charts['legend_position']!='none') and not single_series
        if chart.has_legend:
            chart.legend.position={'bottom':XL_LEGEND_POSITION.BOTTOM,'right':XL_LEGEND_POSITION.RIGHT,'top':XL_LEGEND_POSITION.TOP}[charts['legend_position']]
            chart.legend.include_in_layout=False
            chart.legend.font.name=state['font_family'];chart.legend.font.size=Pt(charts['label_size']);chart.legend.font.color.rgb=muted
        for axis in (chart.category_axis,chart.value_axis):
            axis.has_major_gridlines=charts['gridlines']
            axis.tick_labels.font.name=state['font_family'];axis.tick_labels.font.size=Pt(charts['label_size'])
            axis.tick_labels.font.color.rgb=muted
            try:axis.format.line.color.rgb=border
            except (ValueError,AttributeError):pass
            if charts['gridlines']:
                try:axis.major_gridlines.format.line.color.rgb=border
                except (ValueError,AttributeError):pass
        if unit:
            chart.value_axis.has_title=True
            chart.value_axis.axis_title.text_frame.text=str(unit)
            for paragraph in chart.value_axis.axis_title.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name=state['font_family'];run.font.size=Pt(charts['label_size']);run.font.color.rgb=muted
        else:chart.value_axis.has_title=False
    except (ValueError,KeyError,AttributeError):pass


def restyle_presentation(prs,state):
    """Re-apply the active profile to every generated object and slide background."""
    for slide in prs.slides:
        sid=renderer.slide_identity(slide)
        layout_id=state.get('layouts',{}).get(sid)
        views=None
        if layout_id:
            modules=state.get('layout_modules',{}).get(sid,layouts().get(layout_id)['default_modules'])
            _,views=slot_views(state,layout_id,modules)
            state['layout_slots'][sid]={name:view['rect'] for name,view in views.items()}
        anchors=anchors_for(state,sid)
        background=state['theme']['background']
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(background)
        state.setdefault('backgrounds',{})[sid]=background
        components_to_rebuild=[]
        for shape in list(slide.shapes):
            meta=renderer.metadata(shape)
            if not meta.get('kind'):continue
            if meta['kind']=='text':
                meta['resolved_style']=resolve_style(state,meta,views)
                if meta.get('align'):meta['resolved_style']['align']=meta['align']
                apply_text(shape,meta.get('text') or shape.text,meta['resolved_style'])
                renderer.set_metadata(shape,meta)
            elif meta['kind']=='shape':
                if meta.get('fill_token'):
                    meta['fill']=state['theme'].get(meta['fill_token'],meta['fill'])
                    if meta.get('shape')=='line':shape.line.color.rgb=RGBColor.from_string(meta['fill'])
                    else:shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(meta['fill'])
                renderer.set_metadata(shape,meta)
            elif meta['kind']=='table':apply_table_style(shape,state)
            elif meta['kind']=='chart':apply_chart_style(shape,state,(meta.get('chart') or {}).get('unit'))
            elif meta['kind']=='component':components_to_rebuild.append(shape)
            if meta.get('anchor') and meta['anchor'] in anchors:set_geometry(shape,geometry(meta,anchors),prs)
        for shape in components_to_rebuild:
            meta=renderer.metadata(shape);replacement,g=rebuild_component(slide,shape,state,prs)
            sync_component_children(slide,renderer.element_identity(replacement),g,state,prs,renderer.metadata(replacement)['component'])


class PresentationStore(shared.ArtifactStore):
    def __init__(self,path,temporary_root=None):super().__init__(path,'.pptx','.pptx-state',temporary_root)
    def assert_current(self,state,base_revision=None):
        if state is None:raise ValueError('Presentation not initialized')
        if base_revision is not None and (type(base_revision) is not int or base_revision!=state['revision']):raise ValueError('Revision conflict')
        if not self.path.exists() or hashlib.sha256(self.path.read_bytes()).hexdigest()!=state['sha256']:raise ValueError('External modification; inspect or restore the last snapshot')
    def design_event(self,state):
        return dict(profile_id=state.get('profile_id'),theme=state.get('theme'),font_family=state.get('font_family'),
            layouts=state.get('layouts',{}),layout_variants=state.get('layout_variants',{}),layout_modules=state.get('layout_modules',{}),
            layout_slots=state.get('layout_slots',{}),backgrounds=state.get('backgrounds',{}),
            template_font=state.get('template_font'),template_accent=state.get('template_accent'),
            conflicts=state.get('design_conflicts',[]))
    def commit(self,c,state,data,operation_id,request,changed,kind='presentation_revision'):
        validator.validate_package(data)
        state.update(revision=state['revision']+1,sha256=hashlib.sha256(data).hexdigest(),pending_publish=True,finalized=False)
        prs=Presentation(io.BytesIO(data))
        if changed==['*']:
            state['reviews']={}; state['last_reviews']={}; state['renders']={}; state['repair_attempts']={}
        else:
            for sid in changed:
                old=state.setdefault('reviews',{}).pop(sid,None)
                if old:state.setdefault('last_reviews',{})[sid]=old
        view=renderer.manifest(prs,state)
        event=dict(status='committed',presentation_id=state['presentation_id'],revision=state['revision'],sha256=state['sha256'],operation_id=operation_id,
            slide_id=request.get('slide_id') or state.get('active_slide'),changed_elements=[request.get('element',{}).get('id','*')],changed_slides=changed,
            manifest=view,path=self.path.name,theme=state['theme'],font_family=state['font_family'],profile_id=state.get('profile_id'),
            design=self.design_event(state),committed_at=time.time())
        self.persist_revision(c,state,data,event,request,operation_id,kind)
        return event
    def init(self,presentation_id,template=None,mode='normal',theme=DEFAULT_PROFILE,profile=None):
        identifier(presentation_id)
        profile_id=profile or theme
        if mode not in ('normal','strict'):raise ValueError('Unknown review mode')
        profiles().get(profile_id)
        with self.locked() as (c,state):
            if state:
                if state['presentation_id']!=presentation_id:raise ValueError('Presentation identity mismatch')
                return self.view(c,state)
            if self.path.exists():raise ValueError('File already exists')
            if template:
                validator.validate_package(Path(template).read_bytes()); prs=Presentation(template)
            else:
                prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
            state=dict(presentation_id=presentation_id,revision=-1,mode=mode,slots={},layouts={},layout_variants={},layout_modules={},
                layout_slots={},backgrounds={},reviews={},renders={},repair_attempts={},template=bool(template),template_complex={})
            apply_profile(state,profile_id)
            if template:
                from pptx.opc.constants import RELATIONSHIP_TYPE as RT
                from lxml import etree
                theme_part=prs.slide_masters[0].part.part_related_by(RT.THEME)
                xml=etree.fromstring(theme_part.blob)
                ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
                accent=xml.xpath('string(//a:clrScheme/a:accent1/a:srgbClr/@val)',namespaces=ns)
                if len(accent)==6:state['template_accent']=accent.upper()
                faces=xml.xpath('//a:fontScheme/a:minorFont/a:ea/@typeface',namespaces=ns)
                if faces and faces[0]:state['template_font']=faces[0]
                apply_profile(state,profile_id)
            for slide in prs.slides:
                sid=renderer.slide_identity(slide)
                state['template_complex'][sid]=bool(template)
            state['active_slide']=renderer.slide_identity(prs.slides[0]) if len(prs.slides) else None
            return self.commit(c,state,serialize(prs),'__init__',{'template':str(template)},['*'],'presentation_initialized')
    def view(self,c,state):
        if state is None:raise ValueError('Presentation not initialized')
        data=c.execute('SELECT docx FROM revisions WHERE revision=?',(state['revision'],)).fetchone()[0]
        return dict(presentation_id=state['presentation_id'],revision=state['revision'],sha256=state['sha256'],path=self.path.name,
            manifest=renderer.manifest(Presentation(io.BytesIO(data)),state),active_slide=state.get('active_slide'),finalized=state.get('finalized',False),
            mode=state['mode'],theme=state['theme'],profile_id=state.get('profile_id'),
            history=[{key:value for key,value in json.loads(r[0]).items() if key in ('revision','operation_id','sha256','slide_id','status')} for r in c.execute('SELECT result FROM revisions ORDER BY revision')])
    def status(self):
        with self.locked() as (c,state):return self.view(c,state)
    def track_external(self, presentation_id):
        """Index a valid existing PPTX without changing its package bytes."""
        with self.locked() as (c, state):
            if state is not None:
                return None
            data = self.path.read_bytes()
            validator.validate_package(data)
            prs = Presentation(io.BytesIO(data))
            state = dict(presentation_id=identifier(presentation_id), revision=-1, mode='normal',
                slots={}, layouts={}, layout_variants={}, layout_modules={}, layout_slots={}, backgrounds={},
                reviews={}, renders={}, repair_attempts={}, template=True,
                template_complex={renderer.slide_identity(slide): True for slide in prs.slides},
                active_slide=renderer.slide_identity(prs.slides[0]) if len(prs.slides) else None)
            apply_profile(state,DEFAULT_PROFILE)
            return self.commit(c, state, data, '__external__', {'external': True}, ['*'], 'presentation_initialized')
    def preview(self):
        """Generate display images while retaining pending content/visual checks."""
        current = self.status()
        revision = current['revision']
        reviewer = load('slide_reviewer')
        reviewer.render_snapshot(self, revision)
        with self.locked() as (c, state):
            self.assert_current(state, revision)
            for index, slide in enumerate(current['manifest']['slides']):
                state['renders'][slide['slide_id']] = dict(revision=revision, image=f'slide-{index+1:03d}.png')
            return self.emit_state(c, state, dict(type='presentation_preview',
                presentation_id=state['presentation_id'], revision=revision, path=self.path.name,
                manifest=self.view(c, state)['manifest']))
    def apply_one(self,request,asset_root=None):
        identifier(request['operation_id'])
        if len(json.dumps(request,ensure_ascii=False))>30000:raise ValueError('Bounded operation payload exceeded')
        with self.locked() as (c,state):
            if not state or request['presentation_id']!=state['presentation_id']:raise ValueError('Presentation identity mismatch')
            old=c.execute('SELECT request,result FROM revisions WHERE operation_id=?',(request['operation_id'],)).fetchone()
            if old:
                if json.loads(old[0])!=request:raise ValueError('Operation ID reused with different request')
                return json.loads(old[1])
            self.assert_current(state,request['base_revision']); prs=Presentation(io.BytesIO(self.snapshot_unlocked(c,state['revision'])))
            state=copy.deepcopy(state); normalize_state(state); op=request['operation']; sid=request.get('slide_id'); changed=[sid] if sid else ['*']
            if op=='slide.add':
                for slide in prs.slides:
                    review=state['reviews'].get(renderer.slide_identity(slide),{})
                    if review.get('status')=='blocked' or not review.get('automatic_passed') or (state['mode']=='strict' and review.get('visual_review')!='passed'):raise ValueError('Check all affected slides before creating the next slide')
                identifier(sid)
                if any(renderer.slide_identity(s)==sid for s in prs.slides):raise ValueError('Duplicate slide ID')
                index=request.get('layout_index',next((i for i,layout in enumerate(prs.slide_layouts) if layout.name.lower()=='blank'),0))
                if type(index) is not int or not 0<=index<len(prs.slide_layouts):raise ValueError('Unknown template layout')
                layout_id=request.get('layout','title_body')
                modules=request.get('modules',layouts().get(layout_id)['default_modules'])
                if type(modules) is not int:raise ValueError('modules must be an integer')
                views=set_slide_layout(state,sid,layout_id,modules)
                slide=prs.slides.add_slide(prs.slide_layouts[index]); slide._element.cSld.set('name','da-slide:'+sid)
                state.setdefault('slots',{})[sid]={}
                if state['template']:
                    for placeholder in slide.placeholders:
                        name='title' if placeholder.is_placeholder and placeholder.placeholder_format.type in (1,3) else 'body'
                        state['slots'][sid][name]=(placeholder.left/prs.slide_width,placeholder.top/prs.slide_height,placeholder.width/prs.slide_width,placeholder.height/prs.slide_height)
                state['active_slide']=sid
                if state['template']:state['template_complex'][sid]=True
                slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string(state['theme']['background']); state['backgrounds'][sid]=state['theme']['background']
                state['layout_slots'][sid]={name:view['rect'] for name,view in views.items()}
            elif op in ('presentation.set_theme','presentation.set_profile'):
                target=request.get('profile') or request.get('theme')
                apply_profile(state,target); changed=['*']
                restyle_presentation(prs,state)
            elif op=='slide.reorder':
                slide=find_slide(prs,sid); index=request['index']
                if type(index)is not int or not 0<=index<len(prs.slides):raise ValueError('Invalid slide order')
                node=next(n for n in prs.slides._sldIdLst if int(n.get('id'))==slide.slide_id); prs.slides._sldIdLst.remove(node); prs.slides._sldIdLst.insert(index,node); changed=['*']
            else:
                slide=find_slide(prs,sid)
                if state.get('repair_attempts',{}).get(sid,0)>=3 and state.get('reviews',{}).get(sid,{}).get('status')=='blocked':raise ValueError('Slide repair budget exhausted; restore a known version or request review')
                views=None
                if state.get('layouts',{}).get(sid):
                    views=slot_views(state,state['layouts'][sid],state.get('layout_modules',{}).get(sid,0))[1]
                anchors=anchors_for(state,sid)
                if op=='slide.set_layout':
                    layout_id=request['layout']
                    modules=request.get('modules',state.get('layout_modules',{}).get(sid,layouts().get(layout_id)['default_modules']))
                    if not profiles().allows_layout(state['profile_id'],layout_id):raise ValueError(f'Profile {state["profile_id"]} does not enable layout {layout_id}')
                    target_views=slot_views(state,layout_id,modules)[1]
                    anchored={renderer.metadata(shape).get('anchor'):renderer.metadata(shape) for shape in slide.shapes if renderer.metadata(shape).get('anchor')}
                    missing=sorted(name for name in anchored if name not in target_views)
                    if missing:raise ValueError('Layout '+layout_id+' cannot keep anchored content: missing slot(s) '+', '.join(missing)+'; keep the layout or move those elements to explicit geometry')
                    for name,meta in anchored.items():
                        if meta.get('kind')=='text' and name not in state.get('slots',{}).get(sid,{}):
                            problem=layouts().capacity_error(layout_id,modules,name,dict(type='text',text=meta.get('text','')))
                            if problem:raise ValueError('Layout '+layout_id+' cannot hold the existing text safely: '+problem)
                    set_slide_layout(state,sid,layout_id,modules)
                    anchors=anchors_for(state,sid)
                    for shape in slide.shapes:
                        meta=renderer.metadata(shape)
                        if meta.get('anchor') and meta['anchor'] in anchors:set_geometry(shape,geometry(meta,anchors),prs)
                    views=target_views
                elif op=='element.remove':
                    shape=find_element(slide,request['element']['id'])
                    if renderer.metadata(shape).get('kind')=='component':
                        for child in component_children(slide,request['element']['id']):child._element.getparent().remove(child._element)
                        vector=svg_rel_id(shape)
                    else:vector=None
                    shape._element.getparent().remove(shape._element)
                    if vector:
                        try:slide.part.drop_rel(vector)
                        except KeyError:pass
                elif op=='element.update':
                    element=request['element']; shape=find_element(slide,element['id']); meta=renderer.metadata(shape)
                    if set(element)-{'id','text','geometry','anchor','fill','style','layer','overlap_group','rows','chart','params','texts','component_version'}:raise ValueError('Unsupported element update field')
                    target_anchor=element.get('anchor') or meta.get('anchor')
                    if 'text' in element:check_anchor_content(state,sid,target_anchor,{'type':'text','text':element['text']})
                    if 'rows' in element:check_anchor_content(state,sid,target_anchor,{'type':'table'})
                    if 'chart' in element:check_anchor_content(state,sid,target_anchor,{'type':'chart'})
                    if element.get('anchor') and not ({'text','rows','chart'}&set(element)):
                        check_anchor_content(state,sid,element['anchor'],{'type':meta.get('kind','text'),'text':meta.get('text','')})
                    if meta.get('kind')=='component' and ('params' in element or 'texts' in element or 'component_version' in element):
                        record=dict(meta.get('component') or {})
                        params={**(record.get('params') or {}),**(element.get('params') or {})}
                        texts=dict(record.get('texts') or {})
                        for slot,value in (element.get('texts') or {}).items():
                            if value is None:texts.pop(slot,None)
                            else:texts[slot]=value
                        merged=dict(record,params=params,texts=texts)
                        if 'component_version' in element:merged['version']=element['component_version']
                        meta['component']=merged;renderer.set_metadata(shape,meta)
                        replacement,g=rebuild_component(slide,shape,state,prs)
                        sync_component_children(slide,element['id'],g,state,prs,renderer.metadata(replacement)['component'])
                        shape=replacement;meta=renderer.metadata(replacement)
                    if 'rows' in element:
                        rows=element['rows']
                        if not shape.has_table or len(rows)!=len(shape.table.rows) or any(len(row)!=len(shape.table.columns) or any(not isinstance(text,str) or len(text)>160 for text in row) for row in rows):
                            raise ValueError('Table updates require bounded rows with existing dimensions')
                        for r,row in enumerate(rows):
                            for col,text in enumerate(row):
                                cell=shape.table.cell(r,col);cell.text=text
                                for paragraph in cell.text_frame.paragraphs:
                                    for run in paragraph.runs:run.font.name=state['font_family'];run.font.size=Pt(profile_of(state)['tables']['font_size'])
                        apply_table_style(shape,state);meta['rows']=rows
                    if 'chart' in element:
                        if not shape.has_chart:raise ValueError('Chart data update requires a native chart')
                        if element['chart'].get('type','bar')!=meta.get('chart',{}).get('type','bar'):raise ValueError('Changing chart type requires replacing the chart object')
                        shape.chart.replace_data(chart_data(element['chart']));meta['chart']=element['chart'];apply_chart_style(shape,state,element['chart'].get('unit'))
                    if element.get('geometry') or element.get('anchor'):
                        if meta.get('kind')=='component':check_component_geometry(component_validator_get(meta['component']['component_id']),geometry(element,anchors),prs)
                        set_geometry(shape,geometry(element,anchors),prs)
                        if meta.get('kind')=='component':
                            g=geometry(element,anchors);sync_component_children(slide,element['id'],g,state,prs,meta['component'])
                    if 'text'in element:
                        if not shape.has_text_frame:raise ValueError('Text update requires a text element')
                        if not meta:meta=dict(kind='text',style=element.get('style','body'))
                        meta.update(text=element['text'],style=element.get('style',meta.get('style','body')))
                        meta['resolved_style']=resolve_style(state,meta,views);apply_text(shape,meta['text'],meta['resolved_style'])
                        if meta.get('component_instance'):
                            record=renderer.metadata(find_element(slide,meta['component_instance'])).get('component',{})
                            entry=registry()['components'].get(record.get('component_id'))
                            if entry and record.get('variant') in entry['variants']:
                                slot=entry['variants'][record['variant']]['text_slots'].get(meta.get('component_slot'))
                                if slot and (len(meta['text'])>slot['chars'] or meta['text'].count('\n')+1>slot['lines']):raise ValueError('Component text slot capacity exceeded')
                                record.setdefault('texts',{})[meta['component_slot']]=meta['text']
                                parent=find_element(slide,meta['component_instance']);pmeta=renderer.metadata(parent);pmeta['component']=record;renderer.set_metadata(parent,pmeta)
                    if 'fill'in element:
                        color=element['fill']; check_color(color)
                        if hasattr(shape,'fill'):shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color)
                        else:shape.line.color.rgb=RGBColor.from_string(color)
                        meta['fill']=color
                    if 'style'in element and shape.has_text_frame:
                        meta['style']=element['style']; meta['resolved_style']=resolve_style(state,meta,views); apply_text(shape,shape.text,meta['resolved_style'])
                    meta.update({k:element[k] for k in ('anchor','layer','overlap_group') if k in element}); renderer.set_metadata(shape,meta)
                elif op.startswith('element.add_'):
                    element=request['element']; identifier(element['id'])
                    if any(renderer.element_identity(s)==element['id'] for s in slide.shapes):raise ValueError('Duplicate stable element ID')
                    kind=op.removeprefix('element.add_')
                    check_anchor_content(state,sid,element.get('anchor'),{'type':kind,'text':element.get('text','')})
                    if kind=='component':
                        add_component(slide,element,state,sid,prs,anchors,views)
                        return self.commit(c,state,serialize(prs),request['operation_id'],request,changed)
                    g=geometry(element,anchors); pos=positions(g,prs)
                    meta=dict(element,kind=kind,resolved_style=resolve_style(state,element,views))
                    if meta.get('layer','content') not in ('content','background','decoration'):raise ValueError('Unknown layer intent')
                    if kind=='text':
                        shape=slide.shapes.add_textbox(*pos); apply_text(shape,element['text'],meta['resolved_style'])
                    elif kind=='shape' and element.get('shape')=='line':
                        if g['w']>=g['h']:x1,y1,x2,y2=pos[0],pos[1]+pos[3]//2,pos[0]+pos[2],pos[1]+pos[3]//2
                        else:x1,y1,x2,y2=pos[0]+pos[2]//2,pos[1],pos[0]+pos[2]//2,pos[1]+pos[3]
                        shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x1,y1,x2,y2)
                        meta['shape']='line';meta['fill']=state['theme']['accent'];shape.line.color.rgb=RGBColor.from_string(meta['fill']);shape.line.width=Pt(1.5)
                    elif kind=='shape':
                        shape_type={'rectangle':MSO_AUTO_SHAPE_TYPE.RECTANGLE,'rounded_rectangle':MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,'arrow':MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,'ellipse':MSO_AUTO_SHAPE_TYPE.OVAL}.get(element.get('shape','rectangle'))
                        if shape_type is None:raise ValueError('Unsupported native shape')
                        shape=slide.shapes.add_shape(shape_type,*pos); meta['shape']=element.get('shape','rectangle')
                        meta['fill_token']=element.get('fill_token','card')
                        if meta['fill_token'] not in state['theme']:raise ValueError('Unknown fill token')
                        meta['fill']=state['theme'][meta['fill_token']]
                        shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(meta['fill']); shape.line.fill.background()
                    elif kind=='image':
                        asset=(Path(asset_root or self.path.parent)/element['path']).resolve(); assets=Path(asset_root or self.path.parent).resolve()
                        if not asset.is_relative_to(assets) or not asset.is_file() or asset.stat().st_size>12*1024*1024:raise ValueError('Unsafe or oversized image')
                        from PIL import Image
                        with Image.open(asset) as image:
                            if image.format not in ('PNG','JPEG') or image.width*image.height>40000000:raise ValueError('Unsupported bounded image')
                            if abs(image.width/image.height-g['w']*prs.slide_width/(g['h']*prs.slide_height))>.03:raise ValueError('Image geometry must preserve aspect ratio')
                        shape=slide.shapes.add_picture(str(asset),*pos)
                    elif kind=='table':
                        rows=element['rows']
                        if not isinstance(rows,list) or not 1<=len(rows)<=10 or not 1<=len(rows[0])<=8 or any(len(row)!=len(rows[0]) or any(not isinstance(t,str) or len(t)>160 for t in row) for row in rows):raise ValueError('Bounded uniform native table required')
                        shape=slide.shapes.add_table(len(rows),len(rows[0]),*pos)
                        for r,row in enumerate(rows):
                            for col,text in enumerate(row):
                                cell=shape.table.cell(r,col); cell.text=text
                                for paragraph in cell.text_frame.paragraphs:
                                    for run in paragraph.runs:run.font.name=state['font_family']; run.font.size=Pt(profile_of(state)['tables']['font_size'])
                        apply_table_style(shape,state)
                    elif kind=='chart':
                        chart=element['chart']; categories=chart['categories']; series=chart['series']
                        if not 1<=len(categories)<=12 or any(not isinstance(label,str) or len(label)>80 for label in categories) or not 1<=len(series)<=4 or any(not isinstance(item.get('name'),str) or len(item['name'])>80 for item in series) or any(len(s['values'])!=len(categories) or any(type(n)not in (int,float) or not math.isfinite(n) for n in s['values']) for s in series):raise ValueError('Bounded finite editable chart data required')
                        chart_type={'bar':XL_CHART_TYPE.COLUMN_CLUSTERED,'line':XL_CHART_TYPE.LINE,'pie':XL_CHART_TYPE.PIE}.get(chart.get('type','bar'))
                        if chart_type is None:raise ValueError('Unsupported native chart')
                        shape=slide.shapes.add_chart(chart_type,*pos,chart_data(chart)); apply_chart_style(shape,state,chart.get('unit'))
                    else:raise ValueError('Unsupported native element type')
                    shape.name='da:'+element['id']; renderer.set_metadata(shape,meta)
                else:raise ValueError('Unsupported atomic presentation operation')
            return self.commit(c,state,serialize(prs),request['operation_id'],request,changed)
    def prepare_restore(self,c,state,revision):
        prior=json.loads(c.execute('SELECT result FROM revisions WHERE revision=?',(revision,)).fetchone()[0])
        state=copy.deepcopy(state);design=prior.get('design') or {}
        if design:
            for key in ('theme','font_family','layouts','layout_variants','layout_modules','layout_slots','backgrounds','template_font','template_accent','design_conflicts'):
                if key in design:state[key]=design[key]
            state['profile_id']=design.get('profile_id',state.get('profile_id',DEFAULT_PROFILE));state['theme_name']=state['profile_id']
            try:state['profile']=profiles().get(state['profile_id'])
            except ValueError:apply_profile(state,DEFAULT_PROFILE)
        else:
            state['theme']=prior.get('theme',state.get('theme'));state['font_family']=prior.get('font_family',state.get('font_family'))
            state['profile_id']=state.get('profile_id') or state.get('theme_name') or DEFAULT_PROFILE;state['theme_name']=state['profile_id']
            try:state['profile']=profiles().get(state['profile_id'])
            except ValueError:apply_profile(state,DEFAULT_PROFILE)
            state['layouts']={slide['slide_id']:slide.get('layout','title_body') for slide in prior.get('manifest',{}).get('slides',[])}
            state['backgrounds']={slide['slide_id']:slide.get('background','FFFFFF') for slide in prior.get('manifest',{}).get('slides',[])}
            for sid,layout_id in state['layouts'].items():
                try:modules=layouts().get(layout_id)['default_modules'];key,views=slot_views(state,layout_id,modules)
                except ValueError:continue
                state.setdefault('layout_variants',{})[sid]=key;state.setdefault('layout_modules',{})[sid]=modules
                state.setdefault('layout_slots',{})[sid]={name:view['rect'] for name,view in views.items()}
        state['active_slide']=prior.get('slide_id')
        return state

    def snapshot_unlocked(self,c,revision):return c.execute('SELECT docx FROM revisions WHERE revision=?',(revision,)).fetchone()[0]
    def emit_state(self,c,state,event):
        self.save_state(c,state); c.execute('INSERT INTO events(data) VALUES (?)',(json.dumps(event),)); c.commit(); return event
    def check_slide(self,slide_id,content_notes):
        reviewer=load('slide_reviewer'); return reviewer.check(self,slide_id,content_notes)
    def review(self,slide_id,revision,passed,notes,reviewer='vision_model'):
        if reviewer not in ('vision_model','human') or type(passed)is not bool or not notes.strip():raise ValueError('Actual visual inspection and evidence required')
        with self.locked() as(c,state):
            self.assert_current(state,revision); review=state['reviews'].get(slide_id)
            if not review or not review['automatic_passed'] or review['revision']>revision:raise ValueError('Exact-revision successful formal check required')
            formal=state.get('renders',{}).get(slide_id)
            if not formal or not (self.render_directory(formal['revision'])/formal['image']).is_file():raise ValueError('Formal screenshot is missing')
            review.update(visual_review='passed' if passed else 'failed',status='passed' if passed else 'blocked',visual_evidence=dict(reviewer=reviewer,notes=notes))
            return self.emit_state(c,state,dict(type='presentation_review',presentation_id=state['presentation_id'],revision=revision,slide_id=slide_id,manifest=renderer.manifest(Presentation(io.BytesIO(self.snapshot_unlocked(c,revision))),state),path=self.path.name))
    def finalize(self,base_revision):
        with self.locked() as(c,state):
            self.assert_current(state,base_revision); prs=Presentation(io.BytesIO(self.snapshot_unlocked(c,base_revision)))
            if not len(prs.slides):raise ValueError('Empty presentation')
            for slide in prs.slides:
                review=state['reviews'].get(renderer.slide_identity(slide),{})
                if review.get('status')!='passed' or not review.get('automatic_passed') or review.get('visual_review')!='passed' or not review.get('content_notes'):raise ValueError('Every slide requires actual content and visual acceptance')
            validator.validate_package(self.snapshot_unlocked(c,base_revision)); state['finalized']=True
            return self.emit_state(c,state,dict(type='presentation_finalized',presentation_id=state['presentation_id'],revision=base_revision,finalized=True,path=self.path.name))


def catalog():
    return dict(profiles=profiles().summary(),layouts=layouts().summary(),components=component_renderer.catalog(registry()))


def plan_request(store,request):
    profile_id=request.get('profile') or (store.status().get('profile_id') if store else None) or DEFAULT_PROFILE
    return layout_resolver.plan(profiles(),layouts(),profile_id,request.get('intent'),request.get('content'))


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('command',choices=['init','apply-one','status','restore','check-slide','review-slide','finalize','profiles','layouts','components','catalog','plan']); parser.add_argument('pptx',nargs='?'); parser.add_argument('--request'); parser.add_argument('--presentation-id'); parser.add_argument('--template'); parser.add_argument('--profile'); parser.add_argument('--theme'); parser.add_argument('--slide-id'); parser.add_argument('--revision',type=int); parser.add_argument('--base-revision',type=int); parser.add_argument('--operation-id'); parser.add_argument('--notes',default=''); parser.add_argument('--passed',action='store_true'); parser.add_argument('--reviewer',choices=['human','vision_model'],default='vision_model'); args=parser.parse_args()
    if args.command in ('profiles','layouts','components','catalog'):
        result=catalog() if args.command=='catalog' else {args.command:catalog()[args.command]}
        print(json.dumps(result,ensure_ascii=False)); return
    if not args.pptx:parser.error('a presentation path is required')
    store=PresentationStore(args.pptx)
    if args.command=='plan':result=plan_request(store,json.loads(Path(args.request).read_text()))
    elif args.command=='init':result=store.init(args.presentation_id,args.template,theme=args.theme or args.profile or DEFAULT_PROFILE,profile=args.profile)
    elif args.command=='apply-one':result=store.apply_one(json.loads(Path(args.request).read_text()),Path(args.pptx).parent)
    elif args.command=='status':result=store.status()
    elif args.command=='restore':result=store.restore(args.revision,args.base_revision,args.operation_id)
    elif args.command=='check-slide':result=store.check_slide(args.slide_id,args.notes)
    elif args.command=='review-slide':result=store.review(args.slide_id,args.revision,args.passed,args.notes,args.reviewer)
    else:result=store.finalize(args.base_revision)
    print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':
    try:main()
    except Exception as exc:print(json.dumps(dict(status='rejected',error=str(exc))),file=sys.stderr);sys.exit(1)
