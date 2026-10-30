"""Bounded native styling, presets and relative geometry for generated PPTX objects.

The vocabulary is deliberately small: colours, fonts, strokes, a soft shadow,
a two-colour linear gradient, a few chart and table switches and a handful of
presets. Defaults stay flat; glow, bevel, reflection, 3-D and similar effects
are rejected rather than approximated. Every explicit override is stored in
element metadata so profile changes and data updates re-apply it after the
profile defaults (explicit user style wins).
"""
import io
import math

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

FORBIDDEN_EFFECTS = ('glow', 'bevel', 'reflection', 'soft_edges', '3d', 'three_d', 'inner_shadow', 'neon')
SHADOWS = ('none', 'soft')
CHART_TYPES = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE_MARKERS, 'pie': XL_CHART_TYPE.PIE}
NUMBER_FORMATS = ('General', '0', '0.0', '0.00', '#,##0', '#,##0.0', '0%', '0.0%')
DASHES = {'solid': MSO_LINE_DASH_STYLE.SOLID, 'dash': MSO_LINE_DASH_STYLE.DASH, 'dot': MSO_LINE_DASH_STYLE.ROUND_DOT}
ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}

# Allowed override keys per element kind. Anything else is rejected.
STYLE_KEYS = {
    'text': {'color', 'bold', 'italic', 'font_size', 'font_family', 'align'},
    'shape': {'fill', 'gradient', 'shadow', 'line_color', 'line_width_pt', 'dash', 'text_color', 'bold', 'font_size'},
    'connector': {'line_color', 'line_width_pt', 'dash'},
    'table': {'header_fill', 'header_text', 'header_bold', 'body_fill', 'body_text', 'divider_color', 'font_size', 'zebra'},
    'chart': {'series_colors', 'gridlines', 'data_labels', 'number_format', 'emphasis', 'legend', 'font_size', 'line_width_pt'},
}

PRESETS = {
    'chart.clean': dict(kind='chart', description='light value-axis gridlines, no chart border, explicit markers on lines',
        style=dict(gridlines='light', data_labels=False, legend='auto')),
    'chart.labeled': dict(kind='chart', description='no gridlines; every data point carries its value label',
        style=dict(gridlines='none', data_labels=True, legend='auto')),
    'chart.emphasis': dict(kind='chart', description='one emphasised data point in the accent colour, the rest muted, values labelled',
        style=dict(gridlines='none', data_labels=True, emphasis='max', legend='auto')),
    'table.clean': dict(kind='table', description='soft header, thin horizontal dividers, generous cell padding',
        style=dict(header_fill='card', header_text='text', header_bold=True, body_fill='background', divider_color='border', zebra=False)),
    'table.header_accent': dict(kind='table', description='accent header with white text, thin dividers',
        style=dict(header_fill='accent', header_text='FFFFFF', header_bold=True, body_fill='background', divider_color='border', zebra=False)),
    'table.minimal': dict(kind='table', description='no fills; bold header separated by dividers only',
        style=dict(header_fill='background', header_text='text', header_bold=True, body_fill='background', divider_color='border', zebra=False)),
    'shape.card': dict(kind='shape', description='soft card surface without border',
        style=dict(fill='card', line_color='none')),
    'shape.outline': dict(kind='shape', description='transparent surface with a thin border',
        style=dict(fill='none', line_color='border', line_width_pt=1.0)),
    'shape.accent': dict(kind='shape', description='solid accent surface without border',
        style=dict(fill='accent', line_color='none')),
    'shape.subtle': dict(kind='shape', description='near-background surface with a hairline border',
        style=dict(fill='surface', line_color='border', line_width_pt=.75)),
    'shape.raised': dict(kind='shape', description='soft card surface lifted by a faint shadow',
        style=dict(fill='card', line_color='none', shadow='soft')),
    'shape.gradient': dict(kind='shape', description='gentle two-colour gradient from card to surface',
        style=dict(gradient=dict(start='card', end='surface', angle=90), line_color='none')),
    'text.emphasis': dict(kind='text', description='accent colour, bold', style=dict(color='accent', bold=True)),
    'text.muted': dict(kind='text', description='muted secondary text', style=dict(color='muted', bold=False)),
}


def catalog():
    return {name: dict(kind=item['kind'], description=item['description'], style=dict(item['style'])) for name, item in PRESETS.items()}


def style_kind(meta):
    kind = meta.get('kind')
    if kind == 'shape' and meta.get('shape') == 'line':
        return 'connector'
    return kind


def _hex(value):
    text = value.strip().lstrip('#')
    if len(text) == 6 and all(c in '0123456789abcdefABCDEF' for c in text):
        return text.upper()
    return None


def resolve_color(value, state, allow_none=False):
    """Profile token or RRGGBB. ``none`` only where transparency is meaningful."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError('Colour must be a profile token or RRGGBB value')
    if value.strip().lower() == 'none':
        if allow_none:
            return None
        raise ValueError('This colour cannot be transparent')
    theme = state.get('theme') or {}
    profile_colors = (state.get('profile') or {}).get('colors') or {}
    if value in theme:
        return theme[value]
    if value in profile_colors and isinstance(profile_colors[value], str):
        return profile_colors[value]
    color = _hex(value)
    if color is None:
        raise ValueError(f'Unknown colour {value!r}; use a profile token ({", ".join(sorted(theme))}) or RRGGBB')
    return color


def blend(first, second, ratio):
    a, b = _hex(first), _hex(second)
    channels = [round(int(a[i:i + 2], 16) * (1 - ratio) + int(b[i:i + 2], 16) * ratio) for i in (0, 2, 4)]
    return ''.join(f'{max(0, min(255, c)):02X}' for c in channels)


def _check_forbidden(style):
    for key in style:
        lowered = str(key).lower()
        if any(effect in lowered for effect in FORBIDDEN_EFFECTS):
            raise ValueError(f'Decorative effect {key!r} is not supported; use at most a soft shadow or a two-colour gradient (no glow, bevel, reflection or 3-D)')


def validate_style(kind, style, state):
    """Return a normalised copy of an override dict, or raise ValueError."""
    if not isinstance(style, dict) or not style:
        raise ValueError('style must be a non-empty object')
    _check_forbidden(style)
    allowed = STYLE_KEYS.get(kind)
    if allowed is None:
        raise ValueError(f'Element kind {kind!r} does not support style overrides')
    unknown = sorted(set(style) - allowed)
    if unknown:
        raise ValueError(f'Unsupported {kind} style field(s): {", ".join(unknown)}; allowed: {", ".join(sorted(allowed))}')
    quality = (state.get('profile') or {}).get('quality') or {}
    minimum = float(quality.get('min_font_pt', 10))
    result = {}
    for key, value in style.items():
        if key in ('color', 'text_color', 'header_text', 'body_text'):
            resolve_color(value, state)
        elif key in ('fill', 'line_color', 'header_fill', 'body_fill', 'divider_color'):
            resolve_color(value, state, allow_none=True)
        elif key in ('bold', 'italic', 'header_bold', 'zebra', 'data_labels'):
            if type(value) is not bool:
                raise ValueError(f'{key} must be true or false')
        elif key == 'font_size':
            if type(value) not in (int, float) or not math.isfinite(value) or not minimum <= value <= 96:
                raise ValueError(f'font_size must be between the profile minimum {minimum:g} pt and 96 pt')
        elif key == 'font_family':
            fonts = (state.get('profile') or {}).get('fonts') or {}
            allowed_fonts = {fonts.get('title'), fonts.get('body'), state.get('template_font'), *(fonts.get('fallback') or [])} - {None}
            if value not in allowed_fonts:
                raise ValueError('font_family must be one of the profile/template fonts: ' + ', '.join(sorted(allowed_fonts)))
        elif key == 'align':
            if value not in ALIGN:
                raise ValueError('align must be left, center or right')
        elif key == 'line_width_pt':
            if type(value) not in (int, float) or not math.isfinite(value) or not .25 <= value <= 4:
                raise ValueError('line_width_pt must be between 0.25 and 4 pt')
        elif key == 'dash':
            if value not in DASHES:
                raise ValueError('dash must be solid, dash or dot')
        elif key == 'shadow':
            if value not in SHADOWS:
                raise ValueError('shadow must be none or soft')
        elif key == 'gradient':
            if value is not None and (not isinstance(value, dict) or set(value) - {'start', 'end', 'angle'} or not {'start', 'end'} <= set(value)):
                raise ValueError('gradient is {start, end, angle?} with two colours')
            if value is not None:
                resolve_color(value['start'], state)
                resolve_color(value['end'], state)
                angle = value.get('angle', 90)
                if type(angle) not in (int, float) or not 0 <= angle < 360:
                    raise ValueError('gradient angle must be between 0 and 359 degrees')
        elif key == 'series_colors':
            if not isinstance(value, list) or not 1 <= len(value) <= 6:
                raise ValueError('series_colors must list 1 to 6 colours')
            for item in value:
                resolve_color(item, state)
        elif key == 'gridlines':
            if value not in ('none', 'light'):
                raise ValueError('gridlines must be none or light')
        elif key == 'number_format':
            if value not in NUMBER_FORMATS:
                raise ValueError('number_format must be one of ' + ', '.join(NUMBER_FORMATS))
        elif key == 'emphasis':
            if not (value in ('none', 'max', 'min') or (isinstance(value, str) and 0 < len(value) <= 80) or type(value) is int):
                raise ValueError('emphasis must be none, max, min, a category label or a category index')
        elif key == 'legend':
            if value not in ('auto', 'none', 'bottom', 'right', 'top'):
                raise ValueError('legend must be auto, none, bottom, right or top')
        result[key] = value
    return result


def preset_style(name, kind):
    preset = PRESETS.get(name)
    if preset is None:
        raise ValueError(f'Unknown style preset {name!r}; available: ' + ', '.join(sorted(PRESETS)))
    if preset['kind'] != kind:
        raise ValueError(f'Preset {name} applies to {preset["kind"]} elements, not {kind}')
    return dict(preset['style'])


def effective_style(meta):
    """Preset first, then explicit overrides."""
    kind = style_kind(meta)
    style = {}
    if meta.get('preset') and meta['preset'] in PRESETS and PRESETS[meta['preset']]['kind'] == kind:
        style.update(PRESETS[meta['preset']]['style'])
    style.update(meta.get('style_overrides') or {})
    return style


# ---------------------------------------------------------------- application

def flat(shape):
    """No inherited theme shadow: empty effect list plus a neutral style reference (Impress honours only the latter)."""
    shape.shadow.inherit = False
    style = shape._element.find(qn('p:style'))
    if style is not None:
        reference = style.find(qn('a:effectRef'))
        if reference is not None:
            reference.set('idx', '0')

def _set_line(line, color, width_pt=None, dash=None):
    if color is None:
        line.fill.background()
        return
    line.color.rgb = RGBColor.from_string(color)
    if width_pt is not None:
        line.width = Pt(width_pt)
    if dash is not None:
        line.dash_style = DASHES[dash]


def _apply_runs(text_frame, color=None, bold=None, italic=None, size=None, family=None, align=None):
    for paragraph in text_frame.paragraphs:
        if align is not None:
            paragraph.alignment = ALIGN[align]
        for run in paragraph.runs:
            if color is not None:
                run.font.color.rgb = RGBColor.from_string(color)
            if bold is not None:
                run.font.bold = bold
            if italic is not None:
                run.font.italic = italic
            if size is not None:
                run.font.size = Pt(size)
            if family is not None:
                run.font.name = family


def text_style(meta, state):
    """Resolved text style with overrides merged; used before apply_text."""
    resolved = dict(meta.get('resolved_style') or {})
    style = effective_style(meta)
    if 'color' in style:
        resolved['color'] = resolve_color(style['color'], state)
    for key in ('bold', 'font_size', 'font_family', 'align'):
        if key in style:
            resolved[key] = style[key]
    if 'italic' in style:
        resolved['italic'] = style['italic']
    return resolved


def merge_overrides(current, additions):
    """Solid fill and gradient replace each other; everything else merges."""
    merged = dict(current)
    if 'fill' in additions:
        merged.pop('gradient', None)
    if 'gradient' in additions:
        merged.pop('fill', None)
    merged.update(additions)
    return merged


def _soft_shadow(shape, enabled):
    flat(shape)
    effects = shape._element.spPr.find(qn('a:effectLst'))
    for child in list(effects):
        effects.remove(child)
    if not enabled:
        return
    shadow = OxmlElement('a:outerShdw')
    for key, value in (('blurRad', '76200'), ('dist', '19050'), ('dir', '5400000'), ('algn', 't'), ('rotWithShape', '0')):
        shadow.set(key, value)
    color = OxmlElement('a:srgbClr')
    color.set('val', '000000')
    alpha = OxmlElement('a:alpha')
    alpha.set('val', '16000')
    color.append(alpha)
    shadow.append(color)
    effects.append(shadow)


def apply_shape(shape, meta, state):
    style = effective_style(meta)
    if 'gradient' in style and meta.get('shape') != 'line':
        gradient = style['gradient']
        start, end = resolve_color(gradient['start'], state), resolve_color(gradient['end'], state)
        shape.fill.gradient()
        shape.fill.gradient_angle = gradient.get('angle', 90)
        stops = shape.fill.gradient_stops
        stops[0].color.rgb = RGBColor.from_string(start)
        stops[0].position = 0
        stops[-1].color.rgb = RGBColor.from_string(end)
        stops[-1].position = 1
        # Contrast checks use the mid colour of the gradient.
        meta['fill'] = blend(start, end, .5)
    if 'shadow' in style and meta.get('shape') != 'line':
        _soft_shadow(shape, style['shadow'] == 'soft')
    if 'fill' in style and 'gradient' not in style and meta.get('shape') != 'line':
        color = resolve_color(style['fill'], state, allow_none=True)
        if color is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(color)
        meta['fill'] = color
    if 'line_color' in style or 'line_width_pt' in style or 'dash' in style:
        color = resolve_color(style.get('line_color', 'border'), state, allow_none=True)
        _set_line(shape.line, color, style.get('line_width_pt', 1.0 if color else None), style.get('dash'))
        meta['line'] = color
        if meta.get('shape') == 'line' and color:
            meta['fill'] = color
    if shape.has_text_frame and shape.text_frame.text and any(key in style for key in ('text_color', 'bold', 'font_size')):
        color = resolve_color(style['text_color'], state) if 'text_color' in style else None
        _apply_runs(shape.text_frame, color=color, bold=style.get('bold'), size=style.get('font_size'))
        if color:
            meta['text_color'] = color


def _cell_border(cell, side, color, width_pt=.75):
    tag = {'L': 'a:lnL', 'R': 'a:lnR', 'T': 'a:lnT', 'B': 'a:lnB'}[side]
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn(tag)):
        tc_pr.remove(old)
    line = OxmlElement(tag)
    if color is None:
        line.set('w', '0')
        line.append(OxmlElement('a:noFill'))
    else:
        line.set('w', str(int(width_pt * 12700)))
        solid = OxmlElement('a:solidFill')
        rgb = OxmlElement('a:srgbClr')
        rgb.set('val', color)
        solid.append(rgb)
        line.append(solid)
    return line


def _set_cell_borders(cell, left, right, top, bottom, width_pt=.75):
    tc_pr = cell._tc.get_or_add_tcPr()
    lines = [_cell_border(cell, side, color, width_pt) for side, color in (('L', left), ('R', right), ('T', top), ('B', bottom))]
    # CT_TableCellProperties requires lnL, lnR, lnT, lnB before any fill.
    for line in reversed(lines):
        tc_pr.insert(0, line)


def apply_table(shape, meta, state, profile_tables):
    """Clean default: soft header, thin horizontal dividers, padded cells."""
    style = {**PRESETS['table.clean']['style'],
             'header_fill': profile_tables.get('header_fill_token', 'card'),
             'header_text': profile_tables.get('header_text_token', 'text'),
             'zebra': bool(profile_tables.get('zebra'))}
    style.update(effective_style(meta))
    header_fill = resolve_color(style['header_fill'], state, allow_none=True)
    header_text = resolve_color(style['header_text'], state)
    body_fill = resolve_color(style.get('body_fill', 'background'), state, allow_none=True)
    body_text = resolve_color(style.get('body_text', 'text'), state)
    divider = resolve_color(style.get('divider_color', 'border'), state, allow_none=True)
    zebra_fill = resolve_color('surface', state) if style.get('zebra') else None
    size = style.get('font_size', profile_tables.get('font_size', 14))
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    rows = list(table.rows)
    for r, row in enumerate(rows):
        for cell in row.cells:
            fill = header_fill if r == 0 else (zebra_fill if zebra_fill and r % 2 == 0 else body_fill)
            if fill is None:
                cell.fill.background()
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            cell.margin_left = cell.margin_right = Inches(.08)
            cell.margin_top = cell.margin_bottom = Inches(.04)
            _set_cell_borders(cell, None, None, None, divider, 1.0 if r == 0 else .5)
            _apply_runs(cell.text_frame, color=header_text if r == 0 else body_text,
                bold=bool(style.get('header_bold', True)) if r == 0 else False, size=size)
    meta['table_style'] = dict(header_fill=header_fill, header_text=header_text, body_fill=body_fill, divider=divider)


def _point_emphasis(chart_data, target):
    values = chart_data['series'][0]['values']
    if target in (None, 'none') or not values:
        return None
    if target == 'max':
        return max(range(len(values)), key=lambda i: values[i])
    if target == 'min':
        return min(range(len(values)), key=lambda i: values[i])
    if type(target) is int:
        if not 0 <= target < len(values):
            raise ValueError('emphasis index is outside the chart categories')
        return target
    if target in chart_data['categories']:
        return chart_data['categories'].index(target)
    raise ValueError(f'emphasis category {target!r} is not in the chart')


def apply_chart(shape, meta, state, profile):
    """Profile defaults with explicit overrides: light grid, explicit points and units."""
    charts = profile.get('charts') or {}
    palette = state['theme']
    chart_data = meta.get('chart') or {}
    kind = chart_data.get('type', 'bar')
    style = effective_style(meta)
    default_grid = 'light' if charts.get('gridlines') else 'none'
    gridlines = style.get('gridlines', default_grid)
    categories = len(chart_data.get('categories') or [])
    single = len(chart_data.get('series') or []) == 1
    labels = style.get('data_labels', kind == 'pie' or (gridlines == 'none' and single and categories <= 8))
    if style.get('series_colors'):
        series_colors = [resolve_color(c, state) for c in style['series_colors']]
    elif state.get('effective_palette'):
        accent, muted = palette.get('accent', '246BCE'), palette.get('muted', '536B82')
        series_colors = [accent, muted, blend(accent, muted, 0.45), blend(accent, palette.get('background', 'FFFFFF'), 0.35)]
    else:
        series_colors = list(profile['colors']['chart_series'])
    font_size = style.get('font_size', charts.get('label_size', 12))
    font_family = state.get('font_family')
    muted = palette.get('muted', profile['colors']['muted'])
    border = palette.get('border', profile['colors']['border'])
    chart = shape.chart
    plot = chart.plots[0]
    emphasis = _point_emphasis(chart_data, style.get('emphasis')) if single and kind != 'pie' else None
    line_width = style.get('line_width_pt', charts.get('line_width_pt', 2.5))
    if kind == 'bar':
        try:
            plot.gap_width = int(charts.get('gap_width', 60))
            plot.overlap = 0 if single else -10
        except (AttributeError, ValueError):
            pass
    for index, series in enumerate(chart.series):
        color = series_colors[index % len(series_colors)]
        if kind == 'pie':
            for point_index in range(categories):
                point = series.points[point_index]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(series_colors[point_index % len(series_colors)])
                point.format.line.color.rgb = RGBColor.from_string(palette['background'])
            continue
        if kind == 'line':
            series.format.line.color.rgb = RGBColor.from_string(color)
            series.format.line.width = Pt(line_width)
            series.smooth = False
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 7
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.marker.format.line.color.rgb = RGBColor.from_string(color)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.format.line.fill.background()
        if emphasis is not None:
            soft = blend(color, palette['background'], .55)
            for point_index in range(categories):
                point = series.points[point_index]
                target = color if point_index == emphasis else soft
                if kind == 'line':
                    point.marker.style = XL_MARKER_STYLE.CIRCLE
                    point.marker.size = 10 if point_index == emphasis else 6
                    point.marker.format.fill.solid()
                    point.marker.format.fill.fore_color.rgb = RGBColor.from_string(target)
                    point.marker.format.line.color.rgb = RGBColor.from_string(target)
                else:
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = RGBColor.from_string(target)
    chart.has_title = False
    legend = style.get('legend', 'auto')
    if legend == 'auto':
        legend = 'none' if (single and kind != 'pie') or charts.get('legend_position') == 'none' else charts.get('legend_position', 'bottom')
    chart.has_legend = legend != 'none'
    if chart.has_legend:
        chart.legend.position = {'bottom': XL_LEGEND_POSITION.BOTTOM, 'right': XL_LEGEND_POSITION.RIGHT, 'top': XL_LEGEND_POSITION.TOP}[legend]
        chart.legend.include_in_layout = False
        chart.legend.font.name = font_family
        chart.legend.font.size = Pt(font_size)
        chart.legend.font.color.rgb = RGBColor.from_string(muted)
    plot.has_data_labels = bool(labels)
    if labels:
        data_labels = plot.data_labels
        data_labels.font.name = font_family
        data_labels.font.size = Pt(font_size)
        data_labels.font.color.rgb = RGBColor.from_string(palette['text'] if kind != 'pie' else palette['background'])
        number_format = style.get('number_format')
        if kind == 'pie':
            data_labels.show_percentage = True
            data_labels.show_value = False
            data_labels.number_format = number_format or '0%'
            data_labels.number_format_is_linked = False
            data_labels.position = XL_LABEL_POSITION.CENTER
        else:
            data_labels.show_value = True
            if number_format:
                data_labels.number_format = number_format
                data_labels.number_format_is_linked = False
            data_labels.position = XL_LABEL_POSITION.OUTSIDE_END if kind == 'bar' else XL_LABEL_POSITION.ABOVE
    if kind != 'pie':
        unit = chart_data.get('unit')
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.name = font_family
            axis.tick_labels.font.size = Pt(font_size)
            axis.tick_labels.font.color.rgb = RGBColor.from_string(muted)
            axis.format.line.color.rgb = RGBColor.from_string(border)
            axis.format.line.width = Pt(.75)
        chart.category_axis.has_major_gridlines = False
        chart.value_axis.has_major_gridlines = gridlines == 'light'
        chart.value_axis.has_minor_gridlines = False
        if gridlines == 'light':
            grid = chart.value_axis.major_gridlines.format.line
            grid.color.rgb = RGBColor.from_string(blend(border, palette['background'], .35))
            grid.width = Pt(.5)
        if style.get('number_format'):
            chart.value_axis.tick_labels.number_format = style['number_format']
            chart.value_axis.tick_labels.number_format_is_linked = False
        if unit:
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = str(unit)
            _apply_runs(chart.value_axis.axis_title.text_frame, color=muted, size=font_size, family=font_family)
        else:
            chart.value_axis.has_title = False
    used = categories if kind == 'pie' else len(chart_data.get('series') or [])
    meta['chart_style'] = dict(gridlines=gridlines, data_labels=bool(labels), emphasis=emphasis, legend=legend,
        series_colors=[series_colors[i % len(series_colors)] for i in range(used)])


# ---------------------------------------------------------------- geometry

def adjusted_geometry(current, request, slide_width_emu, slide_height_emu, keep_aspect):
    """Relative move/scale around the element centre in normalised coordinates."""
    x, y, w, h = current
    allowed = {'dx', 'dy', 'dx_in', 'dy_in', 'scale', 'scale_x', 'scale_y', 'rotate'}
    unknown = set(request) - allowed
    if unknown:
        raise ValueError('Unsupported adjust field(s): ' + ', '.join(sorted(unknown)))
    if not set(request) & allowed:
        raise ValueError('element.adjust needs dx/dy, dx_in/dy_in, scale or rotate')
    for key, value in request.items():
        if type(value) not in (int, float) or not math.isfinite(value):
            raise ValueError(f'{key} must be a finite number')
    width_in, height_in = slide_width_emu / 914400, slide_height_emu / 914400
    dx = request.get('dx', 0) + request.get('dx_in', 0) / width_in
    dy = request.get('dy', 0) + request.get('dy_in', 0) / height_in
    if abs(dx) > 1 or abs(dy) > 1:
        raise ValueError('Relative move exceeds the slide size')
    scale = request.get('scale')
    scale_x = request.get('scale_x', scale if scale is not None else 1)
    scale_y = request.get('scale_y', scale if scale is not None else 1)
    if keep_aspect and abs(scale_x - scale_y) > 1e-9:
        raise ValueError('This element keeps its aspect ratio; use one uniform scale value')
    for value in (scale_x, scale_y):
        if not .2 <= value <= 5:
            raise ValueError('scale must be between 0.2 and 5')
    cx, cy = x + w / 2 + dx, y + h / 2 + dy
    nw, nh = w * scale_x, h * scale_y
    nx, ny = cx - nw / 2, cy - nh / 2
    if nx < -1e-6 or ny < -1e-6 or nx + nw > 1 + 1e-6 or ny + nh > 1 + 1e-6:
        raise ValueError('The adjusted element would leave the slide canvas; reduce the adjustment instead of reflowing the page')
    return dict(x=max(0.0, nx), y=max(0.0, ny), w=nw, h=nh)


def rotation(current, delta):
    if type(delta) not in (int, float) or not math.isfinite(delta) or abs(delta) > 360:
        raise ValueError('rotate must be a relative angle between -360 and 360 degrees')
    return round((current + delta) % 360, 3)


# ---------------------------------------------------------------- chart replacement

def chart_payload(chart):
    data = CategoryChartData()
    data.categories = chart['categories']
    for item in chart['series']:
        data.add_series(item['name'], item['values'])
    return data


def verify_workbook(shape, chart):
    """The embedded workbook must hold exactly the logical chart data."""
    from openpyxl import load_workbook
    try:
        blob = shape.chart_part.chart_workbook.xlsx_part.blob
    except AttributeError as exc:
        raise ValueError('Replacement chart has no embedded workbook') from exc
    sheet = load_workbook(io.BytesIO(blob), read_only=True, data_only=True).worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if len(rows) < len(chart['categories']) + 1:
        raise ValueError('Embedded workbook lost chart categories')
    header = rows[0]
    for index, series in enumerate(chart['series'], 1):
        if header[index] != series['name']:
            raise ValueError('Embedded workbook series names do not match the chart')
        for row, value in enumerate(series['values'], 1):
            if rows[row][0] != chart['categories'][row - 1] or not math.isclose(float(rows[row][index]), float(value), rel_tol=1e-9, abs_tol=1e-9):
                raise ValueError('Embedded workbook values do not match the chart')
    plotted = [tuple(series.values) for series in shape.chart.series]
    if plotted != [tuple(float(v) for v in series['values']) for series in chart['series']]:
        raise ValueError('Chart cache values do not match the logical data')
    return True


def replace_chart(slide, shape, new_type, meta):
    """Swap the chart object in place; same position, z-order, name and data.

    Returns the replacement shape and its updated metadata; the caller stores it.
    """
    meta = dict(meta)
    chart = dict(meta.get('chart') or {})
    if new_type not in CHART_TYPES:
        raise ValueError('Supported chart types are bar, line and pie')
    if chart.get('type', 'bar') == new_type:
        raise ValueError(f'Chart is already a {new_type} chart')
    if new_type == 'pie':
        if len(chart['series']) != 1:
            raise ValueError('A pie chart shows exactly one series; keep bar/line or split the data')
        if any(value < 0 for value in chart['series'][0]['values']):
            raise ValueError('A pie chart cannot show negative values')
        if sum(chart['series'][0]['values']) <= 0:
            raise ValueError('A pie chart needs a positive total')
    chart['type'] = new_type
    left, top, width, height, name = shape.left, shape.top, shape.width, shape.height, shape.name
    parent = shape._element.getparent()
    index = parent.index(shape._element)
    old_rid = shape._element.chart_rId
    replacement = slide.shapes.add_chart(CHART_TYPES[new_type], left, top, width, height, chart_payload(chart))
    parent.remove(replacement._element)
    parent.insert(index, replacement._element)
    parent.remove(shape._element)
    try:
        slide.part.drop_rel(old_rid)
    except KeyError:
        pass
    replacement.name = name
    meta['chart'] = chart
    # Type-specific overrides from the former chart stop being meaningful.
    overrides = dict(meta.get('style_overrides') or {})
    if new_type == 'pie':
        overrides.pop('emphasis', None)
        overrides.pop('gridlines', None)
    if overrides:
        meta['style_overrides'] = overrides
    else:
        meta.pop('style_overrides', None)
    if meta.get('preset') and new_type == 'pie' and meta['preset'] == 'chart.emphasis':
        meta.pop('preset')
    verify_workbook(replacement, chart)
    return replacement, meta
