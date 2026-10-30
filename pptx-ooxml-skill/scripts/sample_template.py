"""Build the original consulting-brief deck used to exercise template reuse.

This file is project content. It is not a McKinsey, SlidesCarnival or Canva
template, and nothing here downloads or vendors a third-party deck.
"""
import importlib.util
import json
from pathlib import Path

from lxml import etree

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

def _importer():
    path = Path(__file__).with_name('template_importer.py')
    spec = importlib.util.spec_from_file_location('consulting_brief_importer', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


analyze = _importer().analyze

ACCENT = '0B2545'
TEXT = '1B2430'
MUTED = '5C6B7A'
RUST = 'C45C26'
ROLES = ('cover', 'section', 'comparison', 'data')


def build(directory):
    """Write ``template.pptx``, ``manifest.json``, ``tokens.json`` and ``LICENSE.txt``."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    _cover(prs, blank)
    _section(prs, blank)
    _comparison(prs, blank)
    _data(prs, blank, directory)
    _set_theme(prs)
    package = directory / 'template.pptx'
    prs.save(package)
    (directory / '_photo.png').unlink(missing_ok=True)
    manifest = analyze(Presentation(str(package)), 'consulting-brief', source=str(package))
    manifest['template_id'] = 'consulting-brief'
    manifest['license'] = dict(
        status='recorded',
        redistributable=True,
        note='Original project template created for technical verification. Not a third-party design.',
    )
    for slide, role in zip(manifest['slides'], ROLES):
        slide['role'] = role
    by_id = {slide['slide_id']: slide for slide in manifest['slides']}
    by_id['tpl-004']['slots'] = {'title': 'title', 'chart': 'chart', 'picture': 'photo', 'footnote': 'footnote'}
    manifest['slots_by_slide'] = {slide['slide_id']: slide.get('slots') or {} for slide in manifest['slides']}
    (directory / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    (directory / 'tokens.json').write_text(json.dumps(manifest['tokens'], ensure_ascii=False, indent=2), encoding='utf-8')
    (directory / 'LICENSE.txt').write_text(
        'consulting-brief is original content of this repository.\n'
        'It may be stored and reused inside this project.\n'
        'Third-party templates still need their own redistribution license before they are added here.\n',
        encoding='utf-8')
    return package


def _slide(prs, layout, slide_id):
    slide = prs.slides.add_slide(layout)
    slide._element.cSld.set('name', 'da-slide:' + slide_id)
    return slide


def _text(slide, name, text, x, y, w, h, size, bold=False, color=TEXT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = 'da:' + name
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = 'Liberation Sans'
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


def _cover(prs, layout):
    slide = _slide(prs, layout, 'tpl-001')
    logo = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(11.6), Inches(0.35), Inches(0.7), Inches(0.7))
    logo.name = 'da:logo'
    _fill(logo, ACCENT)
    _text(slide, 'title', '业务概览', 0.8, 1.8, 8.5, 1.1, 40, bold=True)
    subtitle = _text(slide, 'subtitle', '咨询简报', 0.8, 3.1, 8.0, 0.6, 20, bold=True, color=MUTED)
    extra = subtitle.text_frame.paragraphs[0].add_run()
    extra.text = ' 保留'
    extra.font.size = Pt(20)
    extra.font.italic = True
    extra.font.color.rgb = RGBColor.from_string(MUTED)


def _section(prs, layout):
    slide = _slide(prs, layout, 'tpl-002')
    _text(slide, 'title', '章节', 0.8, 2.6, 10, 1.2, 40, bold=True)


def _comparison(prs, layout):
    slide = _slide(prs, layout, 'tpl-003')
    _text(slide, 'title', '方案对比', 0.6, 0.4, 10, 0.8, 28, bold=True)
    _text(slide, 'left_title', '现状', 0.6, 1.6, 5.2, 0.6, 20, bold=True, color=ACCENT)
    _text(slide, 'left_body', '分散的执行记录', 0.6, 2.3, 5.2, 1.6, 16)
    _text(slide, 'right_title', '目标', 6.8, 1.6, 5.2, 0.6, 20, bold=True, color=ACCENT)
    _text(slide, 'right_body', '沿用同一套版式', 6.8, 2.3, 5.2, 1.6, 16)


def _data(prs, layout, directory):
    slide = _slide(prs, layout, 'tpl-004')
    _text(slide, 'title', '业绩', 0.6, 0.35, 8, 0.7, 28, bold=True)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.6), Inches(1.2), Inches(0.18), Inches(4.8))
    band.name = 'da:band'
    _fill(band, ACCENT)
    data = CategoryChartData()
    data.categories = ['Q1', 'Q2', 'Q3']
    data.add_series('Revenue', (10, 20, 30))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.1), Inches(1.4), Inches(6.2), Inches(4.4), data)
    chart.name = 'da:chart'
    image = directory / '_photo.png'
    Image.new('RGB', (200, 100), '#' + ACCENT).save(image)
    photo = slide.shapes.add_picture(str(image), Inches(8.0), Inches(1.6), Inches(4.0), Inches(2.0))
    photo.name = 'da:photo'
    note = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8.0), Inches(4.0), Inches(1.2), Inches(0.4))
    note.name = 'da:note'
    _fill(note, RUST)
    _text(slide, 'footnote', '注', 8.0, 4.7, 1.3, 0.28, 12, color=MUTED)


def _set_theme(prs):
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    theme = etree.fromstring(part.blob)
    scheme = theme.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
    wanted = {'accent1': ACCENT, 'dk1': TEXT, 'lt1': 'FFFFFF', 'dk2': MUTED, 'accent2': MUTED}
    namespace = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    for child in scheme:
        color = wanted.get(child.tag.rsplit('}', 1)[-1])
        if not color:
            continue
        for node in list(child):
            child.remove(node)
        child.append(etree.Element(namespace + 'srgbClr', val=color))
    part.blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)


if __name__ == '__main__':
    destination = Path(__file__).resolve().parents[1] / 'templates' / 'consulting-brief'
    print(build(destination))
