"""Copy one slide from a read-only template package into a working PPTX.

Same-package ``clone_slide`` cannot attach another PPTX's parts. This module
copies blobs into the destination package, reuses masters/layouts/themes from
the same template, duplicates charts and workbooks, and remaps internal slide
jumps only when the target has already been imported.
"""
import copy
import importlib.util
from pathlib import Path

from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import XmlPart

_CLONE = None
MASTER_ID_START = 2147483648
REFUSE_SUBSTRINGS = (
    'oleObject', 'diagramData', 'diagramLayout', 'diagramColors',
    'diagramQuickStyle', 'diagramDrawing',
)


class UnresolvedInternalLink(ValueError):
    """A template slide jump has no unique imported destination."""


def _clone():
    global _CLONE
    if _CLONE is None:
        path = Path(__file__).with_name('template_clone.py')
        spec = importlib.util.spec_from_file_location('pptx_import_clone', path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        _CLONE = module
    return _CLONE


last_import_report = None


def import_slide(dest_prs, source_prs, source, new_slide_id, template_id, source_to_dest, imported_parts, identity):
    """Append a native copy of ``source`` from another package and return ``(slide, report)``.

    ``imported_parts`` is mutated: ``{template_id: {source_partname: dest_partname}}``.
    ``source_to_dest`` maps a template slide id to the dest slide ids already imported
    from that template page (a list; a unique match is required to retarget a jump).
    """
    global last_import_report
    clone = _clone()
    if not isinstance(new_slide_id, str) or not new_slide_id:
        raise ValueError('slide.import requires a new slide_id')
    if dest_prs.slide_width != source_prs.slide_width or dest_prs.slide_height != source_prs.slide_height:
        raise ValueError('Slide size is incompatible with the working PPTX; refusing to stretch template content')
    _assert_importable(source, clone)
    report = dict(shared=[], duplicated=[], omitted=[], external=[], theme_conflict=None)
    existing = {tid for tid, mapping in (imported_parts or {}).items() if mapping}
    if existing and template_id not in existing:
        report['theme_conflict'] = dict(
            existing=sorted(existing), incoming=template_id,
            message='Importing from a different template series; theme and master will not be shared.')
    stored = imported_parts.setdefault(template_id, {})
    ctx = _Context(dest_prs, source_prs, identity, source_to_dest or {}, stored, report, clone)
    dest_layout = _ensure_layout(ctx, source.slide_layout)
    dest = dest_prs.slides.add_slide(dest_layout)
    clone._clear_layout_shapes(dest)
    for shape in source.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shape._element))
    clone._copy_slide_extras(source, dest)
    mapping = _copy_slide_relationships(ctx, source, dest)
    last_import_report = report
    missing = clone._relationship_ids(dest._element) - set(mapping)
    dest_ids = {rel.rId for rel in dest.part.rels.values()}
    unresolved = {item for item in missing if item not in dest_ids}
    if unresolved:
        raise ValueError('Slide import left unresolved relationships: ' + ', '.join(sorted(unresolved)))
    clone._rewrite_rel_ids(dest._element, mapping)
    dest._element.cSld.set('name', 'da-slide:' + new_slide_id)
    clone.assert_internal_jumps(dest_prs, identity)
    return dest, report


class _Context:
    def __init__(self, dest_prs, source_prs, identity, source_to_dest, stored, report, clone):
        self.dest_prs = dest_prs
        self.source_prs = source_prs
        self.identity = identity
        self.source_to_dest = source_to_dest
        self.stored = stored
        self.report = report
        self.clone = clone
        self.dest_package = dest_prs.part.package
        self.reserved = set()
        self.memo = _seed_memo(self.dest_package, stored)
        self.persist_types = clone.SHARE - {RT.SLIDE}


def _seed_memo(dest_package, stored):
    by_name = {str(part.partname): part for part in dest_package.iter_parts()}
    memo = {}
    for source_name, dest_name in list(stored.items()):
        part = by_name.get(dest_name)
        if part is None:
            stored.pop(source_name, None)
            continue
        memo[source_name] = part
    return memo


def _ensure_layout(ctx, source_layout):
    dest_master = copy_part_into(ctx, source_layout.slide_master.part, persist=True)
    dest_layout_part = copy_part_into(ctx, source_layout.part, persist=True)
    _register_master(ctx.dest_prs, dest_master)
    return dest_layout_part.slide_layout


def _register_master(prs, master_part):
    for rel in prs.part.rels.values():
        if not rel.is_external and rel.reltype == RT.SLIDE_MASTER and rel.target_part is master_part:
            return
    r_id = prs.part.relate_to(master_part, RT.SLIDE_MASTER)
    listing = prs.part._element.get_or_add_sldMasterIdLst()
    used = [int(child.get('id') or 0) for child in listing]
    entry = listing._add_sldMasterId(rId=r_id)
    entry.set('id', str(max([MASTER_ID_START - 1] + used) + 1))


def copy_part_into(ctx, source_part, persist=False):
    """Create a dest-package part whose blob is a copy of ``source_part``."""
    key = str(source_part.partname)
    if key in ctx.memo:
        return ctx.memo[key]
    clone = ctx.clone
    name = clone._allocate(ctx.dest_package, source_part.partname, ctx.reserved)
    blob = bytes(source_part.blob)
    try:
        if isinstance(source_part, XmlPart):
            created = type(source_part).load(name, source_part.content_type, ctx.dest_package, blob)
        else:
            created = type(source_part)(name, source_part.content_type, ctx.dest_package, blob)
    except Exception as exc:
        raise ValueError(f'Cannot copy package part {source_part.partname} into the working PPTX') from exc
    ctx.memo[key] = created
    if persist:
        ctx.stored[key] = str(created.partname)
        ctx.report['shared'].append(str(created.partname))
    else:
        ctx.report['duplicated'].append(str(created.partname))
    mapping = {}
    for rel in list(source_part.rels.values()):
        mapped = _map_relationship(ctx, rel, created, on_slide=False)
        if mapped is not None:
            mapping[rel.rId] = mapped
    if isinstance(created, XmlPart):
        clone._rewrite_rel_ids(created._element, mapping)
    return created


def _copy_slide_relationships(ctx, source, dest):
    mapping = {}
    for rel in list(source.part.rels.values()):
        mapped = _map_relationship(ctx, rel, dest.part, on_slide=True)
        if mapped is not None:
            mapping[rel.rId] = mapped
    return mapping


def _map_relationship(ctx, rel, dest_part, on_slide):
    clone = ctx.clone
    if rel.is_external:
        ctx.report['external'].append(rel.reltype.rsplit('/', 1)[-1])
        return dest_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
    if on_slide and rel.reltype in clone.SKIP_ON_SLIDE:
        ctx.report['omitted'].append(rel.reltype.rsplit('/', 1)[-1])
        return None
    if rel.reltype == RT.SLIDE:
        return dest_part.relate_to(_resolve_jump(ctx, rel.target_part), RT.SLIDE)
    action = clone._policy(rel, on_slide=on_slide)
    if action == 'skip':
        ctx.report['omitted'].append(rel.reltype.rsplit('/', 1)[-1])
        return None
    persist = rel.reltype in ctx.persist_types
    if action == 'share' and rel.target_part.package is ctx.dest_package:
        ctx.report['shared'].append(rel.reltype.rsplit('/', 1)[-1])
        target = rel.target_part
    else:
        target = copy_part_into(ctx, rel.target_part, persist=persist)
    return dest_part.relate_to(target, rel.reltype)


def _resolve_jump(ctx, source_part):
    source_slide = next((slide for slide in ctx.source_prs.slides if slide.part is source_part), None)
    if source_slide is None:
        raise UnresolvedInternalLink('unresolved_internal_link: jump target is not a slide in the template')
    source_id = ctx.identity(source_slide)
    dest_ids = list(ctx.source_to_dest.get(source_id) or [])
    if len(dest_ids) != 1:
        raise UnresolvedInternalLink(
            'unresolved_internal_link: template slide ' + source_id + ' is not uniquely imported yet')
    dest_id = dest_ids[0]
    dest_slide = next((slide for slide in ctx.dest_prs.slides if ctx.identity(slide) == dest_id), None)
    if dest_slide is None:
        raise UnresolvedInternalLink('unresolved_internal_link: imported destination is missing for ' + dest_id)
    return dest_slide.part


def _assert_importable(slide, clone):
    seen = set()

    def walk(part):
        if part in seen:
            return
        seen.add(part)
        for rel in list(part.rels.values()):
            reltype = str(rel.reltype)
            if any(token in reltype for token in REFUSE_SUBSTRINGS):
                raise ValueError(
                    'Cannot import slide: unsupported object ' + reltype.rsplit('/', 1)[-1]
                    + '; refusing to convert it to a picture')
            if rel.is_external:
                continue
            if rel.reltype in clone.SKIP_ON_SLIDE:
                continue
            clone._policy(rel, on_slide=False)
            walk(rel.target_part)

    walk(slide.part)
