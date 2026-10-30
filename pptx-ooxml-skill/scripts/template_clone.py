"""Clone a slide without aliasing charts, embeddings or other mutable parts.

Sharing an image part is safe. Sharing a chart or its workbook is not: a later
``chart.replace_data`` would overwrite the source slide. Layouts and masters stay
shared so the clone keeps the template's theme and placeholders.
"""
import copy
import re

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import XmlPart
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CHART_STYLE = 'http://schemas.microsoft.com/office/2011/relationships/chartStyle'
# Shared targets stay one part. A slide jump must keep pointing at the existing slide, not spawn another.
SHARE = {
    RT.IMAGE, RT.SLIDE, RT.SLIDE_LAYOUT, RT.SLIDE_MASTER, RT.THEME, RT.AUDIO, RT.VIDEO, RT.MEDIA,
    RT.CHART_COLOR_STYLE, CHART_STYLE,
}
# Mutable or slide-specific parts get their own copy so later edits do not alias the source.
DUPLICATE = {
    RT.CHART, RT.PACKAGE, RT.OLE_OBJECT, RT.TAGS, RT.CHART_USER_SHAPES,
    RT.DIAGRAM_COLORS, RT.DIAGRAM_DATA, RT.DIAGRAM_LAYOUT,
    getattr(RT, 'DIAGRAM_STYLE', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramStyle'),
}
SKIP_ON_SLIDE = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}
last_clone_report = None


def clone_slide(prs, source, new_slide_id):
    """Append a structural copy of ``source`` and return the new slide.

    The caller inserts it at the requested index and records stable state. The
    new slide's ``cSld/@name`` is ``da-slide:{new_slide_id}``. Shape names are
    copied, so slot ids such as ``title`` still resolve on the clone.
    """
    global last_clone_report
    if not isinstance(new_slide_id, str) or not new_slide_id:
        raise ValueError('slide.clone requires a new slide_id')
    report = dict(shared=[], duplicated=[], omitted=[], external=[])
    dest = prs.slides.add_slide(source.slide_layout)
    _clear_layout_shapes(dest)
    for shape in source.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shape._element))
    _copy_slide_extras(source, dest)
    mapping = _clone_relationships(source, dest, report)
    last_clone_report = report
    missing = _relationship_ids(dest._element) - set(mapping)
    # Ids that already belonged to the destination (its layout relationship) are not in the map.
    dest_ids = {rel.rId for rel in dest.part.rels.values()}
    unresolved = {item for item in missing if item not in dest_ids}
    if unresolved:
        raise ValueError('Slide clone left unresolved relationships: ' + ', '.join(sorted(unresolved)))
    _rewrite_rel_ids(dest._element, mapping)
    dest._element.cSld.set('name', 'da-slide:' + new_slide_id)
    return dest


def _clear_layout_shapes(slide):
    tree = slide.shapes._spTree
    for child in list(tree):
        if etree.QName(child).localname in ('nvGrpSpPr', 'grpSpPr'):
            continue
        tree.remove(child)


def _copy_slide_extras(source, dest):
    for local in ('clrMapOvr', 'transition', 'timing'):
        node = source._element.find(qn(f'p:{local}'))
        if node is None:
            continue
        existing = dest._element.find(qn(f'p:{local}'))
        if existing is not None:
            dest._element.remove(existing)
        dest._element.append(copy.deepcopy(node))
    background = source._element.cSld.find(qn('p:bg'))
    if background is None:
        return
    existing = dest._element.cSld.find(qn('p:bg'))
    if existing is not None:
        dest._element.cSld.remove(existing)
    dest.shapes._spTree.addprevious(copy.deepcopy(background))


def _clone_relationships(source, dest, report):
    mapping = {}
    memo = {}
    reserved = set()
    package = source.part.package
    for rel in list(source.part.rels.values()):
        action = _policy(rel)
        if action == 'skip':
            report['omitted'].append(rel.reltype.rsplit('/', 1)[-1])
            continue
        if action == 'external':
            mapping[rel.rId] = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            report['external'].append(rel.reltype.rsplit('/', 1)[-1])
            continue
        target = rel.target_part
        if action == 'duplicate':
            target = _duplicate_part(package, target, memo, reserved)
            report['duplicated'].append(str(target.partname))
        else:
            report['shared'].append(rel.reltype.rsplit('/', 1)[-1])
        mapping[rel.rId] = dest.part.relate_to(target, rel.reltype)
    return mapping


def _policy(rel):
    if rel.reltype in SKIP_ON_SLIDE:
        return 'skip'
    if rel.is_external:
        return 'external'
    if rel.reltype in SHARE:
        return 'share'
    if rel.reltype in DUPLICATE:
        return 'duplicate'
    raise ValueError('No clone policy for relationship ' + rel.reltype.rsplit('/', 1)[-1] + '; refusing to copy it blindly')


def remove_slide(prs, slide):
    """Drop one slide from the presentation. Unreferenced parts disappear on the next save."""
    rel_id = None
    for rel in prs.part.rels.values():
        if not rel.is_external and rel.reltype == RT.SLIDE and rel.target_part is slide.part:
            rel_id = rel.rId
            break
    if rel_id is None:
        raise ValueError('Slide is not attached to the presentation')
    namespace = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    node = next((item for item in prs.slides._sldIdLst if item.get(namespace) == rel_id), None)
    if node is None:
        raise ValueError('Slide identifier is missing from the presentation')
    prs.slides._sldIdLst.remove(node)
    prs.part.drop_rel(rel_id)


def discard_prototypes(prs, prototype_ids, redirects, identity):
    """Retarget remaining internal slide jumps, then drop the prototype slides.

    ``redirects`` maps a discarded prototype id to a slide that will remain.
    Jumps that already point at a remaining slide are left alone. A jump to a
    prototype with no mapping, or to a missing slide, fails before anything is
    removed. Clone order is never used to guess a destination.
    """
    removing = list(prototype_ids or [])
    mapping = dict(redirects or {})
    by_id = {}
    for slide in prs.slides:
        sid = identity(slide)
        if sid in by_id:
            raise ValueError('Duplicate slide ID')
        by_id[sid] = slide
    missing = [sid for sid in removing if sid not in by_id]
    if missing:
        raise ValueError('Unknown prototype slide: ' + ', '.join(missing))
    for source, dest in mapping.items():
        if not isinstance(source, str) or not isinstance(dest, str):
            raise ValueError('redirects keys and values must be slide ids')
        if source not in removing:
            raise ValueError('redirects source is not a discarded prototype: ' + source)
        if dest not in by_id:
            raise ValueError('redirects target does not exist: ' + dest)
        if dest in removing:
            raise ValueError('redirects target is being discarded: ' + dest)
    redirected = retarget_internal_jumps(prs, set(removing), mapping, identity, by_id)
    for prototype in removing:
        remove_slide(prs, by_id[prototype])
    assert_internal_jumps(prs, identity)
    return dict(discarded=removing, redirected=redirected)


def retarget_internal_jumps(prs, removing, mapping, identity, by_id):
    """Rewrite RT.SLIDE relationships on slides that will remain."""
    redirected = []
    remaining = [slide for slide in prs.slides if identity(slide) not in removing]
    for slide in remaining:
        source_id = identity(slide)
        for rel in list(_internal_slide_rels(slide)):
            target = _slide_for_part(prs, rel.target_part)
            if target is None:
                raise ValueError('Internal slide jump points at a missing slide on ' + source_id)
            target_id = identity(target)
            if target_id not in removing:
                continue
            dest_id = mapping.get(target_id)
            if not dest_id:
                raise ValueError('Internal jump to prototype ' + target_id + ' has no redirect')
            dest = by_id[dest_id]
            old_rid = rel.rId
            new_rid = slide.part.relate_to(dest.part, RT.SLIDE)
            if new_rid != old_rid:
                _rewrite_rel_ids(slide._element, {old_rid: new_rid})
                _drop_unused_rel(slide.part, old_rid)
            redirected.append(dict(from_slide=source_id, prototype=target_id, to_slide=dest_id))
    return redirected


def assert_internal_jumps(prs, identity):
    """Fail if any remaining slide jump targets a part that is not in the deck."""
    known = {slide.part for slide in prs.slides}
    for slide in prs.slides:
        for rel in _internal_slide_rels(slide):
            if rel.target_part not in known:
                raise ValueError('Dangling internal slide jump on ' + identity(slide))


def _internal_slide_rels(slide):
    for rel in list(slide.part.rels.values()):
        if not rel.is_external and rel.reltype == RT.SLIDE:
            yield rel


def _slide_for_part(prs, part):
    for slide in prs.slides:
        if slide.part is part:
            return slide
    return None


def _drop_unused_rel(part, rId):
    if rId not in part.rels:
        return
    if rId in _relationship_ids(part._element):
        raise ValueError('Internal jump relationship ' + rId + ' is still referenced')
    part.rels.pop(rId)


def _duplicate_part(package, part, memo, reserved):
    key = str(part.partname)
    if key in memo:
        return memo[key]
    name = _allocate(package, part.partname, reserved)
    try:
        if isinstance(part, XmlPart):
            created = type(part).load(name, part.content_type, package, part.blob)
        else:
            created = type(part)(name, part.content_type, package, part.blob)
    except Exception as exc:
        raise ValueError(f'Cannot safely clone package part {part.partname}') from exc
    memo[key] = created
    mapping = {}
    for rel in list(part.rels.values()):
        if rel.is_external:
            mapping[rel.rId] = created.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            continue
        action = _policy(rel)
        if action == 'skip':
            continue
        target = rel.target_part
        if action == 'duplicate':
            target = _duplicate_part(package, target, memo, reserved)
        elif action == 'external':
            mapping[rel.rId] = created.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            continue
        mapping[rel.rId] = created.relate_to(target, rel.reltype)
    if isinstance(created, XmlPart):
        _rewrite_rel_ids(created._element, mapping)
    return created


def _allocate(package, partname, reserved):
    template = _numbered_template(str(partname))
    existing = {str(part.partname) for part in package.iter_parts()}
    existing.update(reserved)
    for number in range(1, 10001):
        candidate = template % number
        if candidate not in existing:
            reserved.add(candidate)
            return PackURI(candidate)
    raise ValueError('Could not allocate a package part name')


def _numbered_template(partname):
    replaced, count = re.subn(r'(\d+)(\.[A-Za-z0-9]+)$', r'%d\2', partname)
    if count == 1:
        return replaced
    stem, dot, ext = partname.rpartition('.')
    return stem + '%d.' + ext


def _relationship_ids(element):
    found = set()
    prefix = '{' + REL_NS + '}'
    for node in element.iter():
        for attr, value in node.attrib.items():
            if attr.startswith(prefix):
                found.add(value)
    return found


def _rewrite_rel_ids(element, mapping):
    prefix = '{' + REL_NS + '}'
    for node in element.iter():
        for attr, value in list(node.attrib.items()):
            if attr.startswith(prefix) and value in mapping:
                node.set(attr, mapping[value])
