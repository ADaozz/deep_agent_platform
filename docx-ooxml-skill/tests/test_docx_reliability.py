"""DOCX import, recolor reporting and clone boundaries."""
import importlib.util
import io
import zipfile
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'


def load(name):
    spec = importlib.util.spec_from_file_location('docx_rel_' + name, SCRIPTS / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


atomic = load('atomic_docx')
importer = load('template_importer')
sample = load('sample_template')


def apply(store, **operation):
    state = store.status()
    return store.apply_one(dict(document_id='report', operation_id=f'op-{state["revision"] + 1}',
        base_revision=state['revision'], operation=operation))


def test_plain_docx_is_indexed_without_rewriting_the_file(tmp_path):
    from docx import Document
    path = tmp_path / 'plain.docx'
    document = Document()
    document.add_heading('概述', level=1)
    document.add_paragraph('这是正文。')
    document.add_heading('计划', level=1)
    document.add_paragraph('下一步。')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '事项'
    table.cell(0, 1).text = '状态'
    document.save(path)
    original = path.read_bytes()
    manifest = importer.import_template(path, 'plain-report')
    assert path.read_bytes() == original
    assert [item['text'] for item in manifest['sections']] == ['概述', '计划']
    assert manifest['tables']
    assert manifest['blocks']


def test_unknown_cell_shading_is_not_a_complete_recolor(tmp_path):
    sample.build(tmp_path)
    store = atomic.DocumentStore(tmp_path / 'report.docx')
    store.init('report', tmp_path / 'template.docx')
    result = apply(store, type='document.recolor', palette={'accent': '707780', 'background': 'FFFFFF', 'text': '20242A', 'muted': '79818A'})
    assert result['recolor_complete'] is False
    shades = [item for item in result['unhandled'] if item['color'] == 'F2F4F7']
    assert shades and shades[0]['status'] == 'unmapped' and 'shd' in shades[0]['location']


def test_cell_text_replacement_and_refusals(tmp_path):
    path = _complex_docx(tmp_path / 'complex.docx')
    store = atomic.DocumentStore(tmp_path / 'report.docx')
    store.init('report', path)
    blocks = store.status()['blocks']
    table = next(block['block_id'] for block in blocks if block['type'] == 'tbl' and '旧甲' in block['text'])
    merged = next(block['block_id'] for block in blocks if block['type'] == 'tbl' and '合并' in block['text'])
    mixed = next(block['block_id'] for block in blocks if '混合样式' in block['text'])
    copied = apply(store, type='block.clone', source_block_id=table, block_id='tbl-copy',
        rows=[['事项', '状态'], ['新甲', '完成']])
    assert 'tbl-copy' in copied['changed_blocks']
    document = zipfile.ZipFile(store.path).read('word/document.xml').decode()
    assert document.count('旧乙') == 1
    assert document.count('新甲') == 1
    before = store.path.read_bytes()
    revision = store.status()['revision']
    with pytest.raises(ValueError, match='merged cells'):
        apply(store, type='block.clone', source_block_id=merged, block_id='tbl-merged', rows=[['合并', '右']])
    with pytest.raises(ValueError, match='mixed character styles'):
        apply(store, type='slot.fill', block_id=mixed, text='只保留一种样式')
    assert store.path.read_bytes() == before
    assert store.status()['revision'] == revision


def _complex_docx(path):
    office = load('office_docx')
    atomic_module = load('atomic_docx')
    w = office.W
    document = f'''<w:document xmlns:w="{w}"><w:body>
      <w:tbl><w:tblPr><w:tblW w:w="4000" w:type="dxa"/></w:tblPr>
        <w:tr><w:tc><w:p><w:r><w:t>事项</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>状态</w:t></w:r></w:p></w:tc></w:tr>
        <w:tr><w:tc><w:p><w:r><w:t>旧甲</w:t></w:r></w:p><w:p><w:r><w:t>旧乙</w:t></w:r></w:p></w:tc>
          <w:tc><w:p><w:r><w:t>待定</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
      <w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>合并</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
      <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>混合样式</w:t></w:r><w:r><w:t>普通</w:t></w:r></w:p>
      <w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>
    </w:body></w:document>'''
    parts = office.default_package()
    parts[office.MAIN] = document.encode()
    parts = atomic_module.mark_blocks(parts)
    office.save_parts(parts, path)
    return path
