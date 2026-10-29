"""Local regression tests for OOXML package and rendering gate."""
import hashlib
import importlib.util
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

BASE = Path(__file__).resolve().parents[1]
SCRIPT = BASE / 'scripts' / 'office_docx.py'
spec = importlib.util.spec_from_file_location('office_docx', SCRIPT)
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)
W = mod.W


def cli(*args):
    p = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True, text=True)
    assert p.returncode == 0, (p.stdout, p.stderr)
    return json.loads(p.stdout)


def write_plan(inp, path, ops):
    path.write_text(json.dumps({'source_sha256':mod.sha(inp),'operations':ops}, ensure_ascii=False), encoding='utf8')


def test_blank_insert_and_xml_roundtrip(tmp_path):
    inp, plan, out = (tmp_path / x for x in ('blank.docx','plan.json','out.docx'))
    cli('init','--output',inp)
    xml = f'<w:p xmlns:w="{W}"><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>复杂项目报告</w:t></w:r></w:p>'
    tbl = f'<w:tbl xmlns:w="{W}"><w:tblPr/><w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>表格内容</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
    write_plan(inp,plan,[{'op':'insert_blocks','after_body_index':-1,'blocks':[xml,tbl]}])
    cli('apply',inp,plan,'--output',out)
    info = cli('inspect',out)
    assert info['validation']['ok']
    assert info['body_blocks'][0]['text']=='复杂项目报告'
    assert info['body_blocks'][1]['type']=='table'
    assert '表格内容' in info['body_blocks'][1]['text']


def test_split_runs_precise_replacement_and_preservation(tmp_path):
    from docx import Document
    inp, plan, out = (tmp_path / x for x in ('template.docx','plan.json','modified.docx'))
    d = Document()
    p=d.add_paragraph()
    p.add_run('项目名称：{{')
    p.add_run('NAME').bold=True
    p.add_run('}} 计划不变')
    d.add_paragraph('第二段不变')
    d.save(inp)
    parts=mod.read_parts(inp)
    original_styles=parts['word/styles.xml']
    write_plan(inp,plan,[{'op':'replace_text','paragraph_index':0,'expected_paragraph':'项目名称：{{NAME}} 计划不变','old':'{{NAME}}','new':'深度智能体'}])
    cli('apply',inp,plan,'--output',out)
    assert mod.read_parts(out)['word/styles.xml']==original_styles
    res=cli('inspect',out)
    assert res['paragraphs'][0]['text']=='项目名称：深度智能体 计划不变'
    assert res['paragraphs'][1]['text']=='第二段不变'


def test_reject_bad_xml_and_hash_guard(tmp_path):
    inp, plan, out = (tmp_path / x for x in ('blank.docx','bad.json','bad.docx'))
    cli('init','--output',inp)
    write_plan(inp,plan,[{'op':'insert_blocks','after_body_index':-1,'blocks':['<w:p xmlns:w="'+W+'"><w:r></w:p>']}])
    p=subprocess.run([sys.executable,str(SCRIPT),'apply',str(inp),str(plan),'--output',str(out)],capture_output=True,text=True)
    assert p.returncode != 0
    assert not out.exists()
    plan.write_text(json.dumps({'source_sha256':'WRONG','operations':[{'op':'set_page','top_twips':1200}]}))
    p=subprocess.run([sys.executable,str(SCRIPT),'apply',str(inp),str(plan),'--output',str(out)],capture_output=True,text=True)
    assert p.returncode != 0


def test_image_and_finalization_gate(tmp_path):
    from PIL import Image
    inp, plan, out = (tmp_path / x for x in ('blank.docx','plan.json','result.docx'))
    imgdir=tmp_path/'assets'; imgdir.mkdir()
    Image.new('RGB',(320,160),(123,45,200)).save(imgdir/'diagram.png')
    cli('init','--output',inp)
    head=f'<w:p xmlns:w="{W}"><w:r><w:t>示例图片</w:t></w:r></w:p>'
    write_plan(inp,plan,[{'op':'insert_blocks','after_body_index':-1,'blocks':[head]},{'op':'insert_image','after_body_index':0,'path':'diagram.png','width_cm':7,'alt':'测试图'}])
    cli('apply',inp,plan,'--asset-root',imgdir,'--output',out)
    res=cli('validate',out)
    assert res['ok'],res
    assert len([p for p in mod.read_parts(out) if p.startswith('word/media/')])==1
    if not shutil_which('libreoffice'):
        pytest.skip('LibreOffice unavailable')
    rend=tmp_path/'renders'
    man=cli('render',out,'--output-dir',rend)
    assert man['page_count']>=1
    review=tmp_path/'review.json'
    cli('review-template','--render-dir',rend,'--output',review)
    p=subprocess.run([sys.executable,str(SCRIPT),'finalize',str(out),'--render-dir',str(rend),'--review',str(review),'--output',str(tmp_path/'final.docx')],capture_output=True,text=True)
    assert p.returncode != 0  # no rubber-stamping
    r=json.loads(review.read_text())
    r.update(reviewed_pages=list(range(1, man['page_count']+1)), passed=True, issues=[])
    review.write_text(json.dumps(r))
    cli('finalize',out,'--render-dir',rend,'--review',review,'--output',tmp_path/'final.docx')
    assert (tmp_path/'final.docx').exists()


def shutil_which(cmd):
    from shutil import which
    return which(cmd)


def test_tagged_sdt_edit_and_para_layout(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    inp,plan,out=(tmp_path/x for x in ('source.docx','plan.json','out.docx'))
    d=Document()
    p=d.add_paragraph('标题')
    sdt=OxmlElement('w:sdt')
    prop=OxmlElement('w:sdtPr')
    tag=OxmlElement('w:tag'); tag.set(qn('w:val'),'client_name'); prop.append(tag)
    content=OxmlElement('w:sdtContent')
    run=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='客户'; run.append(t)
    content.append(run)
    sdt.append(prop); sdt.append(content)
    p._p.append(sdt)
    d.save(inp)
    write_plan(inp, plan, [
        {'op':'replace_sdt_text','tag':'client_name','new':'某集团'},
        {'op':'set_paragraph_format','paragraph_index':0,'expected_text':'标题某集团','alignment':'center','space_after_twips':120,'keep_next':True},
    ])
    cli('apply',inp,plan,'--output',out)
    assert cli('inspect',out)['paragraphs'][0]['text']=='标题某集团'
    root=mod.parse_xml(mod.read_parts(out)[mod.MAIN])
    assert root.xpath('string(//w:jc/@w:val)',namespaces={'w':W})=='center'
