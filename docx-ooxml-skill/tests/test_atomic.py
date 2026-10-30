"""Real storage and layout tests, including recovery and optimistic concurrency."""
from concurrent.futures import ProcessPoolExecutor
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pytest

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
spec = importlib.util.spec_from_file_location('atomic', SCRIPTS / 'atomic_docx.py')
a = importlib.util.module_from_spec(spec)
spec.loader.exec_module(a)


def request(revision, operation_id=None, **operation):
    return dict(document_id='report', operation_id=operation_id or f'op-{revision+1}', base_revision=revision,
        operation=operation or dict(type='append_paragraph', block_id=f'p-{revision+1}', text=f'Paragraph {revision+1}'))


def fake_review(monkeypatch, pages=1, passed=True):
    # Inject deterministic PDF/render output, not a fake visual approval.
    # checkpoint loads the helper independently; replace the shared engine renderer
    def render(docx, output, dpi=125):
        import pymupdf as fitz
        doc = fitz.open()
        for i in range(pages):
            page = doc.new_page()
            page.insert_text((72, 72), f'Page {i+1} body')
        doc.save(output / 'preview.pdf')
        for i, page in enumerate(doc, 1):
            page.get_pixmap().save(output / f'page-{i:03d}.png')
        doc.close()
    monkeypatch.setattr(a.o, 'run_render', render)


def test_twenty_operations_idempotency_failure_and_restore(tmp_path, monkeypatch):
    fake_review(monkeypatch)
    store = a.DocumentStore(tmp_path / 'report.docx')
    store.init('report')
    for revision in range(20):
        if store.status()['gate'] != 'open':
            assert store.checkpoint()['gate'] == 'open'
        payload = request(revision)
        result = store.apply_one(payload)
        assert result['revision'] == revision + 1
        assert store.apply_one(payload) == result
        assert store.status()['revision'] == revision + 1
        before = store.path.read_bytes()
        with pytest.raises(ValueError):
            store.apply_one(request(revision + 1, operation_id=f'bad-{revision}', type='insert_block', block_id='bad', xml='<bad>'))
        assert store.path.read_bytes() == before
    state = store.status()
    assert len(state['blocks']) == 20
    assert len(state['history']) == 21  # init + 20 commits
    with pytest.raises(ValueError, match='Revision conflict'):
        store.apply_one(request(0, operation_id='stale-operation'))
    restored = store.restore(5, 20, 'restore-21')
    assert restored['revision'] == 21
    assert store.path.read_bytes() == store.snapshot(5)
    assert store.restore(5, 20, 'restore-21') == restored


def test_conflicts_payload_reuse_and_bounded_protocol(tmp_path):
    store = a.DocumentStore(tmp_path / 'report.docx'); store.init('report')
    result = store.apply_one(request(0))
    with pytest.raises(ValueError, match='Operation ID reused'):
        store.apply_one(request(0, text='different', type='append_paragraph', block_id='other'))
    with pytest.raises(ValueError, match='Revision conflict'):
        store.apply_one(request(0, operation_id='new'))
    for operation in [dict(type='insert_blocks', block_id='bulk', blocks=[]),
        dict(type='append_paragraph', block_id='bulk', text='x\ny'),
        dict(type='append_table', block_id='bulk', rows=[['x']] * 11)]:
        with pytest.raises(ValueError):
            store.apply_one(request(1, **operation))
    assert store.status()['revision'] == 1


def test_publish_crash_is_recovered_on_restart(tmp_path, monkeypatch):
    store = a.DocumentStore(tmp_path / 'report.docx'); store.init('report')
    real = a.publish
    def crash(*args): raise OSError('simulated disk publish failure')
    monkeypatch.setattr(a, 'publish', crash)
    with pytest.raises(OSError): store.apply_one(request(0))
    monkeypatch.setattr(a, 'publish', real)
    reopened = a.DocumentStore(store.path)
    assert reopened.status()['revision'] == 1
    assert reopened.path.read_bytes() == reopened.snapshot(1)
    assert reopened.apply_one(request(0))['revision'] == 1


def test_gate_repairs_do_not_unlock_until_rechecked(tmp_path, monkeypatch):
    fake_review(monkeypatch, pages=2)
    store = a.DocumentStore(tmp_path / 'report.docx'); store.init('report', mode='strict')
    for i in range(3): store.apply_one(request(i))
    with pytest.raises(ValueError, match='checkpoint'): store.apply_one(request(3))
    review = store.checkpoint()
    assert review['gate'] == 'visual_pending'
    assert review['pages'][0]['checks']['visual_review'] == 'pending'
    store.apply_one(request(3, type='replace_text', block_id='p-1', old='Paragraph', new='Updated'))
    assert store.status()['gate'] != 'open'
    with pytest.raises(ValueError, match='checkpoint'): store.apply_one(request(4))
    store.checkpoint(final=True)
    with pytest.raises(ValueError): store.finalize(4)
    store.visual_review(4, 1, True, 'human', 'Inspected page 1: no clipping, text and margins clear')
    store.visual_review(4, 2, True, 'human', 'Inspected page 2: no clipping, text and margins clear')
    assert store.finalize(4)['status'] == 'finalized'
    with pytest.raises(ValueError): store.visual_review(3, 1, True, 'human', 'stale')


def test_failed_visual_review_cannot_be_cleared_by_recheck(tmp_path, monkeypatch):
    fake_review(monkeypatch)
    store = a.DocumentStore(tmp_path / 'report.docx'); store.init('report')
    store.checkpoint(final=True)
    store.visual_review(0, 1, False, 'human', 'Heading visibly overlaps image')
    assert store.checkpoint(final=True)['gate'] == 'failed'
    with pytest.raises(ValueError): store.finalize(0)


def test_stable_template_ids_and_images(tmp_path):
    from docx import Document
    from PIL import Image
    template = tmp_path / 'template.docx'
    doc = Document(); doc.add_paragraph('Template paragraph'); doc.add_table(rows=1, cols=1).cell(0,0).text = 'Cell'; doc.save(template)
    store = a.DocumentStore(tmp_path / 'report.docx'); store.init('report', template)
    original_styles = a.o.read_parts(template)['word/styles.xml']
    ids = {p['text']: p['block_id'] for p in store.status()['blocks']}
    store.apply_one(request(0, type='append_paragraph', block_id='first', text='Inserted before original', after_block_id=ids['Template paragraph']))
    store.apply_one(request(1, type='replace_text', block_id=ids['Cell'], old='Cell', new='Updated cell'))
    assert a.o.read_parts(store.path)['word/styles.xml'] == original_styles
    assert any(b['block_id'] == ids['Cell'] and b['text'] == 'Updated cell' for b in store.status()['blocks'])
    Image.new('RGB', (200,100), 'blue').save(tmp_path / 'image.png')
    store.apply_one(request(2, type='insert_image', block_id='diagram', path='image.png', width_cm=5))
    assert a.o.validate_parts(a.o.read_parts(store.path))['ok']


def test_missing_renderer_blocks_and_corrupt_external_can_restore(tmp_path, monkeypatch):
    store = a.DocumentStore(tmp_path / 'report.docx'); store.init('report')
    monkeypatch.setattr(a.o, 'run_render', lambda *args: (_ for _ in ()).throw(RuntimeError('LibreOffice not installed')))
    assert store.checkpoint()['gate'] == 'failed'
    assert 'LibreOffice' in store.status()['render_error']
    store.path.write_bytes(b'broken ZIP')
    assert store.snapshot(0).startswith(b'PK')
    store.restore(0, 0, 'recover')
    assert store.status()['revision'] == 1


def test_parallel_processes_conflict(tmp_path):
    path = tmp_path / 'report.docx'; a.DocumentStore(path).init('report')
    jobs = []
    for i in range(2):
        payload = tmp_path / f'op-{i}.json'; payload.write_text(json.dumps(request(0, operation_id=f'parallel-{i}')))
        jobs.append(subprocess.Popen([sys.executable, str(SCRIPTS / 'office_docx.py'), 'apply-one', str(path), str(payload)], stdout=subprocess.PIPE, stderr=subprocess.PIPE))
    results = [job.communicate() for job in jobs]
    assert sorted(job.returncode for job in jobs) == [0, 1]
    assert a.DocumentStore(path).status()['revision'] == 1


def test_pdf_layout_failure_cases(tmp_path):
    import pymupdf as fitz
    spec = importlib.util.spec_from_file_location('review', SCRIPTS / 'page_review.py')
    review = importlib.util.module_from_spec(spec); spec.loader.exec_module(review)
    parts = a.o.default_package()
    parts, _, _ = a.mutate(parts, dict(type='append_paragraph', block_id='heading', text='Orphan heading', style='Heading1'), tmp_path)
    parts, _, _ = a.mutate(parts, dict(type='append_table', block_id='table', rows=[['first unique cell'], ['second unique cell']]), tmp_path)
    parts = a.o.apply_plan(parts, dict(operations=[dict(op='set_page', left_twips=100)]), tmp_path)
    doc = fitz.open()
    page = doc.new_page(); page.insert_text((72,72), 'first unique cell'); page.insert_text((72,800), 'Orphan heading')
    page = doc.new_page(); page.insert_text((72,72), 'second unique cell')
    pages = review.analyze_pages(parts, doc, a.o, final=True)
    assert pages[0]['checks']['heading_orphan']
    assert pages[0]['checks']['table_split'] and pages[1]['checks']['table_split']
    assert pages[0]['checks']['margins_abnormal']
    assert all(p['status'] == 'failed' for p in pages)


@pytest.mark.skipif(not shutil.which('libreoffice'), reason='LibreOffice is required for actual layout integration')
def test_real_libreoffice_incremental_chinese_report(tmp_path):
    store = a.DocumentStore(tmp_path / 'report.docx'); store.init('report')
    spec = importlib.util.spec_from_file_location('corpus', SCRIPTS.parents[1] / 'tests' / 'docx_corpus.py')
    corpus = importlib.util.module_from_spec(spec); spec.loader.exec_module(corpus)
    corpus.make_assets(tmp_path)
    for i, operation in enumerate(corpus.formatted_operations(sections=2)):
        if store.status()['gate'] != 'open':
            result = store.checkpoint()
            assert result['gate'] == 'open', result
        store.apply_one(request(i, **operation))
    result = store.checkpoint(final=True)
    assert result['pages'] and all(p['checks']['automatic_passed'] for p in result['pages']), result
    assert all(p['checks']['visual_review'] == 'pending' for p in result['pages'])


def test_image_overflow_and_missing_chinese_font_are_detected(tmp_path, monkeypatch):
    import pymupdf as fitz
    from PIL import Image
    from types import SimpleNamespace
    import io
    spec = importlib.util.spec_from_file_location('review', SCRIPTS / 'page_review.py')
    review = importlib.util.module_from_spec(spec); spec.loader.exec_module(review)
    parts = a.o.default_package()
    parts['word/styles.xml'] = parts['word/styles.xml'].replace(b'Noto Sans CJK SC', b'Missing Chinese Font')
    parts, _, _ = a.mutate(parts, dict(type='append_paragraph', block_id='body', text='中文业务说明'), tmp_path)
    image = io.BytesIO(); Image.new('RGB',(100,100),'blue').save(image,format='PNG')
    doc = fitz.open(); page = doc.new_page()
    page.insert_text((72,72), '中文业务说明', fontname='china-s')
    page.insert_image(fitz.Rect(-30,100,70,200), stream=image.getvalue())
    monkeypatch.setattr(review.subprocess, 'run', lambda *args,**kwargs: SimpleNamespace(stdout='Noto Sans CJK SC'))
    result = review.analyze_pages(parts, doc, a.o, final=True)[0]
    assert result['checks']['content_overflow']
    assert result['checks']['fonts_missing'] == ['Missing Chinese Font']
    assert result['status'] == 'failed'


def test_later_edit_invalidates_prior_page_visual_review(tmp_path, monkeypatch):
    def changing_render(docx, output, dpi=125):
        import pymupdf as fitz
        doc=fitz.open()
        for i in range(2):
            page=doc.new_page(); page.insert_text((72,72), 'Reflow ' + a.o.sha(docx)[:12])
            page.get_pixmap().save(output/f'page-{i+1:03d}.png')
        doc.save(output/'preview.pdf'); doc.close()
    monkeypatch.setattr(a.o,'run_render',changing_render)
    store=a.DocumentStore(tmp_path/'report.docx'); store.init('report')
    store.checkpoint()
    store.visual_review(0,1,True,'human','Test fixture: initial review record')
    store.apply_one(request(0))
    result=store.checkpoint()
    assert result['pages'][0]['checks']['visual_review']=='pending'
    assert result['pages'][0]['status']=='pending'


def test_libreoffice_timeout_is_persisted_without_losing_content(tmp_path, monkeypatch):
    store=a.DocumentStore(tmp_path/'report.docx'); store.init('report'); store.apply_one(request(0))
    data=store.path.read_bytes()
    def timeout(*args,**kwargs): raise subprocess.TimeoutExpired('libreoffice',120)
    monkeypatch.setattr(a.o.subprocess,'run',timeout)
    result=store.checkpoint()
    assert result['gate']=='failed' and 'timed out' in result['render_error']
    assert store.status()['revision']==1 and store.path.read_bytes()==data
